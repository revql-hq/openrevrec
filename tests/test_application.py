"""End-to-end tests use real portable workspaces and real SQLite transactions."""

from __future__ import annotations

import io
import sqlite3
from copy import deepcopy
from decimal import Decimal

import pytest
from openpyxl import load_workbook

from openrevrec.application import Application
from openrevrec.__main__ import main
from openrevrec.demo import load_demo
from openrevrec.interchange import export_bytes, import_bytes, preview_import_bytes, template_bytes
from openrevrec.posting import compare_journals
from openrevrec.reporting import build_review


def app(tmp_path):
    return Application.create(tmp_path / "Test company.orr", "Test company")


def test_cancellable_term_records_assessment_and_review_trigger_through_amendment_and_export(tmp_path):
    application = app(tmp_path)
    application.execute("create_customer", {"id": "cus_1", "name": "Customer"})
    contract = {
        "id": "con_1", "name": "Cancellable service", "customer_id": "cus_1",
        "start_date": "2026-01-01", "end_date": "2026-06-30",
        "consideration": [{"id": "price", "kind": "fixed", "amount": "600.00"}],
        "obligations": [{"id": "service", "name": "Service", "kind": "service", "ssp": "600.00", "method": "monthly", "start_date": "2026-01-01", "end_date": "2026-06-30"}],
        "term_basis": "cancellable", "term_assessment_rationale": "Six months are enforceable under the notice clause.",
        "term_reassessment_trigger": "Review on a cancellation or renewal notice.", "term_review_date": "2026-05-01",
    }
    with pytest.raises(ValueError, match="assessed-term rationale"):
        application.execute("create_contract", {**contract, "term_assessment_rationale": ""}, period="2026-01")
    with pytest.raises(ValueError, match="reassessment trigger"):
        application.execute("create_contract", {**contract, "term_reassessment_trigger": ""}, period="2026-01")
    with pytest.raises(ValueError, match="review date"):
        application.execute("create_contract", {**contract, "term_review_date": "2025-12-31"}, period="2026-01")
    with pytest.raises(ValueError, match="initial assessed accounting end"):
        application.execute("create_contract", {**contract, "obligations": [{**contract["obligations"][0], "end_date": "2026-12-31"}]}, period="2026-01")
    application.execute("create_contract", contract, period="2026-01")
    january = application.state(period="2026-01")
    assert january["contracts"][0]["term_basis"] == "cancellable"
    assert january["report"]["summary"]["revenue"] == "100.00"
    assert next(check for check in application.reports(period="2026-05")["checks"] if check["id"] == "term_reviews")["status"] == "review"
    application.execute("modify_contract", {
        "contract_id": "con_1", "effective_date": "2026-05-01", "treatment": "catch_up", "rationale": "Renewal terms now enforceable",
        "obligations": [{**contract["obligations"][0], "end_date": "2026-12-31"}],
        "term_basis": "evergreen", "term_assessment_rationale": "The renewal creates another enforceable six months.",
        "term_reassessment_trigger": "Review the next notice window.", "term_review_date": "2026-11-01",
    }, period="2026-05")
    state = application.state(period="2026-05")
    assert state["contracts"][0]["activities"][-1]["term_basis"] == "evergreen"
    assert next(check for check in application.reports(period="2026-05")["checks"] if check["id"] == "term_reviews")["status"] == "pass"
    assert application.state(period="2026-01")["report"]["summary"]["revenue"] == january["report"]["summary"]["revenue"]
    assert state["report"]["contracts"][0]["remaining_revenue"] == "350.00"
    book = load_workbook(io.BytesIO(export_bytes(state)))
    assert book["Term assessments"]["C2"].value == "cancellable"
    assert book["Term assessments"]["C3"].value == "evergreen"
    assert book["Term assessments"]["D3"].value == "2026-12-31"
    assert book["Term assessments"]["F3"].value == "Review the next notice window."
    book.close()
    application.execute("modify_contract", {
        "contract_id": "con_1", "effective_date": "2026-07-01", "treatment": "catch_up", "rationale": "The renewed service now has a fixed term",
        "obligations": [{**contract["obligations"][0], "end_date": "2026-12-31"}], "term_basis": "fixed",
    }, period="2026-07")
    assert next(check for check in application.reports(period="2026-11")["checks"] if check["id"] == "term_reviews")["status"] == "pass"


def test_unchanged_term_review_is_auditable_and_reschedules_without_changing_revenue(tmp_path):
    application = app(tmp_path)
    application.execute("create_customer", {"id": "cus_1", "name": "Customer"})
    contract = {
        "id": "con_1", "name": "Cancellable service", "customer_id": "cus_1",
        "start_date": "2026-01-01", "end_date": "2026-12-31",
        "consideration": [{"id": "price", "kind": "fixed", "amount": "1200.00"}],
        "obligations": [{"id": "service", "name": "Service", "kind": "service", "ssp": "1200.00", "method": "monthly", "start_date": "2026-01-01", "end_date": "2026-12-31"}],
        "term_basis": "cancellable", "term_assessment_rationale": "The notice clause makes this year enforceable.",
        "term_reassessment_trigger": "Review each notice window.", "term_review_date": "2026-05-01",
    }
    application.execute("create_contract", contract, period="2026-01")
    may_revenue = application.state(period="2026-05")["report"]["summary"]["revenue"]
    assert next(check for check in application.reports(period="2026-05")["checks"] if check["id"] == "term_reviews")["status"] == "review"
    review = {"contract_id": "con_1", "effective_date": "2026-05-05", "reviewer": "A. Accountant",
              "conclusion": "No cancellation notice; the assessed term is unchanged.",
              "support_memo": "Reviewed the executed agreement and notice register.", "next_review_date": "2026-08-01"}
    with pytest.raises(ValueError, match="next review date"):
        application.execute("record_term_review", {**review, "next_review_date": ""}, period="2026-05")
    with pytest.raises(ValueError, match="must follow"):
        application.execute("record_term_review", {**review, "next_review_date": "2026-05-05"}, period="2026-05")
    result = application.execute("record_term_review", review, period="2026-05")
    state = application.state(period="2026-05")
    assert state["report"]["summary"]["revenue"] == may_revenue
    assert state["term_reviews"][0]["change_set_id"] == result["result"]["change_set_id"]
    assert next(check for check in application.reports(period="2026-05")["checks"] if check["id"] == "term_reviews")["status"] == "pass"
    assert next(check for check in application.reports(period="2026-08")["checks"] if check["id"] == "term_reviews")["status"] == "review"
    book = load_workbook(io.BytesIO(export_bytes(state)))
    assert book["Term reviews"]["C2"].value == "A. Accountant"
    assert book["Term reviews"]["F2"].value == "2026-08-01"
    book.close()
    application.execute("modify_contract", {"contract_id": "con_1", "effective_date": "2026-06-01", "treatment": "catch_up",
        "rationale": "The notice process changed", "obligations": contract["obligations"],
        "term_reassessment_trigger": "Review the revised notice window.", "term_review_date": "2026-07-01"}, period="2026-06")
    assert next(check for check in application.reports(period="2026-07")["checks"] if check["id"] == "term_reviews")["status"] == "review"
    application.execute("close_period", {"period": "2026-05", "review_dispositions": accept_review_items(application, "2026-05")}, period="2026-05")
    with pytest.raises(ValueError, match="Reopen the closed period"):
        application.execute("record_term_review", {**review, "effective_date": "2026-05-10"}, period="2026-05")


def test_evergreen_term_assessment_imports_with_structured_columns(tmp_path):
    application = app(tmp_path)
    book = load_workbook(io.BytesIO(template_bytes()))
    def add(sheet_name, values):
        sheet = book[sheet_name]
        headers = [cell.value for cell in sheet[1]]
        sheet.append([values.get(header) for header in headers])
    add("Customers", {"id": "cus_1", "name": "Customer"})
    add("Contracts", {
        "id": "con_1", "customer_id": "cus_1", "name": "Renewing service", "start_date": "2026-01-01", "end_date": "2026-06-30",
        "term_basis": "evergreen", "term_assessment_rationale": "The first six months are enforceable.",
        "term_reassessment_trigger": "Review the renewal notice.", "term_review_date": "2026-05-01",
    })
    add("Consideration", {"contract_id": "con_1", "id": "price", "kind": "fixed", "amount": "600.00"})
    add("Obligations", {"contract_id": "con_1", "id": "service", "name": "Service", "kind": "service", "ssp": "600.00", "method": "monthly", "start_date": "2026-01-01", "end_date": "2026-06-30"})
    add("Amendments", {
        "source_id": "amend:renewal", "contract_id": "con_1", "effective_date": "2026-04-01", "treatment": "catch_up",
        "replace_obligations": "yes", "rationale": "Renewal notice changed the assessed term",
        "term_basis": "cancellable", "term_assessment_rationale": "The notice makes the year enforceable.",
        "term_reassessment_trigger": "Review the next cancellation window.", "term_review_date": "2026-11-01",
    })
    add("Amendment Obligations", {"amendment_source_id": "amend:renewal", "id": "service", "name": "Service", "kind": "service", "ssp": "600.00", "method": "monthly", "start_date": "2026-01-01", "end_date": "2026-12-31"})
    data = io.BytesIO()
    book.save(data)
    book.close()
    preview = preview_import_bytes(application, data.getvalue(), period="2026-01")
    assert preview["state"]["contracts"][0]["term_reassessment_trigger"] == "Review the renewal notice."
    assert preview["state"]["contracts"][0]["activities"][0]["term_basis"] == "cancellable"
    import_bytes(application, data.getvalue(), period="2026-01", expected_frontier=preview["frontier"], expected_hash=preview["result"]["file_hash"])
    assert application.state(period="2026-01")["contracts"][0]["term_basis"] == "evergreen"


def test_cli_workspace_setup_records_company_currency_and_account_defaults(tmp_path):
    path = tmp_path / "Configured company.orr"
    accounts = '{"revenue":"4100","deferred_revenue":"2310","contract_asset":"1310","billing_clearing":"1110"}'
    assert main(["init", str(path), "--name", "Configured company", "--currency", "EUR", "--account-effective-period", "2026-10", "--accounts-json", accounts]) == 0
    state = Application(path).state(period="2026-10")
    assert state["workspace"]["name"] == "Configured company"
    assert state["workspace"]["currency"] == "EUR"
    assert state["policy"]["accounts"] == {"revenue": "4100", "deferred_revenue": "2310", "contract_asset": "1310", "billing_clearing": "1110"}
    default_path = tmp_path / "Defaults.orr"
    assert main(["init", str(default_path), "--opening-period", "2026-07"]) == 0
    assert Application(default_path).state(period="2026-07")["policy"]["effective_period"] == "2026-07"


def seed(application):
    customer = application.execute("create_customer", {"id": "cus_1", "name": "Customer"})
    contract = {
        "id": "con_1", "name": "Annual service", "customer_id": customer["result"]["id"],
        "start_date": "2026-09-01", "end_date": "2027-08-31",
        "consideration": [{"id": "price_1", "label": "Fixed fee", "kind": "fixed", "amount": "12000.00"}],
        "obligations": [{"id": "pob_1", "name": "Service", "kind": "service", "ssp": "12000.00", "method": "exact_days", "start_date": "2026-09-01", "end_date": "2027-08-31"}],
    }
    application.execute("create_contract", contract, period="2026-09")
    return contract


def test_separate_contract_amendment_link_preserves_two_revenue_schedules(tmp_path):
    application = app(tmp_path)
    application.execute("create_customer", {"id": "cus_1", "name": "Customer"})
    application.execute("create_contract", {
        "id": "original", "customer_id": "cus_1", "name": "Annual service",
        "start_date": "2026-01-01", "end_date": "2026-12-31",
        "consideration": [{"id": "original_price", "kind": "fixed", "amount": "1200"}],
        "obligations": [{"id": "original_service", "name": "Annual service", "kind": "service", "ssp": "1200",
                         "method": "monthly", "start_date": "2026-01-01", "end_date": "2026-12-31"}],
    }, period="2026-01")
    application.execute("create_contract", {
        "id": "added", "customer_id": "cus_1", "name": "Additional service",
        "start_date": "2026-07-01", "end_date": "2026-09-30",
        "consideration": [{"id": "added_price", "kind": "fixed", "amount": "300"}],
        "obligations": [{"id": "added_service", "name": "Additional service", "kind": "service", "ssp": "300",
                         "method": "monthly", "start_date": "2026-07-01", "end_date": "2026-09-30"}],
    }, period="2026-07")
    link = {"contract_id": "original", "added_contract_id": "added", "effective_date": "2026-06-15",
            "additional_consideration": "300", "price_basis": "distinct_at_standalone_price",
            "original_terms_effect": "unchanged",
            "rationale": "The added service is distinct and the amendment increases price by its standalone price."}
    with pytest.raises(ValueError, match="original contract's remaining promises"):
        application.execute("link_modification_contract", {**link, "original_terms_effect": "repriced"}, period="2026-07")
    with pytest.raises(ValueError, match="standalone selling prices"):
        application.execute("link_modification_contract", {**link, "price_basis": ""}, period="2026-07")
    with pytest.raises(ValueError, match="initial transaction price"):
        application.execute("link_modification_contract", {**link, "additional_consideration": "250"}, period="2026-07")
    application.execute("link_modification_contract", link, period="2026-07")
    with pytest.raises(ValueError, match="already linked"):
        application.execute("link_modification_contract", link, period="2026-07")
    july = application.state(period="2026-07")
    support = july["report"]["modification_links"][0]
    assert (support["original_revenue"], support["added_revenue"], support["combined_revenue"]) == ("100.00", "100.00", "200.00")
    assert july["report"]["summary"]["revenue"] == "200.00"
    assert {(row["contract_id"], row["revenue"]) for row in july["report"]["schedule"] if row["period"] == "2026-07"} == {
        ("original", "100.00"), ("added", "100.00")}
    assert sum(row["debit_minor"] - row["credit_minor"] for row in july["report"]["journals"]) == 0
    assert not any("offsetting asset and deferred balances" in warning for warning in july["report"]["warnings"])
    application.execute("record_billing", {"contract_id": "added", "effective_date": "2026-07-01",
                                           "amount": "300.00", "reference": "INV-ADDED"}, period="2026-07")
    july = application.state(period="2026-07")
    assert (july["report"]["modification_links"][0]["original_contract_asset"],
            july["report"]["modification_links"][0]["added_deferred_revenue"]) == ("700.00", "200.00")
    assert any("separately accounted contracts have offsetting asset and deferred balances" in warning for warning in july["report"]["warnings"])
    assert any(item["contract_id"] == "original" and item["related_contract_id"] == "added"
               for item in application.reports(period="2026-07")["exceptions"]["warnings"])
    assert sum(row["debit_minor"] - row["credit_minor"] for row in july["report"]["journals"]) == 0
    with pytest.raises(ValueError, match="mixed-treatment review"):
        application.execute("modify_contract", {"contract_id": "original", "effective_date": "2026-06-15",
            "treatment": "prospective", "consideration": [{"id": "original_price", "kind": "fixed", "amount": "1080"}],
            "rationale": "Discount the remaining original service"}, period="2026-07")
    book = load_workbook(io.BytesIO(export_bytes(july)), read_only=True)
    assert book["Modification links"]["D2"].value == "2026-06-15"
    assert book["Modification links"]["L2"].value == 200
    book.close()


def test_warning_target_keeps_source_ids_when_contracts_have_the_same_name(tmp_path):
    application = app(tmp_path)
    seed(application)
    application.execute("create_contract", {
        "id": "con_2", "name": "Annual service", "customer_id": "cus_1",
        "start_date": "2026-09-01", "end_date": "2026-09-30",
        "consideration": [{"id": "delivery_price", "kind": "fixed", "amount": "100.00"}],
        "obligations": [{"id": "delivery", "name": "Delivery", "kind": "service", "ssp": "100.00",
                         "method": "point_in_time", "start_date": "2026-09-01", "end_date": "2026-09-30"}],
    }, period="2026-09")
    warnings = application.reports(period="2026-09")["exceptions"]["warnings"]
    assert any("satisfaction has not been recorded" in warning["message"] for warning in warnings)
    assert warnings[0]["contract_id"] == "con_2"
    assert warnings[0]["obligation_id"] == "delivery"
    application.execute("create_contract", {
        "id": "con_3", "name": "Annual service", "customer_id": "cus_1",
        "start_date": "2026-09-01", "end_date": "2026-09-30",
        "consideration": [{"id": "other_price", "kind": "fixed", "amount": "100.00"}],
        "obligations": [{"id": "other_delivery", "name": "Delivery", "kind": "service", "ssp": "100.00",
                         "method": "point_in_time", "start_date": "2026-09-01", "end_date": "2026-09-30"}],
    }, period="2026-09")
    review = application.reports(period="2026-09")
    duplicate_message = review["exceptions"]["warnings"]
    assert len(duplicate_message) == 2
    assert {item["contract_id"] for item in duplicate_message} == {"con_2", "con_3"}
    assert next(check for check in review["checks"] if check["id"] == "warnings")["count"] == 2
    book = load_workbook(io.BytesIO(export_bytes(application.state(period="2026-09"))), read_only=True)
    assert [row[0] for row in list(book["Warnings"].values)[1:]] == ["con_2", "con_3"]
    book.close()


def test_closed_warning_target_uses_accepted_snapshot_after_contract_rename(tmp_path):
    application = app(tmp_path)
    application.execute("create_customer", {"id": "cus_1", "name": "Customer"})
    application.execute("create_contract", {
        "id": "con_1", "name": "Delivery service", "customer_id": "cus_1",
        "start_date": "2026-09-01", "end_date": "2026-09-30",
        "consideration": [{"id": "price", "kind": "fixed", "amount": "100.00"}],
        "obligations": [{"id": "delivery", "name": "Delivery", "kind": "service", "ssp": "100.00",
                         "method": "point_in_time", "start_date": "2026-09-01", "end_date": "2026-09-30"}],
    }, period="2026-09")
    application.execute("close_period", {"period": "2026-09", "review_dispositions": accept_review_items(application, "2026-09")}, period="2026-09")
    application.execute("edit_details", {"entity_id": "con_1", "name": "Renamed delivery service"}, period="2026-10")
    closed = application.state(period="2026-09")["report"]
    assert closed["warnings"][0].startswith("Delivery service / Delivery:")
    assert closed["warning_details"][0]["contract_id"] == "con_1"
    assert application.reports(period="2026-09")["exceptions"]["warnings"][0]["obligation_id"] == "delivery"
    older = application.state(period="2026-09")
    older["report"].pop("warning_details")  # An older accepted checkpoint contains message text only.
    assert build_review(older)["exceptions"]["warnings"][0]["contract_id"] is None
    book = load_workbook(io.BytesIO(export_bytes(older)), read_only=True)
    assert book["Warnings"]["D2"].value == closed["warnings"][0]
    book.close()


def test_separate_contract_amendment_link_imports_from_reviewed_workbook(tmp_path):
    application = app(tmp_path)
    book = load_workbook(io.BytesIO(template_bytes()))
    book["Customers"].append(["cus_1", "Customer"])
    book["Contracts"].append(["original", "cus_1", "Original", "2026-01-01", "2026-12-31"])
    book["Contracts"].append(["added", "cus_1", "Added", "2026-07-01", "2026-09-30"])
    book["Consideration"].append(["original", "price", "Service", "fixed", "1200"])
    book["Consideration"].append(["added", "new_price", "Additional", "fixed", "300"])
    book["Obligations"].append(["original", "service", "Service", "service", "1200", "monthly", "2026-01-01", "2026-12-31"])
    book["Obligations"].append(["added", "new_service", "Additional", "service", "300", "monthly", "2026-07-01", "2026-09-30"])
    book["Modification Links"].append(["original", "added", "2026-06-15", "300", "distinct_at_standalone_price",
                                        "unchanged", "Distinct added service at standalone price", "amendment:add-service"])
    data = io.BytesIO()
    book.save(data)
    book.close()
    preview = preview_import_bytes(application, data.getvalue(), period="2026-07")
    assert preview["state"]["report"]["modification_links"][0]["combined_revenue"] == "200.00"
    accepted = import_bytes(application, data.getvalue(), period="2026-07",
                            expected_frontier=preview["frontier"], expected_hash=preview["result"]["file_hash"])
    assert accepted["result"]["controls"]["counts"]["link_modification_contract"] == 1
    assert accepted["state"]["report"]["modification_links"][0]["combined_revenue"] == "200.00"


def test_opening_position_starts_from_reconciled_legacy_balances_without_backfilling_revenue(tmp_path):
    application = app(tmp_path)
    application.execute("create_customer", {"id": "cus_1", "name": "Customer"})
    application.execute("create_contract", {
        "id": "con_1", "name": "Migrated service", "customer_id": "cus_1",
        "start_date": "2026-01-01", "end_date": "2026-12-31",
        "consideration": [{"id": "price_1", "kind": "fixed", "amount": "1200.00"}],
        "obligations": [{"id": "pob_1", "name": "Service", "kind": "service", "ssp": "1200.00", "method": "exact_days", "start_date": "2026-01-01", "end_date": "2026-12-31"}],
    }, period="2026-09")
    opening = {
        "contract_id": "con_1", "effective_date": "2026-09-01", "billed_to_date": "1000.00",
        "contract_asset": "0.00", "deferred_revenue": "200.00",
        "source_name": "Approved legacy workbook", "rationale": "August close agreed to GL",
        "opening_obligations": [{"obligation_id": "pob_1", "recognized_to_date": "800.00"}],
    }
    with pytest.raises(ValueError, match="reconcile"):
        application.execute("record_opening_position", {**opening, "deferred_revenue": "199.00"}, period="2026-09")
    preview = application.preview("record_opening_position", opening, period="2026-09")
    assert preview["state"]["report"]["contracts"][0]["beginning_deferred_revenue"] == "200.00"
    assert application.state(period="2026-08")["report"]["contracts"]  # No accepted cutover yet.
    application.execute("record_opening_position", opening, period="2026-09")
    assert application.state(period="2026-08")["report"]["contracts"] == []
    september = application.state(period="2026-09")["report"]
    contract = september["contracts"][0]
    assert contract["recognized_to_date"] == "898.36"
    assert contract["revenue"] == "98.36"
    assert contract["billed_to_date"] == "1000.00"
    assert contract["deferred_revenue"] == "101.64"
    assert september["schedule"] and all(row["period"] >= "2026-09" for row in september["schedule"])
    assert application.reports("main", "2026-09")["recognition_coverage"][0]["obligation_gaps"] == []
    assert sum(row["debit_minor"] for row in september["journals"]) == sum(row["credit_minor"] for row in september["journals"])
    with pytest.raises(ValueError, match="pre-cutover"):
        application.execute("record_billing", {"contract_id": "con_1", "effective_date": "2026-08-10", "amount": "10.00"}, period="2026-09")
    export = load_workbook(io.BytesIO(export_bytes(application.state(period="2026-09"))))
    assert export["Opening positions"]["B2"].value == "2026-09-01"
    export.close()
    application.execute("set_policy", {"effective_period": "2026-10", "account_overrides": {"contracts": {}, "obligations": {"con_1": {"pob_1": "4100"}}}, "rationale": "Product revenue mapping"}, period="2026-10")
    assert any(row["role"] == "revenue" and row["account"] == "4100" for row in application.state(period="2026-10")["report"]["journals"])


def test_opening_position_preserves_manual_progress_measure(tmp_path):
    application = app(tmp_path)
    application.execute("create_customer", {"id": "cus_1", "name": "Customer"})
    application.execute("create_contract", {
        "id": "con_1", "name": "Implementation", "customer_id": "cus_1",
        "start_date": "2026-01-01", "end_date": "2026-12-31",
        "consideration": [{"id": "price_1", "kind": "fixed", "amount": "1000.00"}],
        "obligations": [{"id": "pob_1", "name": "Build", "kind": "implementation", "ssp": "1000.00", "method": "progress", "start_date": "2026-01-01", "end_date": "2026-12-31"}],
    }, period="2026-09")
    application.execute("record_opening_position", {
        "contract_id": "con_1", "effective_date": "2026-09-01", "billed_to_date": "500.00",
        "contract_asset": "0.00", "deferred_revenue": "0.00", "source_name": "Legacy POC schedule",
        "rationale": "August signed completion estimate", "opening_obligations": [{"obligation_id": "pob_1", "recognized_to_date": "500.00", "measure": "50"}],
    }, period="2026-09")
    application.execute("record_progress", {"contract_id": "con_1", "obligation_id": "pob_1", "effective_date": "2026-09-20", "percentage": "75"}, period="2026-09")
    assert application.state(period="2026-09")["report"]["contracts"][0]["revenue"] == "250.00"


def test_cutover_balance_uses_current_account_without_a_prior_account_transfer(tmp_path):
    application = app(tmp_path)
    application.execute("create_customer", {"id": "cus_1", "name": "Customer"})
    application.execute("create_contract", {
        "id": "con_1", "name": "Migrated service", "customer_id": "cus_1",
        "start_date": "2026-01-01", "end_date": "2026-12-31", "cutover_date": "2026-09-01",
        "consideration": [{"id": "price_1", "kind": "fixed", "amount": "1200.00"}],
        "obligations": [{"id": "pob_1", "name": "Service", "kind": "service", "ssp": "1200.00", "method": "exact_days", "start_date": "2026-01-01", "end_date": "2026-12-31"}],
    }, period="2026-09")
    application.execute("set_policy", {
        "effective_period": "2026-09", "accounts": {"deferred_revenue": "2310"},
        "account_transition": "transfer", "rationale": "September chart change",
    }, period="2026-09")
    application.execute("record_opening_position", {
        "contract_id": "con_1", "effective_date": "2026-09-01", "billed_to_date": "1000.00",
        "contract_asset": "0.00", "deferred_revenue": "200.00",
        "source_name": "Legacy workbook", "rationale": "August GL tie-out",
        "opening_obligations": [{"obligation_id": "pob_1", "recognized_to_date": "800.00"}],
    }, period="2026-09")
    report = application.state(period="2026-09")["report"]
    assert report["account_transitions"] == []
    assert not any(":transfer" in row["id"] for row in report["journals"])
    deferred = [row for row in report["journals"] if row["role"] == "deferred_revenue"]
    assert len(deferred) == 1 and deferred[0]["account"] == "2310"
    assert sum(row["debit_minor"] - row["credit_minor"] for row in report["journals"]) == 0


def test_opening_position_preserves_cumulative_units_and_rejects_completed_shortfall(tmp_path):
    application = app(tmp_path)
    application.execute("create_customer", {"id": "cus_1", "name": "Customer"})
    application.execute("create_contract", {
        "id": "con_1", "name": "Finite unit package", "customer_id": "cus_1", "start_date": "2026-01-01", "end_date": "2026-12-31",
        "consideration": [{"id": "price_1", "kind": "fixed", "amount": "1000.00"}],
        "obligations": [{"id": "pob_1", "name": "Units", "kind": "service", "ssp": "1000.00", "method": "usage", "total_units": "100", "start_date": "2026-01-01", "end_date": "2026-12-31"}],
    }, period="2026-09")
    opening = {"contract_id": "con_1", "effective_date": "2026-09-01", "billed_to_date": "500.00", "contract_asset": "0.00", "deferred_revenue": "100.00",
               "source_name": "Legacy usage ledger", "rationale": "August usage and GL tie-out", "opening_obligations": [{"obligation_id": "pob_1", "recognized_to_date": "400.00", "measure": "40"}]}
    with pytest.raises(ValueError, match="fully satisfied"):
        application.execute("record_opening_position", {**opening, "opening_obligations": [{"obligation_id": "pob_1", "recognized_to_date": "400.00", "measure": "100"}]}, period="2026-09")
    application.execute("record_opening_position", opening, period="2026-09")
    application.execute("record_usage", {"contract_id": "con_1", "obligation_id": "pob_1", "effective_date": "2026-09-10", "quantity": "10"}, period="2026-09")
    assert application.state(period="2026-09")["report"]["contracts"][0]["revenue"] == "100.00"


def test_import_opening_position_previews_and_commits_as_one_reviewed_batch(tmp_path):
    application = app(tmp_path)
    book = load_workbook(io.BytesIO(template_bytes()))
    book["Customers"].append(["cus_1", "Customer"])
    book["Contracts"].append(["con_1", "cus_1", "Migrated", "2026-01-01", "2026-12-31", "Legacy contract"])
    book["Consideration"].append(["con_1", "price_1", "Fixed", "fixed", "1200.00"])
    book["Obligations"].append(["con_1", "pob_1", "Service", "service", "1200.00", "exact_days", "2026-01-01", "2026-12-31"])
    book["Opening Positions"].append(["legacy:con_1", "con_1", "2026-09-01", "1000.00", "0.00", "200.00", "Legacy workbook", "August GL tie-out"])
    book["Opening Obligations"].append(["legacy:con_1", "pob_1", "800.00"])
    data = io.BytesIO()
    book.save(data)
    book.close()
    preview = preview_import_bytes(application, data.getvalue(), period="2026-09")
    assert preview["result"]["controls"]["counts"]["record_opening_position"] == 1
    assert preview["state"]["contracts"][0]["cutover_date"] == "2026-09-01"
    assert preview["state"]["report"]["contracts"][0]["deferred_revenue"] == "101.64"
    assert application.state(period="2026-09")["contracts"] == []
    import_bytes(application, data.getvalue(), period="2026-09", expected_frontier=preview["frontier"], expected_hash=preview["result"]["file_hash"])
    assert application.state(period="2026-09")["report"]["contracts"][0]["deferred_revenue"] == "101.64"


def test_declared_cutover_blocks_close_until_opening_and_allows_new_legacy_contract_after_prior_close(tmp_path):
    application = app(tmp_path)
    application.execute("create_customer", {"id": "cus_1", "name": "Customer"})
    application.execute("create_contract", {
        "id": "con_future", "name": "Future", "customer_id": "cus_1", "start_date": "2026-09-01", "end_date": "2026-09-30",
        "consideration": [{"id": "price_future", "kind": "fixed", "amount": "100.00"}],
        "obligations": [{"id": "pob_future", "name": "Service", "kind": "service", "ssp": "100.00", "method": "exact_days", "start_date": "2026-09-01", "end_date": "2026-09-30"}],
    }, period="2026-08")
    application.execute("close_period", {"period": "2026-08", "review_dispositions": accept_review_items(application, "2026-08")}, period="2026-08")
    application.execute("create_contract", {
        "id": "con_migrated", "name": "Legacy", "customer_id": "cus_1", "start_date": "2026-01-01", "end_date": "2026-12-31", "cutover_date": "2026-09-01",
        "consideration": [{"id": "price_migrated", "kind": "fixed", "amount": "1200.00"}],
        "obligations": [{"id": "pob_migrated", "name": "Service", "kind": "service", "ssp": "1200.00", "method": "exact_days", "start_date": "2026-01-01", "end_date": "2026-12-31"}],
    }, period="2026-09")
    with pytest.raises(ValueError, match="opening position before post-cutover activity"):
        application.execute("record_billing", {"contract_id": "con_migrated", "effective_date": "2026-09-10", "amount": "10.00"}, period="2026-09")
    assert application.state(period="2026-08")["report"]["closed"] is True
    review = application.reports("main", "2026-09")
    assert next(check for check in review["checks"] if check["id"] == "cutover")["status"] == "block"
    with pytest.raises(ValueError, match="Declared cutovers"):
        application.execute("close_period", {"period": "2026-09", "review_dispositions": accept_review_items(application, "2026-09")}, period="2026-09")
    application.execute("record_opening_position", {
        "contract_id": "con_migrated", "effective_date": "2026-09-01", "billed_to_date": "1000.00", "contract_asset": "0.00", "deferred_revenue": "200.00",
        "source_name": "Legacy workbook", "rationale": "August signed balance", "opening_obligations": [{"obligation_id": "pob_migrated", "recognized_to_date": "800.00"}],
    }, period="2026-09")
    application.execute("record_billing", {"contract_id": "con_migrated", "effective_date": "2026-09-10", "amount": "10.00"}, period="2026-09")
    assert next(check for check in application.reports("main", "2026-09")["checks"] if check["id"] == "cutover")["status"] == "pass"


def test_import_migrated_contract_after_closed_legacy_month_preserves_checkpoint(tmp_path):
    application = app(tmp_path)
    application.execute("create_customer", {"id": "cus_1", "name": "Customer"})
    application.execute("create_contract", {
        "id": "con_future", "name": "Future", "customer_id": "cus_1", "start_date": "2026-09-01", "end_date": "2026-09-30",
        "consideration": [{"id": "price_future", "kind": "fixed", "amount": "100.00"}],
        "obligations": [{"id": "pob_future", "name": "Service", "kind": "service", "ssp": "100.00", "method": "exact_days", "start_date": "2026-09-01", "end_date": "2026-09-30"}],
    }, period="2026-08")
    application.execute("close_period", {"period": "2026-08", "review_dispositions": accept_review_items(application, "2026-08")}, period="2026-08")
    before = application.state(period="2026-08")["report"]
    book = load_workbook(io.BytesIO(template_bytes()))
    book["Contracts"].append(["con_legacy", "cus_1", "Migrated", "2026-01-01", "2026-12-31"])
    book["Consideration"].append(["con_legacy", "price_legacy", "Fixed", "fixed", "1200.00"])
    book["Obligations"].append(["con_legacy", "pob_legacy", "Service", "service", "1200.00", "exact_days", "2026-01-01", "2026-12-31"])
    book["Opening Positions"].append(["legacy:con_legacy", "con_legacy", "2026-09-01", "1000.00", "0.00", "200.00", "Legacy close", "August GL tie-out"])
    book["Opening Obligations"].append(["legacy:con_legacy", "pob_legacy", "800.00"])
    data = io.BytesIO()
    book.save(data)
    book.close()
    reviewed = preview_import_bytes(application, data.getvalue(), period="2026-09")
    import_bytes(application, data.getvalue(), period="2026-09", expected_frontier=reviewed["frontier"], expected_hash=reviewed["result"]["file_hash"])
    after = application.state(period="2026-08")["report"]
    assert after["summary"] == before["summary"]
    assert after["journals"] == before["journals"]
    assert after["closed"] is True
    assert application.state(period="2026-09")["report"]["contracts"][1]["recognized_to_date"] == "898.36"


def test_material_right_exercise_command_and_workbook_schedule_later_delivery(tmp_path):
    application = app(tmp_path)
    application.execute("create_customer", {"id": "cus_1", "name": "Customer"})
    contract = {
        "id": "con_1", "name": "Renewal option", "customer_id": "cus_1", "start_date": "2026-01-01", "end_date": "2026-06-30",
        "consideration": [{"id": "fixed", "kind": "fixed", "amount": "600.00"}],
        "obligations": [{"id": "right", "name": "Renewal right", "kind": "material_right", "ssp": "600.00", "method": "point_in_time", "start_date": "2026-01-01", "end_date": "2026-06-30", "exercise_start": "2026-04-01", "exercise_end": "2026-06-30"}],
    }
    application.execute("create_contract", contract, period="2026-01")
    exercise = {"contract_id": "con_1", "obligation_id": "right", "effective_date": "2026-06-15", "delivery_method": "monthly", "delivery_start": "2026-07-01", "delivery_end": "2026-12-31", "rationale": "Renewal option exercised for second-half service"}
    application.execute("record_right_exercise", exercise, period="2026-06")
    assert application.state(period="2026-06")["report"]["summary"]["revenue"] == "0.00"
    assert application.state(period="2026-07")["report"]["summary"]["revenue"] == "100.00"
    assert application.state(period="2026-12")["report"]["summary"]["recognized_to_date"] == "600.00"
    exported = load_workbook(io.BytesIO(export_bytes(application.state(period="2026-07"))), read_only=True)
    assert exported["Right exercises"]["D2"].value == "monthly"
    exported.close()

    imported = app(tmp_path / "import")
    book = load_workbook(io.BytesIO(template_bytes()))
    book["Customers"].append(["cus_1", "Customer"])
    book["Contracts"].append(["con_1", "cus_1", "Renewal option", "2026-01-01", "2026-06-30"])
    book["Consideration"].append(["con_1", "fixed", "Fee", "fixed", "600.00"])
    book["Obligations"].append(["con_1", "right", "Renewal right", "material_right", "600.00", "point_in_time", "2026-01-01", "2026-06-30", None, "2026-04-01", "2026-06-30"])
    book["Right Exercises"].append(["con_1", "right", "2026-06-15", "monthly", "2026-07-01", "2026-12-31", "Reviewed exercise", "rights:1"])
    data = io.BytesIO()
    book.save(data)
    book.close()
    preview = preview_import_bytes(imported, data.getvalue(), period="2026-07")
    assert preview["result"]["controls"]["counts"]["record_right_exercise"] == 1
    assert preview["state"]["report"]["summary"]["revenue"] == "100.00"
    import_bytes(imported, data.getvalue(), period="2026-07", expected_frontier=preview["frontier"], expected_hash=preview["result"]["file_hash"])
    assert imported.state(period="2026-12")["report"]["summary"]["recognized_to_date"] == "600.00"


def test_linked_renewal_keeps_old_right_allocation_separate_from_new_price(tmp_path):
    application = app(tmp_path)
    application.execute("create_customer", {"id": "cus_1", "name": "Customer"})
    application.execute("create_contract", {
        "id": "original", "name": "Original option", "customer_id": "cus_1", "start_date": "2026-01-01", "end_date": "2026-06-30",
        "consideration": [{"id": "fee", "kind": "fixed", "amount": "600.00"}],
        "obligations": [{"id": "right", "name": "Renewal right", "kind": "material_right", "ssp": "600.00", "method": "point_in_time",
                         "start_date": "2026-01-01", "end_date": "2026-06-30", "exercise_start": "2026-04-01", "exercise_end": "2026-06-30"}],
    }, period="2026-01")
    application.execute("create_contract", {
        "id": "renewal", "name": "Renewal service", "customer_id": "cus_1", "start_date": "2026-07-01", "end_date": "2026-12-31",
        "consideration": [{"id": "new_fee", "kind": "fixed", "amount": "300.00"}],
        "obligations": [{"id": "service", "name": "Renewal service", "kind": "service", "ssp": "300.00", "method": "monthly",
                         "start_date": "2026-07-01", "end_date": "2026-12-31"}],
    }, period="2026-06")
    link = {"contract_id": "original", "obligation_id": "right", "renewal_contract_id": "renewal",
            "additional_consideration": "300.00", "price_basis": "new_consideration_only", "rationale": "New fee is separate from the carried right"}
    with pytest.raises(ValueError, match="Exercise the material right"):
        application.execute("link_renewal_contract", link, period="2026-06")
    application.execute("record_right_exercise", {"contract_id": "original", "obligation_id": "right", "effective_date": "2026-06-15",
        "delivery_method": "monthly", "delivery_start": "2026-07-01", "delivery_end": "2026-12-31",
        "rationale": "Option elected for six months"}, period="2026-06")
    with pytest.raises(ValueError, match="only new consideration"):
        application.execute("link_renewal_contract", {**link, "price_basis": ""}, period="2026-06")
    with pytest.raises(ValueError, match="initial transaction price"):
        application.execute("link_renewal_contract", {**link, "additional_consideration": "900.00"}, period="2026-06")
    application.execute("record_billing", {"contract_id": "original", "effective_date": "2026-01-01", "amount": "600.00"}, period="2026-01")
    application.execute("record_billing", {"contract_id": "renewal", "effective_date": "2026-06-15", "amount": "300.00"}, period="2026-06")
    application.execute("link_renewal_contract", link, period="2026-07")
    with pytest.raises(ValueError, match="already has a linked renewal"):
        application.execute("link_renewal_contract", link, period="2026-07")
    june = application.state(period="2026-06")["report"]
    july = application.state(period="2026-07")["report"]
    assert june["summary"]["revenue"] == "0.00"
    assert july["summary"]["transaction_price"] == "900.00"
    assert july["summary"]["revenue"] == "150.00"
    assert july["summary"]["deferred_revenue"] == "750.00"
    support = july["renewal_links"][0]
    assert (support["original_right_allocation"], support["initial_new_consideration"], support["combined_consideration"],
            support["right_revenue"], support["renewal_revenue"], support["combined_revenue"]) == (
                "600.00", "300.00", "900.00", "100.00", "50.00", "150.00")
    assert sum(row["debit_minor"] - row["credit_minor"] for row in july["journals"]) == 0
    exported = load_workbook(io.BytesIO(export_bytes(application.state(period="2026-07"))), read_only=True)
    assert exported["Renewal links"]["J2"].value == 900
    assert exported["Renewal links"]["M2"].value == 150
    exported.close()


def test_linked_renewal_imports_after_contracts_and_exercise(tmp_path):
    application = app(tmp_path)
    book = load_workbook(io.BytesIO(template_bytes()))
    book["Customers"].append(["customer", "Customer"])
    book["Contracts"].append(["original", "customer", "Original", "2026-01-01", "2026-06-30"])
    book["Contracts"].append(["renewal", "customer", "Renewal", "2026-07-01", "2026-12-31"])
    book["Consideration"].append(["original", "old", "Old fee", "fixed", "600.00"])
    book["Consideration"].append(["renewal", "new", "New fee", "fixed", "300.00"])
    book["Obligations"].append(["original", "right", "Renewal right", "material_right", "600.00", "point_in_time", "2026-01-01", "2026-06-30", None, "2026-04-01", "2026-06-30"])
    book["Obligations"].append(["renewal", "service", "Renewal service", "service", "300.00", "monthly", "2026-07-01", "2026-12-31"])
    book["Right Exercises"].append(["original", "right", "2026-06-15", "monthly", "2026-07-01", "2026-12-31", "Option elected", "exercise:1"])
    book["Renewal Links"].append(["original", "right", "renewal", "300.00", "new_consideration_only", "New fee excludes old right", "link:1"])
    data = io.BytesIO()
    book.save(data)
    book.close()
    preview = preview_import_bytes(application, data.getvalue(), period="2026-07")
    assert preview["result"]["controls"]["counts"]["link_renewal_contract"] == 1
    assert preview["state"]["report"]["renewal_links"][0]["combined_consideration"] == "900.00"
    import_bytes(application, data.getvalue(), period="2026-07", expected_frontier=preview["frontier"], expected_hash=preview["result"]["file_hash"])
    assert application.state(period="2026-07")["report"]["summary"]["revenue"] == "150.00"


def test_material_right_exercise_cannot_rewrite_an_accepted_expiry_month(tmp_path):
    application = app(tmp_path)
    application.execute("create_customer", {"id": "cus_1", "name": "Customer"})
    application.execute("create_contract", {
        "id": "con_1", "name": "Option", "customer_id": "cus_1", "start_date": "2026-01-01", "end_date": "2026-06-30",
        "consideration": [{"id": "fixed", "kind": "fixed", "amount": "600.00"}],
        "obligations": [{"id": "right", "name": "Option", "kind": "material_right", "ssp": "600.00", "method": "point_in_time", "start_date": "2026-01-01", "end_date": "2026-06-30", "exercise_start": "2026-04-01", "exercise_end": "2026-06-30"}],
    }, period="2026-06")
    application.execute("close_period", {"period": "2026-06", "review_dispositions": accept_review_items(application, "2026-06")}, period="2026-06")
    with pytest.raises(ValueError, match="closed"):
        application.execute("record_right_exercise", {"contract_id": "con_1", "obligation_id": "right", "effective_date": "2026-06-15", "delivery_method": "monthly", "delivery_start": "2026-07-01", "delivery_end": "2026-12-31", "rationale": "Late source exercise"}, period="2026-06")
    assert application.state(period="2026-06")["report"]["summary"]["revenue"] == "600.00"


def test_specific_variable_and_credit_allocation_reconcile_and_reassessment_stays_targeted(tmp_path):
    application = app(tmp_path)
    application.execute("create_customer", {"id": "cus_1", "name": "Customer"})
    base = {
        "id": "con_1", "name": "Hosting and implementation", "customer_id": "cus_1", "start_date": "2026-01-01", "end_date": "2026-12-31",
        "consideration": [
            {"id": "fixed", "kind": "fixed", "amount": "1000.00"},
            {"id": "bonus", "kind": "variable", "amount": "300.00", "included_amount": "200.00", "allocation_scope": "specific", "target_obligation_ids": ["implementation"], "allocation_rationale": "Reviewed delivery bonus relates specifically to implementation"},
            {"id": "credit", "kind": "credit", "amount": "-100.00", "allocation_scope": "specific", "target_obligation_ids": ["hosting"], "allocation_rationale": "Reviewed hosting-only price concession"},
        ],
        "obligations": [
            {"id": "implementation", "name": "Implementation", "kind": "implementation", "ssp": "1000.00", "method": "exact_days", "start_date": "2026-01-01", "end_date": "2026-12-31"},
            {"id": "hosting", "name": "Hosting", "kind": "service", "ssp": "1000.00", "method": "exact_days", "start_date": "2026-01-01", "end_date": "2026-12-31"},
        ],
    }
    with pytest.raises(ValueError, match="rationale"):
        application.execute("create_contract", {**base, "consideration": [base["consideration"][0], {**base["consideration"][1], "allocation_rationale": ""}, base["consideration"][2]]}, period="2026-02")
    application.execute("create_contract", base, period="2026-02")
    january = application.state(period="2026-01")["report"]["contracts"][0]
    assert january["transaction_price"] == "1100.00"
    assert {row["obligation_id"]: row["amount"] for row in january["allocation"]} == {"implementation": "700.00", "hosting": "400.00"}
    application.execute("reassess_variable_consideration", {"contract_id": "con_1", "component_id": "bonus", "effective_date": "2026-02-01", "included_amount": "300.00", "rationale": "Bonus now fully included"}, period="2026-02")
    february = application.state(period="2026-02")["report"]
    assert {row["obligation_id"]: row["amount"] for row in february["contracts"][0]["allocation"]} == {"implementation": "800.00", "hosting": "400.00"}
    assert sum(Decimal(row["amount"]) for row in february["contracts"][0]["allocation"]) == Decimal("1200.00")
    assert all(row["obligation_id"] == "implementation" for row in february["catch_ups"])
    assert sum(row["debit_minor"] for row in february["journals"]) == sum(row["credit_minor"] for row in february["journals"])
    exported = load_workbook(io.BytesIO(export_bytes(application.state(period="2026-02"))))
    assert exported["Component allocation"]["F3"].value == "specific"
    exported.close()
    application.execute("modify_contract", {"contract_id": "con_1", "effective_date": "2026-03-01", "treatment": "prospective", "rationale": "Revised price for remaining distinct service", "consideration": base["consideration"]}, period="2026-03")
    march = application.state(period="2026-03")["report"]
    assert march["contracts"][0]["transaction_price"] == "1100.00"
    assert not [row for row in march["catch_ups"] if row["period"] == "2026-03"]
    assert application.state(period="2026-02")["report"]["summary"]["revenue"] == february["summary"]["revenue"]
    assert sum(Decimal(row["amount"]) for row in march["contracts"][0]["allocation"]) == Decimal("1100.00")


def test_specific_allocation_import_uses_structured_target_ids(tmp_path):
    application = app(tmp_path)
    book = load_workbook(io.BytesIO(template_bytes()))
    book["Customers"].append(["cus_1", "Customer"])
    book["Contracts"].append(["con_1", "cus_1", "Two deliverables", "2026-01-01", "2026-12-31"])
    book["Consideration"].append(["con_1", "fixed", "Base", "fixed", "1000.00"])
    book["Consideration"].append(["con_1", "bonus", "Delivery bonus", "variable", "200.00", "200.00", None, None, None, "Amount fully included", "specific", "pob_a", "Bonus belongs to the first deliverable"])
    for identifier in ("pob_a", "pob_b"):
        book["Obligations"].append(["con_1", identifier, identifier, "service", "1000.00", "exact_days", "2026-01-01", "2026-12-31"])
    data = io.BytesIO()
    book.save(data)
    book.close()
    preview = preview_import_bytes(application, data.getvalue(), period="2026-01")
    allocation = preview["state"]["report"]["contracts"][0]["allocation"]
    assert {item["obligation_id"]: item["amount"] for item in allocation} == {"pob_a": "700.00", "pob_b": "500.00"}
    assert preview["state"]["report"]["contracts"][0]["allocation_components"][1]["target_obligation_ids"] == ["pob_a"]


def test_prospective_modification_preserves_earned_targeted_bonus_and_allocates_only_remaining_service(tmp_path):
    application = app(tmp_path)
    application.execute("create_customer", {"id": "cus_1", "name": "Customer"})
    fixed = {"id": "fixed", "kind": "fixed", "amount": "1200.00"}
    bonus = {"id": "bonus", "kind": "variable", "amount": "180.00", "included_amount": "120.00",
             "allocation_scope": "specific", "target_obligation_ids": ["a"],
             "allocation_rationale": "The bonus relates specifically to service A."}
    obligations = [{"id": identifier, "name": identifier, "kind": "service", "ssp": "600.00", "method": "monthly",
                    "start_date": "2026-01-01", "end_date": "2026-12-31"} for identifier in ("a", "b")]
    application.execute("create_contract", {"id": "con_1", "name": "Two services", "customer_id": "cus_1",
        "start_date": "2026-01-01", "end_date": "2026-12-31", "consideration": [fixed, bonus], "obligations": obligations}, period="2026-01")
    june = application.state(period="2026-06")["report"]
    assert june["summary"]["recognized_to_date"] == "660.00"
    application.execute("close_period", {"period": "2026-06", "review_dispositions": accept_review_items(application, "2026-06")}, period="2026-06")
    application.execute("modify_contract", {"contract_id": "con_1", "effective_date": "2026-07-01",
        "treatment": "prospective", "rationale": "Distinct remaining months with a revised base and bonus",
        "consideration": [{**fixed, "amount": "1500.00"}, {**bonus, "included_amount": "180.00"}]}, period="2026-07")
    july = application.state(period="2026-07")["report"]
    assert july["summary"]["revenue"] == "170.00"
    assert july["catch_ups"] == []
    assert {row["obligation_id"]: row["amount"] for row in july["contracts"][0]["allocation"]} == {"a": "930.00", "b": "750.00"}
    assert next(row for row in july["contracts"][0]["allocation_components"] if row["component_id"] == "bonus")["recognized_to_date"] == "80.00"
    exported = load_workbook(io.BytesIO(export_bytes(application.state(period="2026-07"))))
    assert exported["Component allocation"]["J3"].value == 80
    exported.close()
    assert application.state(period="2026-06")["report"]["summary"]["revenue"] == june["summary"]["revenue"]
    december = application.state(period="2026-12")["report"]
    assert december["summary"]["recognized_to_date"] == december["summary"]["transaction_price"] == "1680.00"
    assert sum(Decimal(row["revenue"]) for row in december["schedule"]) == Decimal("1680.00")
    assert sum(row["debit_minor"] for row in december["journals"]) == sum(row["credit_minor"] for row in december["journals"])
    with pytest.raises(ValueError, match="original variable promise"):
        application.execute("reassess_variable_consideration", {"contract_id": "con_1", "component_id": "bonus",
            "effective_date": "2026-08-01", "included_amount": "170.00", "rationale": "Later estimate changed"}, period="2026-08")
    with pytest.raises(ValueError, match="pre-modification targeted variable"):
        application.execute("modify_contract", {"contract_id": "con_1", "effective_date": "2026-08-01",
            "treatment": "catch_up", "rationale": "Later estimate changed",
            "consideration": [{**fixed, "amount": "1500.00"}, {**bonus, "included_amount": "170.00"}]}, period="2026-08")
    september = application.state(period="2026-09")["report"]["summary"]["recognized_to_date"]
    application.execute("modify_contract", {"contract_id": "con_1", "effective_date": "2026-10-01",
        "treatment": "prospective", "rationale": "A further price change for distinct remaining months",
        "consideration": [{**fixed, "amount": "1600.00"}, {**bonus, "included_amount": "180.00"}]}, period="2026-10")
    assert application.state(period="2026-09")["report"]["summary"]["recognized_to_date"] == september
    revised = application.state(period="2026-12")["report"]
    assert revised["summary"]["recognized_to_date"] == revised["summary"]["transaction_price"] == "1780.00"
    assert sum(Decimal(row["revenue"]) for row in revised["schedule"]) == Decimal("1780.00")


def test_prospective_targeted_allocation_rejects_unattributed_prior_adjustment(tmp_path):
    application = app(tmp_path)
    application.execute("create_customer", {"id": "cus_1", "name": "Customer"})
    bonus = {"id": "bonus", "kind": "variable", "amount": "120.00", "included_amount": "120.00",
             "allocation_scope": "specific", "target_obligation_ids": ["service"],
             "allocation_rationale": "Bonus relates to service."}
    application.execute("create_contract", {"id": "con_1", "name": "Service", "customer_id": "cus_1",
        "start_date": "2026-01-01", "end_date": "2026-12-31",
        "consideration": [{"id": "fixed", "kind": "fixed", "amount": "1200.00"}, bonus],
        "obligations": [{"id": "service", "name": "Service", "kind": "service", "ssp": "1200.00", "method": "monthly",
                         "start_date": "2026-01-01", "end_date": "2026-12-31"}]}, period="2026-01")
    application.execute("record_adjustment", {"contract_id": "con_1", "obligation_id": "service",
        "effective_date": "2026-02-01", "amount": "10.00", "rationale": "Manual correction"}, period="2026-02")
    with pytest.raises(ValueError, match="component attribution"):
        application.execute("modify_contract", {"contract_id": "con_1", "effective_date": "2026-03-01",
            "treatment": "prospective", "rationale": "Remaining distinct services", "consideration": [{"id": "fixed", "kind": "fixed", "amount": "1300.00"}, bonus]}, period="2026-03")


def test_prospective_modification_carries_a_targeted_credit_without_reversing_earned_discount(tmp_path):
    application = app(tmp_path)
    application.execute("create_customer", {"id": "cus_1", "name": "Customer"})
    credit = {"id": "credit", "kind": "credit", "amount": "-120.00", "allocation_scope": "specific",
              "target_obligation_ids": ["a"], "allocation_rationale": "The concession applies to service A."}
    obligations = [{"id": identifier, "name": identifier, "kind": "service", "ssp": "600.00", "method": "monthly",
                    "start_date": "2026-01-01", "end_date": "2026-12-31"} for identifier in ("a", "b")]
    application.execute("create_contract", {"id": "con_1", "name": "Two services", "customer_id": "cus_1",
        "start_date": "2026-01-01", "end_date": "2026-12-31",
        "consideration": [{"id": "fixed", "kind": "fixed", "amount": "1200.00"}, credit],
        "obligations": obligations}, period="2026-01")
    assert application.state(period="2026-06")["report"]["summary"]["recognized_to_date"] == "540.00"
    application.execute("modify_contract", {"contract_id": "con_1", "effective_date": "2026-07-01",
        "treatment": "prospective", "rationale": "Higher base price for distinct remaining months",
        "consideration": [{"id": "fixed", "kind": "fixed", "amount": "1500.00"}, credit]}, period="2026-07")
    july = application.state(period="2026-07")["report"]
    assert july["summary"]["revenue"] == "140.00"
    assert next(row for row in july["contracts"][0]["allocation_components"] if row["component_id"] == "credit")["recognized_to_date"] == "-70.00"
    assert {row["obligation_id"]: row["amount"] for row in july["contracts"][0]["allocation"]} == {"a": "630.00", "b": "750.00"}
    december = application.state(period="2026-12")["report"]
    assert december["summary"]["recognized_to_date"] == december["summary"]["transaction_price"] == "1380.00"


def test_prospective_targeted_component_uses_remaining_ssp_when_one_target_is_complete(tmp_path):
    application = app(tmp_path)
    application.execute("create_customer", {"id": "cus_1", "name": "Customer"})
    bonus = {"id": "bonus", "kind": "variable", "amount": "120.00", "included_amount": "120.00",
             "allocation_scope": "specific", "target_obligation_ids": ["a", "b"],
             "allocation_rationale": "Bonus relates to both services."}
    application.execute("create_contract", {"id": "con_1", "name": "Two services", "customer_id": "cus_1",
        "start_date": "2026-01-01", "end_date": "2026-12-31",
        "consideration": [{"id": "fixed", "kind": "fixed", "amount": "1200.00"}, bonus],
        "obligations": [
            {"id": "a", "name": "A", "kind": "service", "ssp": "600.00", "method": "monthly", "start_date": "2026-01-01", "end_date": "2026-06-30"},
            {"id": "b", "name": "B", "kind": "service", "ssp": "600.00", "method": "monthly", "start_date": "2026-01-01", "end_date": "2026-12-31"},
        ]}, period="2026-01")
    assert application.state(period="2026-06")["report"]["summary"]["recognized_to_date"] == "990.00"
    application.execute("modify_contract", {"contract_id": "con_1", "effective_date": "2026-07-01",
        "treatment": "prospective", "rationale": "Repriced distinct remaining service B",
        "consideration": [{"id": "fixed", "kind": "fixed", "amount": "1500.00"}, bonus]}, period="2026-07")
    july = application.state(period="2026-07")["report"]
    assert july["summary"]["revenue"] == "105.00"
    assert {row["obligation_id"]: row["amount"] for row in july["contracts"][0]["allocation"]} == {"a": "660.00", "b": "960.00"}
    assert july["contracts"][0]["allocation_components"][1]["recognized_to_date"] == "95.00"


def test_variable_bonus_targets_one_service_month_and_later_reassessment_catches_up(tmp_path):
    application = app(tmp_path)
    application.execute("create_customer", {"id": "cus_1", "name": "Customer"})
    contract = {
        "id": "con_1", "name": "Monthly series with June bonus", "customer_id": "cus_1",
        "start_date": "2026-01-01", "end_date": "2026-12-31",
        "consideration": [
            {"id": "fixed", "kind": "fixed", "amount": "1200.00"},
            {"id": "bonus", "kind": "variable", "amount": "150.00", "included_amount": "120.00", "allocation_scope": "specific",
             "target_obligation_ids": ["service"], "target_period": "2026-06",
             "allocation_rationale": "The variable bonus relates only to the distinct June service in the series and meets the allocation objective."},
        ],
        "obligations": [{"id": "service", "name": "Monthly service", "kind": "service", "ssp": "1200.00", "method": "monthly", "start_date": "2026-01-01", "end_date": "2026-12-31"}],
    }
    with pytest.raises(ValueError, match="Target service month"):
        application.execute("create_contract", {**contract, "consideration": [contract["consideration"][0], {**contract["consideration"][1], "target_period": "2027-06"}]}, period="2026-01")
    with pytest.raises(ValueError, match="one specifically targeted"):
        application.execute("create_contract", {**contract, "consideration": [contract["consideration"][0], {**contract["consideration"][1], "target_obligation_ids": ["service", "other"]}]}, period="2026-01")
    with pytest.raises(ValueError, match="time-based service"):
        application.execute("create_contract", {**contract, "obligations": [{**contract["obligations"][0], "method": "point_in_time"}]}, period="2026-01")
    with pytest.raises(ValueError, match="reviewed cutover allocation"):
        application.execute("create_contract", {**contract, "cutover_date": "2026-07-01"}, period="2026-01")
    application.execute("create_contract", contract, period="2026-01")
    assert application.state(period="2026-01")["report"]["summary"]["revenue"] == "100.00"
    june = application.state(period="2026-06")
    assert june["report"]["summary"]["revenue"] == "220.00"
    assert june["report"]["contracts"][0]["allocation"][0]["amount"] == "1320.00"
    assert june["report"]["contracts"][0]["allocation_components"][1]["target_period"] == "2026-06"
    exported = load_workbook(io.BytesIO(export_bytes(june)))
    assert exported["Component allocation"]["I3"].value == "2026-06"
    exported.close()
    application.execute("reassess_variable_consideration", {"contract_id": "con_1", "component_id": "bonus", "effective_date": "2026-07-01", "included_amount": "150.00", "rationale": "Final June outcome confirmed."}, period="2026-07")
    july = application.state(period="2026-07")["report"]
    assert july["summary"]["revenue"] == "130.00"
    assert july["contracts"][0]["allocation"][0]["amount"] == "1350.00"
    assert [(row["obligation_id"], row["catch_up"]) for row in july["catch_ups"]] == [("service", "30.00")]
    assert application.state(period="2026-06")["report"]["summary"]["revenue"] == "220.00"
    assert sum(row["debit_minor"] for row in july["journals"]) == sum(row["credit_minor"] for row in july["journals"])
    with pytest.raises(ValueError, match="already satisfied services"):
        application.execute("modify_contract", {"contract_id": "con_1", "effective_date": "2026-08-01", "treatment": "prospective", "rationale": "Remove bonus", "consideration": [contract["consideration"][0]]}, period="2026-08")
    application.execute("modify_contract", {"contract_id": "con_1", "effective_date": "2026-08-01", "treatment": "catch_up", "rationale": "Bonus reversed after final contract review", "consideration": [contract["consideration"][0]]}, period="2026-08")
    august = application.state(period="2026-08")["report"]
    assert august["summary"]["revenue"] == "-50.00"
    assert august["contracts"][0]["allocation"][0]["amount"] == "1200.00"
    assert [(row["obligation_id"], row["catch_up"]) for row in august["catch_ups"] if row["period"] == "2026-08"] == [("service", "-150.00")]
    assert application.state(period="2026-09")["report"]["summary"]["revenue"] == "100.00"
    december = application.state(period="2026-12")["report"]
    assert december["summary"]["recognized_to_date"] == december["summary"]["transaction_price"] == "1200.00"
    assert sum(Decimal(row["revenue"]) for row in december["schedule"]) == Decimal("1200.00")


def test_prospective_modification_keeps_earned_service_month_bonus_and_rejects_repricing_it(tmp_path):
    application = app(tmp_path)
    application.execute("create_customer", {"id": "cus_1", "name": "Customer"})
    base = {"id": "fixed", "kind": "fixed", "amount": "1200.00"}
    bonus = {"id": "bonus", "kind": "variable", "amount": "120.00", "included_amount": "120.00",
             "allocation_scope": "specific", "target_obligation_ids": ["service"], "target_period": "2026-06",
             "allocation_rationale": "The bonus belongs to the distinct June service."}
    application.execute("create_contract", {"id": "con_1", "name": "Monthly service", "customer_id": "cus_1",
        "start_date": "2026-01-01", "end_date": "2026-12-31", "consideration": [base, bonus],
        "obligations": [{"id": "service", "name": "Service", "kind": "service", "ssp": "1200.00", "method": "monthly",
                         "start_date": "2026-01-01", "end_date": "2026-12-31"}]}, period="2026-01")
    assert application.state(period="2026-06")["report"]["summary"]["revenue"] == "220.00"
    application.execute("modify_contract", {"contract_id": "con_1", "effective_date": "2026-07-01",
        "treatment": "prospective", "rationale": "Higher base fee for distinct remaining months",
        "consideration": [{**base, "amount": "1500.00"}, bonus]}, period="2026-07")
    assert application.state(period="2026-07")["report"]["summary"]["revenue"] == "150.00"
    assert next(row for row in application.state(period="2026-07")["report"]["contracts"][0]["allocation_components"] if row["component_id"] == "bonus")["recognized_to_date"] == "120.00"
    assert application.state(period="2026-06")["report"]["summary"]["revenue"] == "220.00"
    december = application.state(period="2026-12")["report"]
    assert december["summary"]["recognized_to_date"] == december["summary"]["transaction_price"] == "1620.00"
    with pytest.raises(ValueError, match="pre-modification targeted variable"):
        application.execute("modify_contract", {"contract_id": "con_1", "effective_date": "2026-08-01",
            "treatment": "prospective", "rationale": "Reduce the completed June bonus",
            "consideration": [{**base, "amount": "1500.00"}, {**bonus, "included_amount": "90.00"}]}, period="2026-08")


def test_post_modification_june_bonus_reassessment_catches_up_original_month(tmp_path):
    application = app(tmp_path)
    application.execute("create_customer", {"id": "cus_1", "name": "Customer"})
    fixed = {"id": "fixed", "kind": "fixed", "amount": "1200.00"}
    bonus = {"id": "bonus", "kind": "variable", "amount": "150.00", "included_amount": "120.00",
             "allocation_scope": "specific", "target_obligation_ids": ["service"], "target_period": "2026-06",
             "allocation_rationale": "The bonus belongs to the June service."}
    application.execute("create_contract", {"id": "con_1", "name": "Monthly service", "customer_id": "cus_1",
        "start_date": "2026-01-01", "end_date": "2026-12-31", "consideration": [fixed, bonus],
        "obligations": [{"id": "service", "name": "Service", "kind": "service", "ssp": "1200.00",
                         "method": "monthly", "start_date": "2026-01-01", "end_date": "2026-12-31"}]}, period="2026-01")
    application.execute("close_period", {"period": "2026-06", "review_dispositions": accept_review_items(application, "2026-06")}, period="2026-06")
    application.execute("modify_contract", {"contract_id": "con_1", "effective_date": "2026-07-01",
        "treatment": "prospective", "rationale": "Higher base fee for distinct remaining months",
        "consideration": [{**fixed, "amount": "1500.00"}, bonus]}, period="2026-07")
    application.execute("reassess_variable_consideration", {"contract_id": "con_1", "component_id": "bonus",
        "effective_date": "2026-08-01", "included_amount": "150.00", "rationale": "Final June outcome confirmed"}, period="2026-08")
    june = application.state(period="2026-06")["report"]
    august = application.state(period="2026-08")["report"]
    assert june["summary"]["revenue"] == "220.00"
    assert august["summary"]["revenue"] == "180.00"
    assert [(row["obligation_id"], row["catch_up"]) for row in august["catch_ups"]] == [("service", "30.00")]
    assert august["contracts"][0]["allocation_components"][1]["recognized_to_date"] == "150.00"
    assert [(row["component_id"], row["obligation_id"], row["allocated_change"], row["recognized_to_date"])
            for row in august["contracts"][0]["original_promise_changes"]] == [("bonus", "service", "30.00", "30.00")]
    book = load_workbook(io.BytesIO(export_bytes(application.state(period="2026-08"))))
    assert book["Original promise changes"]["H2"].value == 30
    book.close()
    application.execute("reassess_variable_consideration", {"contract_id": "con_1", "component_id": "bonus",
        "effective_date": "2026-09-01", "included_amount": "90.00", "rationale": "Final settlement reduced the June outcome"}, period="2026-09")
    september = application.state(period="2026-09")["report"]
    assert september["summary"]["revenue"] == "90.00"
    assert [(row["obligation_id"], row["catch_up"]) for row in september["catch_ups"] if row["period"] == "2026-09"] == [("service", "-60.00")]
    december = application.state(period="2026-12")["report"]
    assert december["summary"]["recognized_to_date"] == december["summary"]["transaction_price"] == "1590.00"


def test_post_modification_relative_variable_change_uses_original_obligations_and_survives_next_amendment(tmp_path):
    application = app(tmp_path)
    application.execute("create_customer", {"id": "cus_1", "name": "Customer"})
    fixed = {"id": "fixed", "kind": "fixed", "amount": "1200.00"}
    bonus = {"id": "bonus", "kind": "variable", "amount": "180.00", "included_amount": "120.00"}
    obligations = [
        {"id": "a", "name": "A", "kind": "service", "ssp": "600.00", "method": "monthly", "start_date": "2026-01-01", "end_date": "2026-06-30"},
        {"id": "b", "name": "B", "kind": "service", "ssp": "600.00", "method": "monthly", "start_date": "2026-07-01", "end_date": "2026-12-31"},
    ]
    application.execute("create_contract", {"id": "con_1", "name": "Two services", "customer_id": "cus_1",
        "start_date": "2026-01-01", "end_date": "2026-12-31", "consideration": [fixed, bonus],
        "obligations": obligations}, period="2026-01")
    application.execute("modify_contract", {"contract_id": "con_1", "effective_date": "2026-07-01",
        "treatment": "prospective", "rationale": "Reprice the remaining distinct service",
        "consideration": [{**fixed, "amount": "1500.00"}, bonus],
        "obligations": [{**obligations[1], "ssp": "1200.00"}]}, period="2026-07")
    application.execute("reassess_variable_consideration", {"contract_id": "con_1", "component_id": "bonus",
        "effective_date": "2026-08-01", "included_amount": "180.00", "rationale": "Final original-contract estimate"}, period="2026-08")
    august = application.state(period="2026-08")["report"]
    assert august["summary"]["revenue"] == "200.00"
    assert {(row["obligation_id"], row["catch_up"]) for row in august["catch_ups"]} == {("a", "30.00"), ("b", "5.00")}
    assert {row["obligation_id"]: row["allocated_change"] for row in august["contracts"][0]["original_promise_changes"]} == {"a": "30.00", "b": "30.00"}
    assert {row["obligation_id"]: row["amount"] for row in august["contracts"][0]["allocation"]} == {"a": "690.00", "b": "990.00"}
    application.execute("modify_contract", {"contract_id": "con_1", "effective_date": "2026-10-01",
        "treatment": "prospective", "rationale": "Second price change for remaining B service",
        "consideration": [{**fixed, "amount": "1800.00"}, {**bonus, "included_amount": "180.00"}]}, period="2026-10")
    october = application.state(period="2026-10")["report"]
    assert october["summary"]["revenue"] == "265.00"
    december = application.state(period="2026-12")["report"]
    assert december["summary"]["recognized_to_date"] == december["summary"]["transaction_price"] == "1980.00"
    assert sum(Decimal(row["revenue"]) for row in december["schedule"]) == Decimal("1980.00")
    with pytest.raises(ValueError, match="later amendment cannot replace"):
        application.execute("modify_contract", {"contract_id": "con_1", "effective_date": "2026-11-01",
            "treatment": "catch_up", "rationale": "Reprice the original bonus without attribution",
            "consideration": [{**fixed, "amount": "1800.00"}, {**bonus, "included_amount": "170.00"}]}, period="2026-11")


def test_post_modification_reassessment_rejects_changed_original_service_path(tmp_path):
    application = app(tmp_path)
    application.execute("create_customer", {"id": "cus_1", "name": "Customer"})
    variable = {"id": "bonus", "kind": "variable", "amount": "180.00", "included_amount": "120.00"}
    obligations = [{"id": "service", "name": "Service", "kind": "service", "ssp": "1200.00",
                    "method": "monthly", "start_date": "2026-01-01", "end_date": "2026-12-31"}]
    application.execute("create_contract", {"id": "con_1", "name": "Service", "customer_id": "cus_1",
        "start_date": "2026-01-01", "end_date": "2026-12-31", "consideration": [variable],
        "obligations": obligations}, period="2026-01")
    application.execute("modify_contract", {"contract_id": "con_1", "effective_date": "2026-07-01",
        "treatment": "prospective", "rationale": "Extend the remaining distinct service",
        "obligations": [{**obligations[0], "end_date": "2027-03-31"}]}, period="2026-07")
    with pytest.raises(ValueError, match="original variable promise"):
        application.execute("reassess_variable_consideration", {"contract_id": "con_1", "component_id": "bonus",
            "effective_date": "2026-08-01", "included_amount": "150.00", "rationale": "Estimate changed"}, period="2026-08")


def test_repeated_original_promise_reassessments_allocate_cumulative_pennies(tmp_path):
    application = app(tmp_path)
    application.execute("create_customer", {"id": "cus_1", "name": "Customer"})
    fixed = {"id": "fixed", "kind": "fixed", "amount": "1200.00"}
    bonus = {"id": "bonus", "kind": "variable", "amount": "180.00", "included_amount": "120.00"}
    obligations = [{"id": identifier, "name": identifier, "kind": "service", "ssp": "600.00", "method": "monthly",
                    "start_date": "2026-01-01", "end_date": "2026-12-31"} for identifier in ("a", "b")]
    application.execute("create_contract", {"id": "con_1", "name": "Two services", "customer_id": "cus_1",
        "start_date": "2026-01-01", "end_date": "2026-12-31", "consideration": [fixed, bonus],
        "obligations": obligations}, period="2026-01")
    application.execute("modify_contract", {"contract_id": "con_1", "effective_date": "2026-07-01",
        "treatment": "prospective", "rationale": "Revised remaining service price",
        "consideration": [{**fixed, "amount": "1500.00"}, bonus]}, period="2026-07")
    for month, included in (("2026-08", "120.01"), ("2026-09", "120.02")):
        application.execute("reassess_variable_consideration", {"contract_id": "con_1", "component_id": "bonus",
            "effective_date": f"{month}-01", "included_amount": included, "rationale": "Updated original bonus estimate"}, period=month)
    allocations = {row["obligation_id"]: row["amount"] for row in application.state(period="2026-12")["report"]["contracts"][0]["allocation"]}
    assert allocations == {"a": "810.01", "b": "810.01"}


def test_period_target_import_and_cutover_boundary(tmp_path):
    application = app(tmp_path)
    book = load_workbook(io.BytesIO(template_bytes()))
    def append(sheet_name, values):
        sheet = book[sheet_name]
        headers = [cell.value for cell in sheet[1]]
        sheet.append([values.get(header) for header in headers])
    append("Customers", {"id": "cus_1", "name": "Customer"})
    append("Contracts", {"id": "con_1", "customer_id": "cus_1", "name": "Series", "start_date": "2026-01-01", "end_date": "2026-12-31"})
    append("Consideration", {"contract_id": "con_1", "id": "fixed", "kind": "fixed", "amount": "1200.00"})
    append("Consideration", {"contract_id": "con_1", "id": "bonus", "kind": "variable", "amount": "120.00", "included_amount": "120.00", "allocation_scope": "specific", "target_obligation_ids": "service", "target_period": "2026-06", "allocation_rationale": "June outcome relates to the distinct June service."})
    append("Obligations", {"contract_id": "con_1", "id": "service", "name": "Service", "kind": "service", "ssp": "1200.00", "method": "monthly", "start_date": "2026-01-01", "end_date": "2026-12-31"})
    data = io.BytesIO()
    book.save(data)
    book.close()
    preview = preview_import_bytes(application, data.getvalue(), period="2026-06")
    assert preview["state"]["report"]["summary"]["revenue"] == "220.00"
    import_bytes(application, data.getvalue(), period="2026-06", expected_frontier=preview["frontier"], expected_hash=preview["result"]["file_hash"])
    assert application.state(period="2026-06")["contracts"][0]["consideration"][1]["target_period"] == "2026-06"
    with pytest.raises(ValueError, match="reviewed cutover allocation"):
        application.execute("record_opening_position", {"contract_id": "con_1", "effective_date": "2026-07-01", "source_name": "Legacy", "rationale": "Migration", "billed_to_date": "0.00", "contract_asset": "0.00", "deferred_revenue": "0.00", "opening_obligations": [{"obligation_id": "service", "recognized_to_date": "720.00"}]}, period="2026-07")


def test_retargeted_service_month_moves_catch_up_between_obligations(tmp_path):
    application = app(tmp_path)
    application.execute("create_customer", {"id": "cus_1", "name": "Customer"})
    fixed = {"id": "fixed", "kind": "fixed", "amount": "1200.00"}
    bonus = {"id": "bonus", "kind": "variable", "amount": "120.00", "included_amount": "120.00", "allocation_scope": "specific", "target_obligation_ids": ["a"], "target_period": "2026-06", "allocation_rationale": "June service outcome relates to A."}
    obligations = [{"id": identifier, "name": identifier, "kind": "service", "ssp": "600.00", "method": "monthly", "start_date": "2026-01-01", "end_date": "2026-12-31"} for identifier in ("a", "b")]
    application.execute("create_contract", {"id": "con_1", "name": "Two series", "customer_id": "cus_1", "start_date": "2026-01-01", "end_date": "2026-12-31", "consideration": [fixed, bonus], "obligations": obligations}, period="2026-01")
    assert {row["obligation_id"]: row["revenue"] for row in application.state(period="2026-06")["report"]["schedule"] if row["period"] == "2026-06"} == {"a": "170.00", "b": "50.00"}
    application.execute("modify_contract", {"contract_id": "con_1", "effective_date": "2026-07-01", "treatment": "catch_up", "rationale": "Reviewed conclusion now attributes June bonus to series B", "consideration": [fixed, {**bonus, "target_obligation_ids": ["b"], "allocation_rationale": "June service outcome relates to B."}]}, period="2026-07")
    july = application.state(period="2026-07")["report"]
    assert july["summary"]["revenue"] == "100.00"
    assert {row["obligation_id"]: row["catch_up"] for row in july["catch_ups"] if row["period"] == "2026-07"} == {"a": "-120.00", "b": "120.00"}
    assert {row["obligation_id"]: row["amount"] for row in july["contracts"][0]["allocation"]} == {"a": "600.00", "b": "720.00"}


def accept_review_items(application, period):
    review = application.reports("main", period)
    return {check["id"]: {"disposition": "accepted", "reason": f"Reviewed {check['label']} against source support."} for check in review["checks"] if check["status"] == "review"}


def test_commands_persist_exact_history_and_are_idempotent(tmp_path):
    application = app(tmp_path)
    seed(application)
    first = application.execute("record_billing", {"contract_id": "con_1", "effective_date": "2026-09-01", "amount": "12000.00", "reference": "INV-1"}, idempotency_key="invoice-1", period="2026-09")
    second = application.execute("record_billing", {"contract_id": "con_1", "effective_date": "2026-09-01", "amount": "12000.00", "reference": "INV-1"}, idempotency_key="invoice-1", period="2026-09")
    assert second["result"]["replayed"] is True
    assert first["result"]["change_set_id"] == second["result"]["change_set_id"]
    assert second["state"]["report"]["summary"]["billings"] == "12000.00"
    with pytest.raises(ValueError, match="different input"):
        application.execute("record_billing", {"contract_id": "con_1", "effective_date": "2026-09-01", "amount": "1.00"}, idempotency_key="invoice-1")


def test_billing_credit_links_the_original_invoice_and_does_not_change_price(tmp_path):
    application = app(tmp_path)
    seed(application)
    original = application.execute("record_billing", {"contract_id": "con_1", "effective_date": "2026-09-01", "amount": "200.00", "reference": "INV-1"}, period="2026-09")["result"]["change_set_id"]
    with pytest.raises(ValueError, match="original invoice"):
        application.execute("record_billing", {"contract_id": "con_1", "effective_date": "2026-09-15", "amount": "-50.00", "reference": "CM-1", "rationale": "Invoice correction"}, period="2026-09")
    credit = application.execute("record_billing", {"contract_id": "con_1", "effective_date": "2026-09-15", "amount": "-50.00", "reference": "CM-1", "applies_to_change_set_id": original, "rationale": "Invoice correction"}, period="2026-09")
    assert credit["state"]["report"]["summary"]["billings"] == "150.00"
    assert credit["state"]["report"]["summary"]["transaction_price"] == "12000.00"
    with pytest.raises(ValueError, match="only on negative billing"):
        application.execute("record_billing", {"contract_id": "con_1", "effective_date": "2026-09-16", "amount": "50.00", "applies_to_change_set_id": original}, period="2026-09")


def test_scenario_billing_credit_link_points_to_applied_invoice(tmp_path):
    application = app(tmp_path)
    seed(application)
    scenario = application.execute("create_scenario", {"name": "Billing credit review"})["result"]["id"]
    invoice = application.execute("record_billing", {"contract_id": "con_1", "effective_date": "2026-09-01", "amount": "100.00", "reference": "INV-S"}, scenario_id=scenario, period="2026-09")["result"]["change_set_id"]
    application.execute("record_billing", {"contract_id": "con_1", "effective_date": "2026-09-05", "amount": "-10.00", "applies_to_change_set_id": invoice, "rationale": "Credit 10"}, scenario_id=scenario, period="2026-09")
    accepted = application.execute("apply_scenario", {"scenario_id": scenario}, scenario_id=scenario, period="2026-09")["state"]
    billings = [row for row in accepted["change_sets"] if row["command"] == "record_billing"]
    assert len(billings) == 2
    assert next(row for row in billings if row["payload"]["amount"] == "-10.00")["payload"]["applies_to_change_set_id"] == next(row for row in billings if row["payload"]["amount"] == "100.00")["id"]


def test_preview_rolls_back_every_write(tmp_path):
    application = app(tmp_path)
    seed(application)
    before = application.state(period="2026-09")
    preview = application.preview("record_billing", {"contract_id": "con_1", "effective_date": "2026-09-01", "amount": "12000.00"}, period="2026-09")
    after = application.state(period="2026-09")
    assert preview["state"]["report"]["summary"]["billings"] == "12000.00"
    assert after["frontier"] == before["frontier"]
    assert after["report"]["summary"]["billings"] == "0.00"


def test_scenario_isolation_compare_and_apply(tmp_path):
    application = app(tmp_path)
    seed(application)
    scenario = application.execute("create_scenario", {"name": "Proposed price change"})["result"]["id"]
    application.execute("modify_contract", {
        "contract_id": "con_1", "effective_date": "2026-10-01", "treatment": "prospective",
        "consideration": [{"id": "price_1", "label": "Fixed fee", "kind": "fixed", "amount": "18000.00"}],
        "rationale": "Customer expanded the remaining service scope.",
    }, scenario_id=scenario, period="2026-10")
    main = application.state("main", "2026-10")
    proposed = application.state(scenario, "2026-10")
    assert main["report"]["summary"]["transaction_price"] == "12000.00"
    assert proposed["report"]["summary"]["transaction_price"] == "18000.00"
    assert application.compare(scenario, "2026-10")["affected_periods"]
    application.execute("apply_scenario", {"scenario_id": scenario}, period="2026-10")
    accepted = application.state("main", "2026-10")
    assert accepted["report"]["summary"]["transaction_price"] == "18000.00"
    assert next(s for s in accepted["scenarios"] if s["id"] == scenario)["status"] == "applied"
    applied = [c for c in accepted["change_sets"] if c["command"] == "modify_contract"]
    assert applied[0]["originating_change_set_id"]


def test_prebuilt_reports_reconcile_balances_and_surface_scenario_impact(tmp_path):
    application = app(tmp_path)
    seed(application)
    application.execute("record_billing", {"contract_id": "con_1", "effective_date": "2026-09-01", "amount": "12000.00"}, period="2026-09")
    scenario = application.execute("create_scenario", {"name": "Price review"})["result"]["id"]
    application.execute("modify_contract", {
        "contract_id": "con_1", "effective_date": "2026-10-01", "treatment": "prospective",
        "consideration": [{"id": "price_1", "label": "Fixed fee", "kind": "fixed", "amount": "18000.00"}],
        "rationale": "Review expanded scope.",
    }, scenario_id=scenario, period="2026-10")

    review = application.reports("main", "2026-09")
    row = review["rollforward"][0]
    assert row["opening_contract_asset"] == "0.00"
    assert row["opening_deferred_revenue"] == "0.00"
    assert row["closing_deferred_revenue"] == "11013.70"
    assert review["recognition_coverage"][0]["status"] == "covered"
    assert review["scenario_impacts"][0]["name"] == "Price review"
    assert review["scenario_impacts"][0]["affected_periods"]
    assert next(check for check in review["checks"] if check["id"] == "scenarios")["status"] == "pass"
    assert next(check for check in application.reports("main", "2026-10")["checks"] if check["id"] == "scenarios")["status"] == "review"
    assert next(check for check in review["checks"] if check["id"] == "journal")["status"] == "pass"


def test_close_readiness_flags_entries_recorded_after_period_end(tmp_path):
    application = app(tmp_path)
    customer = application.execute("create_customer", {"id": "cus_late", "name": "Late customer"})
    application.execute("create_contract", {
        "id": "con_late", "name": "August service", "customer_id": customer["result"]["id"],
        "start_date": "2020-08-01", "end_date": "2020-08-31",
        "consideration": [{"id": "price_late", "label": "Fixed", "kind": "fixed", "amount": "3100.00"}],
        "obligations": [{"id": "pob_late", "name": "Service", "kind": "service", "ssp": "3100.00", "method": "exact_days", "start_date": "2020-08-01", "end_date": "2020-08-31"}],
    }, period="2020-08")
    cutoff = next(check for check in application.reports("main", "2020-08")["checks"] if check["id"] == "cutoff")
    assert cutoff["status"] == "review"
    assert cutoff["count"] == 1
    (application.workspace.path / "attachments" / "later.txt").write_text("Later support")
    application.execute("attach_evidence", {"name": "later.txt", "path": "attachments/later.txt", "entity_id": "con_late"}, period="2020-08")
    assert next(check for check in application.reports("main", "2020-08")["checks"] if check["id"] == "cutoff")["count"] == 1


def test_close_requires_explicit_acceptance_of_each_open_review_check(tmp_path):
    application = app(tmp_path)
    seed(application)
    application.execute("set_policy", {"effective_period": "2026-09", "accounts": {"revenue": "4100"}, "rationale": "Reviewed mapping"}, period="2026-09")
    dispositions = accept_review_items(application, "2026-09")
    assert "evidence" in dispositions
    with pytest.raises(ValueError, match="Review every open close check"):
        application.execute("close_period", {"period": "2026-09"}, period="2026-09")
    with pytest.raises(ValueError, match="acceptance reason"):
        application.execute("close_period", {"period": "2026-09", "review_dispositions": {**dispositions, "evidence": {"disposition": "accepted", "reason": ""}}}, period="2026-09")
    close = application.execute("close_period", {"period": "2026-09", "review_dispositions": dispositions}, period="2026-09")
    recorded = next(row for row in close["state"]["change_sets"] if row["command"] == "close_period")
    assert recorded["payload"]["review_dispositions"] == dispositions


def test_external_controls_compare_to_independent_totals_and_freeze_at_close(tmp_path):
    application = app(tmp_path)
    seed(application)
    application.execute("record_billing", {"contract_id": "con_1", "effective_date": "2026-09-01", "amount": "12000.00"}, period="2026-09")
    summary = application.state(period="2026-09")["report"]["summary"]
    payload = {"period": "2026-09", "source_name": "ERP September close", "rationale": "Matched the source invoice register and GL trial balance.",
               "billings": summary["billings"], "contract_asset": summary["contract_asset"], "deferred_revenue": summary["deferred_revenue"]}
    application.execute("record_control_totals", payload, period="2026-09")
    review = application.reports(period="2026-09")
    assert next(check for check in review["checks"] if check["id"] == "external_controls")["status"] == "pass"
    changed = {**payload, "deferred_revenue": "1.00", "rationale": "Revised GL extract shows a difference."}
    application.execute("record_control_totals", changed, period="2026-09")
    review = application.reports(period="2026-09")
    assert review["external_control_comparison"]["deferred_revenue"]["difference"] != "0.00"
    assert next(check for check in review["checks"] if check["id"] == "external_controls")["status"] == "review"
    application.execute("close_period", {"period": "2026-09", "review_dispositions": accept_review_items(application, "2026-09")}, period="2026-09")
    with pytest.raises(ValueError, match="Reopen the period"):
        application.execute("record_control_totals", payload, period="2026-09")


def test_source_population_catches_missing_records_even_when_money_totals_match(tmp_path):
    application = app(tmp_path)
    seed(application)
    application.execute("edit_details", {"entity_id": "con_1", "reference": "CON-1"}, period="2026-09")
    billing = application.execute("record_billing", {"contract_id": "con_1", "effective_date": "2026-09-01", "amount": "12000.00", "reference": "INV-1"}, period="2026-09")
    summary = application.state(period="2026-09")["report"]["summary"]
    application.execute("record_control_totals", {"period": "2026-09", "source_name": "Trial balance", "rationale": "Independent GL and billing totals", **{key: summary[key] for key in ("billings", "contract_asset", "deferred_revenue")}}, period="2026-09")
    manifest = {"period": "2026-09", "source_name": "Contract register and invoice ledger", "rationale": "Active agreements and September invoices, including credits.",
                "contract_references": ["CON-1", "CON-2"], "billing_references": [{"contract_reference": "CON-1", "invoice_reference": "INV-1"}, {"contract_reference": "CON-2", "invoice_reference": "INV-2"}]}
    application.execute("record_population_manifest", manifest, period="2026-09")
    review = application.reports(period="2026-09")
    assert next(check for check in review["checks"] if check["id"] == "external_controls")["status"] == "pass"
    assert next(check for check in review["checks"] if check["id"] == "source_population")["status"] == "review"
    assert review["population_comparison"]["missing_contracts"] == ["CON-2"]
    assert review["population_comparison"]["missing_billings"] == [("CON-2", "INV-2")]
    with pytest.raises(ValueError, match="duplicate references"):
        application.execute("record_population_manifest", {**manifest, "contract_references": ["CON-1", "CON-1"]}, period="2026-09")
    with pytest.raises(ValueError, match="duplicate contract and invoice"):
        application.execute("record_population_manifest", {**manifest, "billing_references": [manifest["billing_references"][0]] * 2}, period="2026-09")
    scenario = application.execute("create_scenario", {"name": "Alternative population"})["result"]["id"]
    with pytest.raises(ValueError, match="belong to Main"):
        application.execute("record_population_manifest", manifest, scenario_id=scenario, period="2026-09")
    application.execute("record_population_manifest", {**manifest, "contract_references": ["CON-1"], "billing_references": [manifest["billing_references"][0]]}, period="2026-09")
    review = application.reports(period="2026-09")
    assert next(check for check in review["checks"] if check["id"] == "source_population")["status"] == "pass"
    book = load_workbook(io.BytesIO(export_bytes(application.state(period="2026-09"), review)))
    assert book["Source contracts"]["A2"].value == "CON-1"
    assert book["Source invoices"]["B2"].value == "INV-1"
    book.close()
    application.execute("close_period", {"period": "2026-09", "review_dispositions": accept_review_items(application, "2026-09")}, period="2026-09")
    with pytest.raises(ValueError, match="Reopen the period"):
        application.execute("record_population_manifest", manifest, period="2026-09")
    application.execute("edit_details", {"entity_id": "con_1", "reference": "CON-REVISED"}, period="2026-10")
    closed = application.reports(period="2026-09")
    assert next(check for check in closed["checks"] if check["id"] == "source_population")["status"] == "pass"
    assert closed["population_comparison"]["missing_contracts"] == []
    assert billing["result"]["change_set_id"]


def test_source_population_uses_imported_billing_source_id_when_invoice_reference_is_absent(tmp_path):
    application = app(tmp_path)
    seed(application)
    application.execute("edit_details", {"entity_id": "con_1", "reference": "CON-1"}, period="2026-09")
    book = load_workbook(io.BytesIO(template_bytes()))
    sheet = book["Billing"]
    headers = [cell.value for cell in sheet[1]]
    values = {"contract_id": "con_1", "effective_date": "2026-09-02", "amount": "125.00", "source_id": "billing:125"}
    sheet.append([values.get(header) for header in headers])
    data = io.BytesIO()
    book.save(data)
    book.close()
    preview = preview_import_bytes(application, data.getvalue(), period="2026-09")
    import_bytes(application, data.getvalue(), period="2026-09", expected_frontier=preview["frontier"], expected_hash=preview["result"]["file_hash"])
    application.execute("record_population_manifest", {"period": "2026-09", "source_name": "ERP register", "rationale": "All active contracts and September billing transactions.",
                "contract_references": ["CON-1"], "billing_references": [{"contract_reference": "CON-1", "invoice_reference": "billing:125"}]}, period="2026-09")
    comparison = application.reports(period="2026-09")["population_comparison"]
    assert comparison["missing_billings"] == []
    assert comparison["unidentified_billings"] == []


def test_source_population_flags_duplicate_workspace_invoice_identities(tmp_path):
    application = app(tmp_path)
    seed(application)
    application.execute("edit_details", {"entity_id": "con_1", "reference": "CON-1"}, period="2026-09")
    for day in ("01", "02"):
        application.execute("record_billing", {"contract_id": "con_1", "effective_date": f"2026-09-{day}", "amount": "50.00", "reference": "INV-1"}, period="2026-09")
    application.execute("record_population_manifest", {"period": "2026-09", "source_name": "Invoice ledger", "rationale": "All September transactions.",
                "contract_references": ["CON-1"], "billing_references": [{"contract_reference": "CON-1", "invoice_reference": "INV-1"}]}, period="2026-09")
    review = application.reports(period="2026-09")
    assert review["population_comparison"]["duplicate_billings"] == [("CON-1", "INV-1")]
    assert next(check for check in review["checks"] if check["id"] == "source_population")["status"] == "review"


def test_period_associated_task_is_a_close_item_without_due_date(tmp_path):
    application = app(tmp_path)
    seed(application)
    application.execute("add_note", {"kind": "task", "body": "Confirm source invoice population", "period": "2026-10"})
    assert next(check for check in application.reports(period="2026-09")["checks"] if check["id"] == "tasks")["status"] == "pass"
    october = application.reports(period="2026-10")
    assert next(check for check in october["checks"] if check["id"] == "tasks")["count"] == 1


def test_close_cutoff_excludes_normal_post_period_entry_but_flags_later_entry(tmp_path, monkeypatch):
    application = app(tmp_path)
    seed(application)
    monkeypatch.setattr("openrevrec.application.now", lambda: "2026-10-05T12:00:00+00:00")
    application.execute("record_billing", {"contract_id": "con_1", "effective_date": "2026-09-30", "amount": "100.00"}, period="2026-09")
    summary = application.state(period="2026-09")["report"]["summary"]
    controls = {"period": "2026-09", "source_name": "ERP", "rationale": "Close register", "billings": summary["billings"], "contract_asset": summary["contract_asset"], "deferred_revenue": summary["deferred_revenue"]}
    application.execute("record_control_totals", {**controls, "close_cutoff_date": "2026-10-10"}, period="2026-09")
    assert next(check for check in application.reports(period="2026-09")["checks"] if check["id"] == "cutoff")["status"] == "pass"
    application.execute("record_control_totals", {**controls, "close_cutoff_date": "2026-10-01"}, period="2026-09")
    assert next(check for check in application.reports(period="2026-09")["checks"] if check["id"] == "cutoff")["count"] == 1


def test_scenario_rebase_detects_same_entity_conflict(tmp_path):
    application = app(tmp_path)
    seed(application)
    scenario = application.execute("create_scenario", {"name": "Alternative"})["result"]["id"]
    application.execute("record_billing", {"contract_id": "con_1", "effective_date": "2026-10-01", "amount": "1000.00"}, scenario_id=scenario)
    application.execute("record_billing", {"contract_id": "con_1", "effective_date": "2026-10-02", "amount": "1000.00"})
    with pytest.raises(ValueError, match="same accounting records"):
        application.execute("rebase_scenario", {"scenario_id": scenario})
    with pytest.raises(ValueError, match="Main has advanced"):
        application.execute("apply_scenario", {"scenario_id": scenario})


def test_scenario_rebase_accepts_distinct_identified_invoices_on_one_contract(tmp_path):
    application = app(tmp_path)
    seed(application)
    scenario = application.execute("create_scenario", {"name": "Second invoice"})["result"]["id"]
    application.execute("record_billing", {"contract_id": "con_1", "effective_date": "2026-10-02",
                                           "amount": "100.00", "reference": "INV-102"}, scenario_id=scenario, period="2026-10")
    application.execute("record_billing", {"contract_id": "con_1", "effective_date": "2026-10-01",
                                           "amount": "200.00", "reference": "INV-101"}, period="2026-10")
    assert application.compare(scenario, "2026-10")["conflicts"] == []
    application.execute("rebase_scenario", {"scenario_id": scenario}, scenario_id=scenario, period="2026-10")
    rebased = application.state(scenario, "2026-10")
    assert rebased["report"]["summary"]["billings"] == "300.00"
    application.execute("apply_scenario", {"scenario_id": scenario}, period="2026-10")
    accepted = application.state(period="2026-10")
    assert accepted["report"]["summary"]["billings"] == "300.00"
    assert sum(row["debit_minor"] for row in accepted["report"]["journals"]) == sum(
        row["credit_minor"] for row in accepted["report"]["journals"])


@pytest.mark.parametrize("main_amount,main_reference", [("200.00", "INV-102"), ("-20.00", "")])
def test_scenario_rebase_retains_invoice_collision_and_unidentified_credit_review(tmp_path, main_amount, main_reference):
    application = app(tmp_path)
    seed(application)
    original = application.execute("record_billing", {"contract_id": "con_1", "effective_date": "2026-09-01",
                                                     "amount": "500.00", "reference": "INV-100"}, period="2026-09")["result"]["change_set_id"]
    scenario = application.execute("create_scenario", {"name": "Second invoice"})["result"]["id"]
    application.execute("record_billing", {"contract_id": "con_1", "effective_date": "2026-10-02",
                                           "amount": "100.00", "reference": "INV-102"}, scenario_id=scenario, period="2026-10")
    main_invoice = {"contract_id": "con_1", "effective_date": "2026-10-01", "amount": main_amount,
                    "reference": main_reference}
    if Decimal(main_amount) < 0:
        main_invoice.update(applies_to_change_set_id=original, rationale="Correct original invoice")
    application.execute("record_billing", main_invoice, period="2026-10")
    assert len(application.compare(scenario, "2026-10")["conflicts"]) == 1
    with pytest.raises(ValueError, match="same accounting records"):
        application.execute("rebase_scenario", {"scenario_id": scenario}, scenario_id=scenario, period="2026-10")


def test_scenario_rebase_combines_new_invoice_with_identified_credit_to_older_invoice(tmp_path):
    application = app(tmp_path)
    seed(application)
    original = application.execute("record_billing", {"contract_id": "con_1", "effective_date": "2026-09-01",
                                                     "amount": "500.00", "reference": "INV-100"}, period="2026-09")["result"]["change_set_id"]
    scenario = application.execute("create_scenario", {"name": "Second invoice"})["result"]["id"]
    application.execute("record_billing", {"contract_id": "con_1", "effective_date": "2026-10-02",
                                           "amount": "100.00", "reference": "INV-102"}, scenario_id=scenario, period="2026-10")
    application.execute("record_billing", {"contract_id": "con_1", "effective_date": "2026-10-01",
                                           "amount": "-20.00", "reference": "CM-100", "applies_to_change_set_id": original,
                                           "rationale": "Credit against the September invoice"}, period="2026-10")
    assert application.compare(scenario, "2026-10")["conflicts"] == []
    application.execute("rebase_scenario", {"scenario_id": scenario}, scenario_id=scenario, period="2026-10")
    assert application.state(scenario, "2026-10")["report"]["summary"]["billings"] == "80.00"
    application.execute("apply_scenario", {"scenario_id": scenario}, period="2026-10")
    accepted = application.state(period="2026-10")
    assert accepted["report"]["summary"]["billings"] == "80.00"
    assert any(activity.get("applies_to_change_set_id") == original for activity in accepted["contracts"][0]["activities"])
    assert sum(row["debit_minor"] - row["credit_minor"] for row in accepted["report"]["journals"]) == 0


def test_scenario_rebase_combines_identified_credit_with_new_main_invoice(tmp_path):
    application = app(tmp_path)
    seed(application)
    original = application.execute("record_billing", {"contract_id": "con_1", "effective_date": "2026-09-01",
                                                     "amount": "500.00", "reference": "INV-100"}, period="2026-09")["result"]["change_set_id"]
    scenario = application.execute("create_scenario", {"name": "Invoice credit"})["result"]["id"]
    application.execute("record_billing", {"contract_id": "con_1", "effective_date": "2026-10-02",
                                           "amount": "-20.00", "reference": "CM-100", "applies_to_change_set_id": original,
                                           "rationale": "Credit against the September invoice"}, scenario_id=scenario, period="2026-10")
    application.execute("record_billing", {"contract_id": "con_1", "effective_date": "2026-10-01",
                                           "amount": "100.00", "reference": "INV-102"}, period="2026-10")
    assert application.compare(scenario, "2026-10")["conflicts"] == []
    application.execute("rebase_scenario", {"scenario_id": scenario}, scenario_id=scenario, period="2026-10")
    application.execute("apply_scenario", {"scenario_id": scenario}, period="2026-10")
    accepted = application.state(period="2026-10")
    assert accepted["report"]["summary"]["billings"] == "80.00"
    assert any(activity.get("applies_to_change_set_id") == original for activity in accepted["contracts"][0]["activities"])
    assert sum(row["debit_minor"] - row["credit_minor"] for row in accepted["report"]["journals"]) == 0


def test_scenario_rebase_combines_distinct_identified_credits_to_different_workspace_invoices(tmp_path):
    application = app(tmp_path)
    seed(application)
    first = application.execute("record_billing", {"contract_id": "con_1", "effective_date": "2026-09-01",
                                                   "amount": "100.00", "reference": "INV-101"}, period="2026-09")["result"]["change_set_id"]
    second = application.execute("record_billing", {"contract_id": "con_1", "effective_date": "2026-09-02",
                                                    "amount": "200.00", "reference": "INV-102"}, period="2026-09")["result"]["change_set_id"]
    scenario = application.execute("create_scenario", {"name": "Credit first invoice"})["result"]["id"]
    application.execute("record_billing", {"contract_id": "con_1", "effective_date": "2026-10-02", "amount": "-10.00",
                                           "reference": "CM-101", "applies_to_change_set_id": first,
                                           "rationale": "Credit the first invoice."}, scenario_id=scenario, period="2026-10")
    application.execute("record_billing", {"contract_id": "con_1", "effective_date": "2026-10-03", "amount": "-20.00",
                                           "reference": "CM-102", "applies_to_change_set_id": second,
                                           "rationale": "Credit the second invoice."}, period="2026-10")
    assert application.compare(scenario, "2026-10")["conflicts"] == []
    application.execute("rebase_scenario", {"scenario_id": scenario}, scenario_id=scenario, period="2026-10")
    application.execute("apply_scenario", {"scenario_id": scenario}, period="2026-10")
    accepted = application.state(period="2026-10")
    assert accepted["report"]["summary"]["billings"] == "-30.00"
    assert {activity.get("applies_to_change_set_id") for activity in accepted["contracts"][0]["activities"]
            if activity["type"] == "billing" and Decimal(activity["amount"]) < 0} == {first, second}
    assert sum(row["debit_minor"] - row["credit_minor"] for row in accepted["report"]["journals"]) == 0


def test_scenario_rebase_keeps_same_original_or_duplicate_credit_in_conflict(tmp_path):
    application = app(tmp_path)
    seed(application)
    original = application.execute("record_billing", {"contract_id": "con_1", "effective_date": "2026-09-01",
                                                      "amount": "100.00", "reference": "INV-101"}, period="2026-09")["result"]["change_set_id"]
    scenario = application.execute("create_scenario", {"name": "Credit original"})["result"]["id"]
    application.execute("record_billing", {"contract_id": "con_1", "effective_date": "2026-10-02", "amount": "-10.00",
                                           "reference": "CM-101", "applies_to_change_set_id": original,
                                           "rationale": "Scenario credit."}, scenario_id=scenario, period="2026-10")
    application.execute("record_billing", {"contract_id": "con_1", "effective_date": "2026-10-03", "amount": "-20.00",
                                           "reference": "CM-102", "applies_to_change_set_id": original,
                                           "rationale": "Main credit to the same invoice."}, period="2026-10")
    assert len(application.compare(scenario, "2026-10")["conflicts"]) == 1
    with pytest.raises(ValueError, match="same accounting records"):
        application.execute("rebase_scenario", {"scenario_id": scenario}, scenario_id=scenario, period="2026-10")


def test_scenario_rebase_keeps_duplicate_credit_reference_in_conflict(tmp_path):
    application = app(tmp_path)
    seed(application)
    first = application.execute("record_billing", {"contract_id": "con_1", "effective_date": "2026-09-01",
                                                   "amount": "100.00", "reference": "INV-101"}, period="2026-09")["result"]["change_set_id"]
    second = application.execute("record_billing", {"contract_id": "con_1", "effective_date": "2026-09-02",
                                                    "amount": "200.00", "reference": "INV-102"}, period="2026-09")["result"]["change_set_id"]
    scenario = application.execute("create_scenario", {"name": "Credit first invoice"})["result"]["id"]
    application.execute("record_billing", {"contract_id": "con_1", "effective_date": "2026-10-02", "amount": "-10.00",
                                           "reference": "CM-DUPLICATE", "applies_to_change_set_id": first,
                                           "rationale": "Scenario credit."}, scenario_id=scenario, period="2026-10")
    application.execute("record_billing", {"contract_id": "con_1", "effective_date": "2026-10-03", "amount": "-20.00",
                                           "reference": "CM-DUPLICATE", "applies_to_change_set_id": second,
                                           "rationale": "Main credit with duplicate credit number."}, period="2026-10")
    assert len(application.compare(scenario, "2026-10")["conflicts"]) == 1


def test_scenario_rebase_keeps_credits_to_duplicate_source_invoices_in_conflict(tmp_path):
    application = app(tmp_path)
    seed(application)
    first = application.execute("record_billing", {"contract_id": "con_1", "effective_date": "2026-09-01",
                                                   "amount": "100.00", "reference": "INV-DUPLICATE"}, period="2026-09")["result"]["change_set_id"]
    second = application.execute("record_billing", {"contract_id": "con_1", "effective_date": "2026-09-02",
                                                    "amount": "100.00", "reference": "INV-DUPLICATE"}, period="2026-09")["result"]["change_set_id"]
    scenario = application.execute("create_scenario", {"name": "Credit first source record"})["result"]["id"]
    application.execute("record_billing", {"contract_id": "con_1", "effective_date": "2026-10-02", "amount": "-10.00",
                                           "reference": "CM-101", "applies_to_change_set_id": first,
                                           "rationale": "Credit first source record."}, scenario_id=scenario, period="2026-10")
    application.execute("record_billing", {"contract_id": "con_1", "effective_date": "2026-10-03", "amount": "-20.00",
                                           "reference": "CM-102", "applies_to_change_set_id": second,
                                           "rationale": "Credit second source record."}, period="2026-10")
    assert len(application.compare(scenario, "2026-10")["conflicts"]) == 1


def test_scenario_rebase_keeps_workspace_credit_and_new_invoice_with_same_source_reference_in_conflict(tmp_path):
    application = app(tmp_path)
    seed(application)
    original = application.execute("record_billing", {"contract_id": "con_1", "effective_date": "2026-09-01",
                                                      "amount": "100.00", "reference": "INV-101"}, period="2026-09")["result"]["change_set_id"]
    scenario = application.execute("create_scenario", {"name": "Credit old invoice"})["result"]["id"]
    application.execute("record_billing", {"contract_id": "con_1", "effective_date": "2026-10-02", "amount": "-10.00",
                                           "reference": "CM-101", "applies_to_change_set_id": original,
                                           "rationale": "Credit old invoice."}, scenario_id=scenario, period="2026-10")
    application.execute("record_billing", {"contract_id": "con_1", "effective_date": "2026-10-03", "amount": "100.00",
                                           "reference": "INV-101"}, period="2026-10")
    assert len(application.compare(scenario, "2026-10")["conflicts"]) == 1


@pytest.mark.parametrize("second_original,conflicts", [("LEGACY-102", False), ("LEGACY-101", True)])
def test_scenario_rebase_external_credits_need_distinct_originals(tmp_path, second_original, conflicts):
    application = app(tmp_path)
    seed(application)
    scenario = application.execute("create_scenario", {"name": "Legacy credit"})["result"]["id"]
    application.execute("record_billing", {"contract_id": "con_1", "effective_date": "2026-10-02", "amount": "-10.00",
                                           "reference": "CM-101", "applies_to_reference": "LEGACY-101",
                                           "rationale": "Credit first legacy invoice."}, scenario_id=scenario, period="2026-10")
    application.execute("record_billing", {"contract_id": "con_1", "effective_date": "2026-10-03", "amount": "-20.00",
                                           "reference": "CM-102", "applies_to_reference": second_original,
                                           "rationale": "Credit another legacy invoice."}, period="2026-10")
    assert bool(application.compare(scenario, "2026-10")["conflicts"]) is conflicts
    if not conflicts:
        application.execute("rebase_scenario", {"scenario_id": scenario}, scenario_id=scenario, period="2026-10")
        application.execute("apply_scenario", {"scenario_id": scenario}, period="2026-10")
        accepted = application.state(period="2026-10")
        assert accepted["report"]["summary"]["billings"] == "-30.00"
        assert sum(row["debit_minor"] - row["credit_minor"] for row in accepted["report"]["journals"]) == 0


@pytest.mark.parametrize("new_invoice_reference,conflicts", [("INV-NEW", False), ("INV-OLD", True)])
def test_scenario_rebase_external_original_credit_requires_distinct_new_invoice(tmp_path, new_invoice_reference, conflicts):
    application = app(tmp_path)
    seed(application)
    scenario = application.execute("create_scenario", {"name": "Credit legacy invoice"})["result"]["id"]
    application.execute("record_billing", {"contract_id": "con_1", "effective_date": "2026-10-02", "amount": "-10.00",
                                           "reference": "CM-OLD", "applies_to_reference": "INV-OLD",
                                           "rationale": "Credit an invoice from the legacy billing system."}, scenario_id=scenario, period="2026-10")
    application.execute("record_billing", {"contract_id": "con_1", "effective_date": "2026-10-03", "amount": "100.00",
                                           "reference": new_invoice_reference}, period="2026-10")
    assert bool(application.compare(scenario, "2026-10")["conflicts"]) is conflicts
    if conflicts:
        with pytest.raises(ValueError, match="same accounting records"):
            application.execute("rebase_scenario", {"scenario_id": scenario}, scenario_id=scenario, period="2026-10")
    else:
        application.execute("rebase_scenario", {"scenario_id": scenario}, scenario_id=scenario, period="2026-10")
        application.execute("apply_scenario", {"scenario_id": scenario}, period="2026-10")
        accepted = application.state(period="2026-10")
        assert accepted["report"]["summary"]["billings"] == "90.00"
        assert any(activity.get("applies_to_reference") == "INV-OLD" for activity in accepted["contracts"][0]["activities"])
        assert sum(row["debit_minor"] - row["credit_minor"] for row in accepted["report"]["journals"]) == 0


def test_scenario_rebase_combines_new_invoice_with_credit_to_distinct_external_original(tmp_path):
    application = app(tmp_path)
    seed(application)
    scenario = application.execute("create_scenario", {"name": "Second invoice"})["result"]["id"]
    application.execute("record_billing", {"contract_id": "con_1", "effective_date": "2026-10-02",
                                           "amount": "100.00", "reference": "INV-102"}, scenario_id=scenario, period="2026-10")
    application.execute("record_billing", {"contract_id": "con_1", "effective_date": "2026-10-01",
                                           "amount": "-20.00", "reference": "CM-100", "applies_to_reference": "LEGACY-100",
                                           "rationale": "Credit against a pre-workspace invoice"}, period="2026-10")
    assert application.compare(scenario, "2026-10")["conflicts"] == []
    application.execute("rebase_scenario", {"scenario_id": scenario}, scenario_id=scenario, period="2026-10")
    application.execute("apply_scenario", {"scenario_id": scenario}, period="2026-10")
    accepted = application.state(period="2026-10")
    assert accepted["report"]["summary"]["billings"] == "80.00"
    assert any(activity.get("applies_to_reference") == "LEGACY-100" for activity in accepted["contracts"][0]["activities"])
    assert sum(row["debit_minor"] - row["credit_minor"] for row in accepted["report"]["journals"]) == 0


def test_scenario_rebase_keeps_incomparable_invoice_identifiers_for_review(tmp_path):
    application = app(tmp_path)
    seed(application)
    scenario = application.execute("create_scenario", {"name": "Second invoice"})["result"]["id"]
    application.execute("record_billing", {"contract_id": "con_1", "effective_date": "2026-10-02",
                                           "amount": "100.00", "reference": "INV-102"}, scenario_id=scenario, period="2026-10")
    application.execute("record_billing", {"contract_id": "con_1", "effective_date": "2026-10-01",
                                           "amount": "200.00", "import_source_identity": ["Billing", "source:101"]}, period="2026-10")
    assert len(application.compare(scenario, "2026-10")["conflicts"]) == 1


def test_scenario_invoice_rebase_still_reviews_global_account_change(tmp_path):
    application = app(tmp_path)
    seed(application)
    scenario = application.execute("create_scenario", {"name": "Invoice"})["result"]["id"]
    application.execute("record_billing", {"contract_id": "con_1", "effective_date": "2026-10-02",
                                           "amount": "100.00", "reference": "INV-102"}, scenario_id=scenario, period="2026-10")
    application.execute("set_policy", {"effective_period": "2026-10", "accounts": {"billing_clearing": "1110"},
                                       "rationale": "New clearing account"}, period="2026-10")
    assert [row["command"] for row in application.compare(scenario, "2026-10")["conflicts"]] == ["set_policy"]
    with pytest.raises(ValueError, match="same accounting records"):
        application.execute("rebase_scenario", {"scenario_id": scenario}, scenario_id=scenario, period="2026-10")


def test_scenario_rebase_ignores_descriptive_main_edit_and_lists_proposals(tmp_path):
    application = app(tmp_path)
    seed(application)
    scenario = application.execute("create_scenario", {"name": "Alternative"})["result"]["id"]
    proposed = application.execute("record_billing", {"contract_id": "con_1", "effective_date": "2026-10-01", "amount": "100"}, scenario_id=scenario, period="2026-10")["result"]["change_set_id"]
    application.execute("edit_details", {"entity_id": "con_1", "description": "Updated contact note"})
    comparison = application.compare(scenario, "2026-10")
    assert comparison["proposals"][0]["id"] == proposed
    assert comparison["conflicts"] == []
    application.execute("rebase_scenario", {"scenario_id": scenario}, scenario_id=scenario, period="2026-10")
    assert application.state(scenario, "2026-10")["report"]["summary"]["billings"] == "100.00"


def test_archive_and_restore_keep_scenario_history(tmp_path):
    application = app(tmp_path)
    seed(application)
    scenario = application.execute("create_scenario", {"name": "Saved alternative"})["result"]["id"]
    application.execute("record_billing", {"contract_id": "con_1", "effective_date": "2026-09-01", "amount": "2000.00"}, scenario_id=scenario)
    application.execute("archive_scenario", {"scenario_id": scenario})
    assert next(s for s in application.state()["scenarios"] if s["id"] == scenario)["status"] == "archived"
    application.execute("restore_scenario", {"scenario_id": scenario})
    state = application.state(scenario, "2026-09")
    assert next(s for s in state["scenarios"] if s["id"] == scenario)["status"] == "active"
    assert state["report"]["summary"]["billings"] == "2000.00"


def test_close_checkpoint_blocks_change_and_reopen_is_explicit(tmp_path):
    application = app(tmp_path)
    seed(application)
    application.execute("record_billing", {"contract_id": "con_1", "effective_date": "2026-09-01", "amount": "12000.00"}, period="2026-09")
    open_positions = application.state(period="2026-09")["report"]["account_positions"]
    closed = application.execute("close_period", {"period": "2026-09", "rationale": "Reconciled to billing support.", "review_dispositions": accept_review_items(application, "2026-09")}, period="2026-09")
    assert closed["state"]["report"]["closed"] is True
    assert closed["state"]["report"]["account_positions"] == open_positions
    assert closed["result"]["backup_path"].endswith(".orr")
    with pytest.raises(ValueError, match="affects closed period"):
        application.execute("record_billing", {"contract_id": "con_1", "effective_date": "2026-09-30", "amount": "1.00"}, period="2026-09")
    application.execute("reopen_period", {"period": "2026-09", "rationale": "Late invoice requires correction."}, period="2026-09")
    updated = application.execute("record_billing", {"contract_id": "con_1", "effective_date": "2026-09-30", "amount": "1.00"}, period="2026-09")
    assert updated["state"]["report"]["summary"]["billings"] == "12001.00"


def test_sqlite_rejects_direct_history_changes(tmp_path):
    application = app(tmp_path)
    seed(application)
    with application.workspace.connect() as db:
        with pytest.raises(sqlite3.IntegrityError, match="immutable"):
            db.execute("DELETE FROM change_sets")
    result = application.query("SELECT command, entity_id FROM change_sets ORDER BY version")
    assert result["columns"] == ["command", "entity_id"]
    assert len(result["rows"]) == 2
    with pytest.raises(ValueError):
        application.query("DELETE FROM change_sets")


def test_metadata_edits_and_company_setup_do_not_change_accounting(tmp_path):
    application = app(tmp_path)
    seed(application)
    before = application.state(period="2026-09")["report"]
    application.execute("edit_details", {
        "entity_id": "cus_1", "name": "Renamed customer", "email": "accounting@example.com",
        "reference": "CRM-42", "description": "Internal operating note",
    })
    application.execute("set_policy", {
        "name": "Renamed company", "accounts": {"revenue": "4100"}, "rationale": "Initial company setup",
    })
    state = application.state(period="2026-09")
    assert state["workspace"]["name"] == "Renamed company"
    assert state["customers"][0]["email"] == "accounting@example.com"
    assert state["customers"][0]["reference"] == "CRM-42"
    assert state["report"]["summary"] == before["summary"]


def test_accounting_activity_view_matches_sql_inspector_starter_query(tmp_path):
    application = app(tmp_path)
    seed(application)
    application.execute("record_billing", {"contract_id": "con_1", "effective_date": "2026-09-01", "amount": "12.00"})
    result = application.query("SELECT scenario_id, activity_type, contract_id, effective_date, recorded_at FROM accounting_activity ORDER BY recorded_at DESC LIMIT 25")
    assert result["columns"] == ["scenario_id", "activity_type", "contract_id", "effective_date", "recorded_at"]
    assert result["rows"][0][1:3] == ["record_billing", "con_1"]


def test_template_import_is_atomic_and_export_is_valid(tmp_path):
    application = app(tmp_path)
    book = load_workbook(io.BytesIO(template_bytes()))
    book["Customers"].append(["cus_import", "Imported customer", "", "C-1"])
    book["Contracts"].append(["con_import", "cus_import", "Imported contract", "2026-09-01", "2026-09-30", "Reviewed contract"])
    book["Consideration"].append(["con_import", "price_import", "Fixed fee", "fixed", "1000.00", "", "", "", "", ""])
    book["Obligations"].append(["con_import", "pob_import", "September service", "service", "1000.00", "exact_days", "2026-09-01", "2026-09-30", "", "", "", ""])
    data = io.BytesIO(); book.save(data); book.close()
    result = import_bytes(application, data.getvalue(), "import.xlsx", period="2026-09")
    assert result["result"]["imported"] == 2
    assert len(result["state"]["contracts"]) == 1
    review = application.reports("main", "2026-09")
    sheets = load_workbook(io.BytesIO(export_bytes(result["state"], review))).sheetnames
    assert sheets[:5] == ["Workspace", "Summary", "Close readiness", "External controls", "Contract rollforward"]
    assert "Recognition coverage" in sheets

    bad = load_workbook(io.BytesIO(template_bytes()))
    bad["Customers"].append(["cus_second", "Second", "", ""])
    bad["Contracts"].append(["con_bad", "missing", "Bad", "2026-09-01", "2026-09-30", ""])
    bad["Consideration"].append(["con_bad", "p", "Fixed", "fixed", "10", "", "", "", "", ""])
    bad["Obligations"].append(["con_bad", "o", "Service", "service", "10", "exact_days", "2026-09-01", "2026-09-30", "", "", "", ""])
    broken = io.BytesIO(); bad.save(broken); bad.close()
    with pytest.raises(ValueError, match="No accounting rows were imported"):
        import_bytes(application, broken.getvalue(), "bad.xlsx")
    assert not any(c["id"] == "cus_second" for c in application.state()["customers"])


def test_import_preview_is_read_only_and_rejects_repeated_source_rows(tmp_path):
    application = app(tmp_path)
    seed(application)
    book = load_workbook(io.BytesIO(template_bytes()))
    book["Billing"].append(["con_1", "2026-09-01", "100.00", "INV-1", "", "", "", "erp:invoice-line-1"])
    data = io.BytesIO(); book.save(data); book.close()
    before = application.state(period="2026-09")

    preview = preview_import_bytes(application, data.getvalue(), period="2026-09")
    assert preview["result"]["imported"] == 1
    assert preview["state"]["report"]["summary"]["billings"] == "100.00"
    assert application.state(period="2026-09") == before

    import_bytes(application, data.getvalue(), period="2026-09", expected_frontier=preview["frontier"], expected_hash=preview["result"]["file_hash"])
    assert application.state(period="2026-09")["report"]["summary"]["billings"] == "100.00"
    with pytest.raises(ValueError, match="already imported"):
        import_bytes(application, data.getvalue(), period="2026-09")
    assert application.state(period="2026-09")["report"]["summary"]["billings"] == "100.00"

    changed = load_workbook(io.BytesIO(data.getvalue()))
    changed["Billing"]["C2"] = "125.00"
    revised = io.BytesIO(); changed.save(revised); changed.close()
    with pytest.raises(ValueError, match="already imported"):
        preview_import_bytes(application, revised.getvalue(), period="2026-09")


def test_imported_billing_credit_retains_external_original_reference(tmp_path):
    application = app(tmp_path)
    seed(application)
    book = load_workbook(io.BytesIO(template_bytes()))
    book["Billing"].append(["con_1", "2026-09-15", "-25.00", "CM-25", "", "LEGACY-INV-20", "Credit memo for legacy invoice", "erp:credit-25"])
    output = io.BytesIO(); book.save(output); book.close()
    reviewed = preview_import_bytes(application, output.getvalue(), period="2026-09")
    assert reviewed["state"]["report"]["summary"]["billings"] == "-25.00"
    committed = import_bytes(application, output.getvalue(), period="2026-09", expected_frontier=reviewed["frontier"], expected_hash=reviewed["result"]["file_hash"])
    credit = next(row for row in committed["state"]["change_sets"] if row["command"] == "record_billing")
    assert credit["payload"]["applies_to_reference"] == "LEGACY-INV-20"


def test_import_commit_requires_same_file_and_frontier_as_preview(tmp_path):
    application = app(tmp_path)
    seed(application)
    book = load_workbook(io.BytesIO(template_bytes()))
    book["Billing"].append(["con_1", "2026-09-01", "100.00", "INV-1", "", "", "", ""])
    data = io.BytesIO(); book.save(data); book.close()
    preview = preview_import_bytes(application, data.getvalue(), period="2026-09")
    with pytest.raises(ValueError, match="workbook changed"):
        import_bytes(application, data.getvalue(), period="2026-09", expected_hash="wrong")
    application.execute("add_note", {"entity_id": "con_1", "body": "Reviewer note"}, period="2026-09")
    with pytest.raises(ValueError, match="workspace changed"):
        import_bytes(application, data.getvalue(), period="2026-09", expected_frontier=preview["frontier"], expected_hash=preview["result"]["file_hash"])
    assert application.state(period="2026-09")["report"]["summary"]["billings"] == "0.00"


def test_import_preview_shows_future_billing_period_even_when_selected_month_is_unchanged(tmp_path):
    application = app(tmp_path)
    seed(application)
    book = load_workbook(io.BytesIO(template_bytes()))
    book["Billing"].append(["con_1", "2026-12-01", "300.00", "INV-DEC", "", "", "", ""])
    data = io.BytesIO(); book.save(data); book.close()
    preview = preview_import_bytes(application, data.getvalue(), period="2026-09")
    assert preview["comparison"]["summary"]["billings"] == "0.00"
    assert "2026-12" in preview["comparison"]["affected_periods"]
    assert next(row for row in preview["period_impacts"] if row["period"] == "2026-12")["delta"]["billings"] == "300.00"


def test_import_source_identity_survives_scenario_application(tmp_path):
    application = app(tmp_path)
    seed(application)
    scenario = application.execute("create_scenario", {"name": "Import review"})["result"]["id"]
    book = load_workbook(io.BytesIO(template_bytes()))
    book["Billing"].append(["con_1", "2026-09-01", "100.00", "INV-SCENARIO", "", "", "", "erp:invoice-scenario"])
    data = io.BytesIO(); book.save(data); book.close()
    import_bytes(application, data.getvalue(), scenario_id=scenario, period="2026-09")
    application.execute("apply_scenario", {"scenario_id": scenario}, scenario_id=scenario, period="2026-09")
    assert application.state(period="2026-09")["report"]["summary"]["billings"] == "100.00"
    with pytest.raises(ValueError, match="already imported"):
        import_bytes(application, data.getvalue(), scenario_id="main", period="2026-09")


def test_usage_source_correction_replaces_measure_without_erasing_history(tmp_path):
    application = app(tmp_path)
    application.execute("create_customer", {"id": "cus_1", "name": "Customer"})
    application.execute("create_contract", {
        "id": "con_1", "name": "Finite units", "customer_id": "cus_1",
        "start_date": "2026-01-01", "end_date": "2026-12-31",
        "consideration": [{"id": "price", "kind": "fixed", "amount": "1000"}],
        "obligations": [{"id": "units", "name": "Units", "kind": "service", "ssp": "1000", "method": "usage", "total_units": "100", "start_date": "2026-01-01", "end_date": "2026-12-31"}],
    }, period="2026-01")
    expected_gap = application.reports(period="2026-01")
    assert expected_gap["recognition_coverage"][0]["unscheduled"] == "1000.00"
    assert next(check for check in expected_gap["checks"] if check["id"] == "coverage")["status"] == "pass"
    assert not expected_gap["warnings"]
    original = application.execute("record_usage", {"contract_id": "con_1", "obligation_id": "units", "effective_date": "2026-01-10", "quantity": "50"}, period="2026-01")["result"]["change_set_id"]
    application.execute("record_usage", {"contract_id": "con_1", "obligation_id": "units", "effective_date": "2026-02-10", "quantity": "10"}, period="2026-02")
    assert application.state(period="2026-02")["report"]["summary"]["recognized_to_date"] == "600.00"
    correction = application.execute("correct_activity", {
        "target_change_set_id": original, "rationale": "Source usage file overstated January by ten units",
        "replacement": {"contract_id": "con_1", "obligation_id": "units", "effective_date": "2026-01-10", "quantity": "40"},
    }, period="2026-02")
    assert correction["state"]["report"]["summary"]["recognized_to_date"] == "500.00"
    assert correction["state"]["report"]["summary"]["revenue"] == "100.00"
    assert correction["state"]["contracts"][0]["activities"][0]["corrects"] == original
    assert any(change["id"] == original for change in correction["state"]["change_sets"])
    with pytest.raises(ValueError, match="current"):
        application.execute("correct_activity", {"target_change_set_id": original, "rationale": "Duplicate correction", "replacement": {"contract_id": "con_1", "obligation_id": "units", "effective_date": "2026-01-10", "quantity": "40"}}, period="2026-02")


def test_metered_usage_correction_reprices_open_period_and_closed_period_is_protected(tmp_path):
    application = app(tmp_path)
    application.execute("create_customer", {"id": "cus_1", "name": "Customer"})
    application.execute("create_contract", {
        "id": "con_1", "name": "Metered service", "customer_id": "cus_1",
        "start_date": "2026-01-01", "end_date": "2026-12-31",
        "consideration": [{"id": "meter", "kind": "metered", "amount": "0", "unit_rate": "2.50",
                           "pricing_basis": "right_to_invoice", "rounding_period": "calendar_month",
                           "rationale": "Each measured unit has the same customer value."}],
        "obligations": [{"id": "units", "name": "Units", "kind": "service", "ssp": "1", "method": "metered",
                         "start_date": "2026-01-01", "end_date": "2026-12-31"}],
    }, period="2026-01")
    original = application.execute("record_usage", {
        "contract_id": "con_1", "obligation_id": "units", "effective_date": "2026-01-10", "quantity": "100",
    }, period="2026-01")["result"]["change_set_id"]
    application.execute("record_usage", {
        "contract_id": "con_1", "obligation_id": "units", "effective_date": "2026-02-10", "quantity": "200",
    }, period="2026-02")
    assert application.state(period="2026-02")["report"]["summary"]["recognized_to_date"] == "750.00"
    correction = application.execute("correct_activity", {
        "target_change_set_id": original, "rationale": "Meter source overstated January by ten units",
        "replacement": {"contract_id": "con_1", "obligation_id": "units", "effective_date": "2026-01-10", "quantity": "90"},
    }, period="2026-02")
    assert correction["state"]["report"]["summary"]["recognized_to_date"] == "725.00"
    assert correction["state"]["report"]["summary"]["revenue"] == "500.00"
    application.execute("close_period", {
        "period": "2026-01", "review_dispositions": accept_review_items(application, "2026-01"),
    }, period="2026-01")
    with pytest.raises(ValueError, match="closed period"):
        application.execute("record_usage", {
            "contract_id": "con_1", "obligation_id": "units", "effective_date": "2026-01-20", "quantity": "1",
        }, period="2026-02")


def test_metered_rate_changes_use_delivery_date_and_round_once_per_month(tmp_path):
    application = app(tmp_path)
    application.execute("create_customer", {"id": "cus_1", "name": "Customer"})
    application.execute("create_contract", {
        "id": "con_1", "name": "Metered service", "customer_id": "cus_1",
        "start_date": "2026-01-01", "end_date": "2026-12-31",
        "consideration": [{"id": "meter", "kind": "metered", "amount": "0", "unit_rate": "0.015",
                           "pricing_basis": "right_to_invoice", "rounding_period": "calendar_month",
                           "rationale": "Invoice value tracks delivered units."}],
        "obligations": [{"id": "units", "name": "Units", "kind": "service", "ssp": "0", "method": "metered",
                         "start_date": "2026-01-01", "end_date": "2026-12-31"}],
    }, period="2026-01")
    application.execute("record_usage", {"contract_id": "con_1", "obligation_id": "units",
                                          "effective_date": "2026-01-10", "quantity": "1001"}, period="2026-01")
    application.execute("record_usage", {"contract_id": "con_1", "obligation_id": "units",
                                          "effective_date": "2026-01-15", "quantity": "9"}, period="2026-01")
    january_before = application.state(period="2026-01")["report"]
    assert january_before["summary"]["revenue"] == "15.15"
    with pytest.raises(ValueError, match="conclusion"):
        application.execute("record_rate_change", {"contract_id": "con_1", "component_id": "meter",
            "effective_date": "2026-01-15", "unit_rate": "0.017"}, period="2026-01")
    first_change = application.execute("record_rate_change", {"contract_id": "con_1", "component_id": "meter",
        "effective_date": "2026-01-15", "unit_rate": "0.017",
        "rationale": "Approved schedule confirms the revised price reflects delivered units."}, period="2026-01")["result"]["change_set_id"]
    january = application.state(period="2026-01")["report"]
    assert january["summary"]["revenue"] == "15.17"
    contract = january["contracts"][0]
    assert [(item["effective_date"], item["unit_rate"]) for item in contract["metered_rate_history"]] == [
        ("2026-01-01", "0.015"), ("2026-01-15", "0.017")]
    assert [(item["quantity"], item["unit_rate"]) for item in contract["metered_usage_valuation"]] == [
        ("1001", "0.015"), ("9", "0.017")]
    assert contract["metered_monthly_values"] == [{"period": "2026-01", "quantity": "1010",
        "unrounded_value": "15.168", "revenue": "15.17"}]
    assert sum(item["debit_minor"] - item["credit_minor"] for item in january["journals"]) == 0
    with pytest.raises(ValueError, match="one metered rate change"):
        application.execute("record_rate_change", {"contract_id": "con_1", "component_id": "meter",
            "effective_date": "2026-01-15", "unit_rate": "0.018", "rationale": "Duplicate date"}, period="2026-01")
    with pytest.raises(ValueError, match="different from the preceding rate"):
        application.execute("record_rate_change", {"contract_id": "con_1", "component_id": "meter",
            "effective_date": "2026-02-10", "unit_rate": "0.017", "rationale": "Unchanged rate"}, period="2026-02")
    application.execute("record_usage", {"contract_id": "con_1", "obligation_id": "units",
                                          "effective_date": "2026-02-01", "quantity": "1"}, period="2026-02")
    application.execute("record_usage", {"contract_id": "con_1", "obligation_id": "units",
                                          "effective_date": "2026-02-12", "quantity": "20"}, period="2026-02")
    second_change = application.execute("record_rate_change", {"contract_id": "con_1", "component_id": "meter",
        "effective_date": "2026-02-10", "unit_rate": "0.019",
        "rationale": "Approved February price remains directly tied to units delivered."}, period="2026-02")["result"]["change_set_id"]
    february = application.state(period="2026-02")["report"]
    assert february["summary"]["revenue"] == "0.40"
    assert february["summary"]["recognized_to_date"] == "15.57"
    support = load_workbook(io.BytesIO(export_bytes(application.state(period="2026-02"))), read_only=True)
    monthly_row = list(support["Metered monthly revenue"].values)[-1]
    assert monthly_row[:4] == ("con_1", "2026-02", "21", "0.397")
    assert monthly_row[4] == 0.4
    assert [(row[4], row[5], row[6]) for row in list(support["Metered usage valuation"].values)[-2:]] == [
        ("1", "0.017", "0.017"), ("20", "0.019", "0.380")]
    assert list(support["Metered rate history"].values)[-1][4] == second_change
    support.close()
    application.execute("close_period", {"period": "2026-01", "review_dispositions": accept_review_items(application, "2026-01")}, period="2026-01")
    corrected = application.execute("correct_activity", {"target_change_set_id": second_change,
        "rationale": "Corrected approved February rate source", "replacement": {"contract_id": "con_1",
        "component_id": "meter", "effective_date": "2026-02-10", "unit_rate": "0.023"}}, period="2026-02")
    assert corrected["state"]["report"]["summary"]["revenue"] == "0.48"
    assert application.state(period="2026-01")["report"]["summary"]["revenue"] == "15.17"
    with pytest.raises(ValueError, match="closed period"):
        application.execute("correct_activity", {"target_change_set_id": first_change,
            "rationale": "Retrospective correction", "replacement": {"contract_id": "con_1",
            "component_id": "meter", "effective_date": "2026-01-15", "unit_rate": "0.020"}}, period="2026-02")
    with pytest.raises(ValueError, match="closed period"):
        application.execute("record_rate_change", {"contract_id": "con_1", "component_id": "meter",
            "effective_date": "2026-01-25", "unit_rate": "0.021",
            "rationale": "Late rate change after the accepted close"}, period="2026-02")


def test_metered_rate_and_units_import_from_structured_workbook(tmp_path):
    application = app(tmp_path)
    book = load_workbook(io.BytesIO(template_bytes()))
    book["Customers"].append(["cus_1", "Customer"])
    book["Contracts"].append(["con_1", "cus_1", "Metered service", "2026-01-01", "2026-12-31"])
    sheet = book["Consideration"]
    headers = [cell.value for cell in sheet[1]]
    component = {"contract_id": "con_1", "id": "meter", "label": "Service units", "kind": "metered",
                 "amount": "0", "unit_rate": "1.25", "pricing_basis": "right_to_invoice", "rounding_period": "calendar_month",
                 "rationale": "Each measured unit has the same customer value."}
    sheet.append([component.get(key) for key in headers])
    book["Obligations"].append(["con_1", "units", "Units", "service", "1", "metered", "2026-01-01", "2026-12-31"])
    book["Usage"].append(["con_1", "units", "2026-01-10", "100", "MTR-1", "", "meter:jan-1"])
    data = io.BytesIO()
    book.save(data)
    book.close()
    preview = preview_import_bytes(application, data.getvalue(), period="2026-01")
    assert preview["state"]["report"]["summary"]["revenue"] == "125.00"
    committed = import_bytes(application, data.getvalue(), period="2026-01")
    assert committed["state"]["report"]["summary"]["revenue"] == "125.00"
    support = load_workbook(io.BytesIO(export_bytes(committed["state"])), read_only=True)
    allocation = list(support["Component allocation"].values)
    assert allocation[1][-3:] == ("1.25", "right_to_invoice", "calendar_month")
    support.close()


def test_metered_rate_change_import_and_source_correction(tmp_path):
    application = app(tmp_path)
    book = load_workbook(io.BytesIO(template_bytes()))
    book["Customers"].append(["cus_1", "Customer"])
    book["Contracts"].append(["con_1", "cus_1", "Metered service", "2026-01-01", "2026-12-31"])
    consideration = {"contract_id": "con_1", "id": "meter", "label": "Service units", "kind": "metered",
                     "amount": "0", "unit_rate": "0.015", "pricing_basis": "right_to_invoice",
                     "rounding_period": "calendar_month", "rationale": "Invoice value tracks delivered units."}
    headers = [cell.value for cell in book["Consideration"][1]]
    book["Consideration"].append([consideration.get(key) for key in headers])
    book["Obligations"].append(["con_1", "units", "Units", "service", "0", "metered", "2026-01-01", "2026-12-31"])
    book["Rate Changes"].append(["con_1", "meter", "2026-01-15", "0.017",
                                 "New approved unit rate reflects delivered value.", "rate:jan-15"])
    book["Usage"].append(["con_1", "units", "2026-01-10", "1001", "MTR-1", "", "usage:jan-10"])
    book["Usage"].append(["con_1", "units", "2026-01-15", "9", "MTR-2", "", "usage:jan-15"])
    data = io.BytesIO()
    book.save(data)
    book.close()
    preview = preview_import_bytes(application, data.getvalue(), period="2026-01")
    assert preview["state"]["report"]["summary"]["revenue"] == "15.17"
    committed = import_bytes(application, data.getvalue(), period="2026-01")
    assert committed["state"]["report"]["summary"]["revenue"] == "15.17"
    correction_book = load_workbook(io.BytesIO(template_bytes()))
    correction = {"source_id": "correction:rate-1", "target_source_id": "rate:jan-15",
                  "activity_type": "rate_change", "contract_id": "con_1", "component_id": "meter",
                  "effective_date": "2026-01-15", "unit_rate": "0.019",
                  "rationale": "Corrected approved rate schedule"}
    correction_headers = [cell.value for cell in correction_book["Corrections"][1]]
    correction_book["Corrections"].append([correction.get(key) for key in correction_headers])
    corrected_data = io.BytesIO()
    correction_book.save(corrected_data)
    correction_book.close()
    corrected = import_bytes(application, corrected_data.getvalue(), period="2026-01")
    assert corrected["state"]["report"]["summary"]["revenue"] == "15.19"
    assert [row["type"] for row in corrected["state"]["contracts"][0]["activities"]].count("rate_change") == 1
    assert corrected["state"]["contracts"][0]["activities"][0]["unit_rate"] == "0.019"


def test_customer_external_reference_is_unique_within_its_source_system(tmp_path):
    application = app(tmp_path)
    application.execute("create_customer", {"id": "first", "name": "Legal entity A", "reference": "42", "source_system": "CRM"})
    application.execute("create_customer", {"id": "second", "name": "Legal entity B", "reference": "42", "source_system": "ERP"})
    with pytest.raises(ValueError, match="source system and external reference"):
        application.execute("create_customer", {"name": "Duplicate", "reference": "42", "source_system": "CRM"})
    with pytest.raises(ValueError, match="source system and external reference"):
        application.execute("edit_details", {"entity_id": "second", "source_system": "CRM"})


def test_structured_amendment_precedes_usage_and_correction_in_one_import(tmp_path):
    application = app(tmp_path)
    seed(application)
    book = load_workbook(io.BytesIO(template_bytes()))
    book["Amendments"].append(["erp:amend-1", "con_1", "2026-10-01", "prospective", "no", "yes", "Added finite-use service"])
    book["Amendment Obligations"].append(["erp:amend-1", "pob_1", "Service", "service", "12000", "exact_days", "2026-09-01", "2027-08-31", "", "", "", ""])
    book["Amendment Obligations"].append(["erp:amend-1", "usage_new", "Usage", "service", "12000", "usage", "2026-10-01", "2027-08-31", "100", "", "", ""])
    book["Usage"].append(["con_1", "usage_new", "2026-10-02", "10", "MTR-1", "", "erp:units-1"])
    book["Corrections"].append(["erp:correction-1", "", "erp:units-1", "usage", "con_1", "usage_new", "2026-10-02", "", "", "8", "MTR-1", "Corrected meter file"])
    data = io.BytesIO(); book.save(data); book.close()
    preview = preview_import_bytes(application, data.getvalue(), period="2026-10")
    assert preview["result"]["imported"] == 3
    assert application.state(period="2026-10")["contracts"][0]["activities"] == []
    imported = import_bytes(application, data.getvalue(), period="2026-10", expected_frontier=preview["frontier"], expected_hash=preview["result"]["file_hash"])
    activities = imported["state"]["contracts"][0]["activities"]
    assert [item["type"] for item in activities] == ["modification", "usage"]
    assert activities[1]["quantity"] == "8"
    assert activities[1]["corrects"] == imported["result"]["rows"][1]["change_set_id"]


def test_future_activity_preview_moves_to_effective_month_and_shows_billing(tmp_path):
    application = app(tmp_path)
    seed(application)
    before = application.state(period="2026-09")
    preview = application.preview("record_billing", {"contract_id": "con_1", "effective_date": "2026-12-01", "amount": "1200"}, period="2026-09")
    assert preview["selected_period"] == "2026-09"
    assert preview["focus_period"] == "2026-12"
    assert preview["comparison"]["summary"]["billings"] == "1200.00"
    assert "2026-12" in preview["comparison"]["affected_periods"]
    assert application.state(period="2026-09") == before


def test_demo_walkthrough_has_five_contracts_scenario_and_balanced_journal(tmp_path):
    application = app(tmp_path)
    result = load_demo(application)
    state = result["state"]
    assert len(state["contracts"]) == 5
    assert len([s for s in state["scenarios"] if s["id"] != "main"]) == 1
    assert state["report"]["catch_ups"]
    assert sum(j["debit_minor"] for j in state["report"]["journals"]) == sum(j["credit_minor"] for j in state["report"]["journals"])


def test_apply_preview_from_scenario_shows_main_impact_without_writes(tmp_path):
    application = app(tmp_path)
    seed(application)
    sid = application.execute("create_scenario", {"name": "Billing"})["result"]["id"]
    application.execute("record_billing", {"contract_id": "con_1", "effective_date": "2026-09-01", "amount": "100"}, scenario_id=sid)
    before = application.state(sid, "2026-09")
    preview = application.preview("apply_scenario", {"scenario_id": sid}, scenario_id=sid, period="2026-09")
    assert preview["comparison"]["summary"]["billings"] == "100.00"
    assert preview["state"]["scenario_id"] == "main"
    assert application.state(sid, "2026-09") == before
    assert application.state("main", "2026-09")["report"]["summary"]["billings"] == "0.00"


def test_closed_period_allows_new_future_contract_but_blocks_backdated_billing(tmp_path):
    application = app(tmp_path)
    contract = seed(application)
    closed = application.execute("close_period", {"period": "2026-09", "review_dispositions": accept_review_items(application, "2026-09")}, period="2026-09")["state"]["report"]
    future = deepcopy(contract)
    future.update(id="future", start_date="2026-10-01")
    future["obligations"][0]["start_date"] = "2026-10-01"
    application.execute("create_contract", future, period="2026-10")
    after = application.state(period="2026-09")["report"]
    assert after["contracts"] == closed["contracts"]
    assert after["summary"] == closed["summary"]
    with pytest.raises(ValueError, match="closed period"):
        application.execute("record_billing", {"contract_id": "future", "effective_date": "2026-09-30", "amount": "100"})


def test_closed_period_protects_obligation_schedule_even_when_totals_match(tmp_path):
    application = app(tmp_path)
    contract = seed(application)
    obligations = [deepcopy(contract["obligations"][0]), deepcopy(contract["obligations"][0])]
    obligations[1]["id"] = "pob_2"
    for obligation in obligations:
        obligation["method"] = "monthly"
    application.execute("modify_contract", {"contract_id": "con_1", "effective_date": "2026-09-01", "treatment": "catch_up", "obligations": obligations, "rationale": "Split services"})
    application.execute("close_period", {"period": "2026-09", "review_dispositions": accept_review_items(application, "2026-09")})
    obligations[0]["ssp"] = "36000"
    with pytest.raises(ValueError, match="closed period"):
        application.execute("modify_contract", {"contract_id": "con_1", "effective_date": "2026-09-01", "treatment": "catch_up", "obligations": obligations, "rationale": "Reallocate services"})


@pytest.mark.parametrize("command,payload", [
    ("create_customer", {"id": 42, "name": "Invalid ID"}),
    ("create_customer", {"name": "Invalid text", "email": {"bad": True}}),
    ("create_contract", {"name": "Invalid terms", "customer_id": "cus_1", "start_date": "2026-09-01", "end_date": "2027-08-31", "consideration": [None], "obligations": []}),
    ("set_policy", {"accounts": []}),
    ("create_scenario", []),
    ("add_note", {"id": "note_1", "body": "Duplicate"}),
    ("edit_note", {"id": "note_1", "body": {"invalid": True}}),
    ("edit_note", {"id": "note_1", "due_date": 123}),
    ("edit_note", {"id": "note_1", "due_date": "2026-02-30"}),
])
def test_malformed_commands_are_validation_errors_and_leave_history_unchanged(tmp_path, command, payload):
    application = app(tmp_path)
    seed(application)
    application.execute("add_note", {"id": "note_1", "kind": "task", "body": "Review"})
    before = application.state()
    with pytest.raises(ValueError):
        application.execute(command, payload)
    assert application.state() == before
    application.reports()


def test_workspace_connections_close_on_context_exit(tmp_path):
    application = app(tmp_path)
    with application.workspace.connect() as db:
        connection = db
    with pytest.raises(sqlite3.ProgrammingError, match="closed"):
        connection.execute("SELECT 1")


@pytest.mark.parametrize("headers,row", [
    (["id", "name", "name"], ["cus", "Original", "Overwritten"]),
    (["id", "name", ""], ["cus", "Name", "Silently discarded"]),
])
def test_import_rejects_ambiguous_columns(tmp_path, headers, row):
    application = app(tmp_path)
    book = load_workbook(io.BytesIO(template_bytes()))
    sheet = book["Customers"]
    sheet.delete_rows(1)
    sheet.append(headers)
    sheet.append(row)
    data = io.BytesIO()
    book.save(data)
    book.close()
    with pytest.raises(ValueError):
        import_bytes(application, data.getvalue())
    assert application.state()["customers"] == []


def test_sql_bounds_large_values_and_returns_json_safe_blobs(tmp_path):
    application = app(tmp_path)
    assert application.query("SELECT x'00ff' AS evidence")["rows"] == [[{"hex": "00ff"}]]
    with pytest.raises(ValueError, match="too big"):
        application.query("SELECT zeroblob(1000001)")


def test_close_backup_preserves_precommand_state_and_retry_creates_no_backup(tmp_path):
    application = app(tmp_path)
    seed(application)
    before = application.state(period="2026-09")
    close_payload = {"period": "2026-09", "review_dispositions": accept_review_items(application, "2026-09")}
    closed = application.execute("close_period", close_payload, idempotency_key="close-september", period="2026-09")
    backup = Application(closed["result"]["backup_path"])
    assert backup.state(period="2026-09") == before
    backups = list((application.workspace.path / "backups").iterdir())
    retried = application.execute("close_period", close_payload, idempotency_key="close-september", period="2026-09")
    assert retried["result"]["replayed"] is True
    assert list((application.workspace.path / "backups").iterdir()) == backups


def test_apply_retry_returns_main_and_does_not_duplicate_activity(tmp_path):
    application = app(tmp_path)
    seed(application)
    sid = application.execute("create_scenario", {"name": "Billing"})["result"]["id"]
    application.execute("record_billing", {"contract_id": "con_1", "effective_date": "2026-09-01", "amount": "100"}, scenario_id=sid)
    first = application.execute("apply_scenario", {"scenario_id": sid}, scenario_id=sid, idempotency_key="apply-billing", period="2026-09")
    retry = application.execute("apply_scenario", {"scenario_id": sid}, scenario_id=sid, idempotency_key="apply-billing", period="2026-09")
    assert retry["state"] == first["state"]
    assert retry["state"]["scenario_id"] == "main"
    assert retry["state"]["report"]["summary"]["billings"] == "100.00"


def test_account_mapping_applies_from_its_effective_month_without_rewriting_prior_journals(tmp_path):
    application = app(tmp_path)
    seed(application)
    september = application.state(period="2026-09")
    application.execute("set_policy", {
        "effective_period": "2026-10", "accounts": {"revenue": "4100"},
        "rationale": "New revenue account from October",
    }, period="2026-10")
    earlier = application.state(period="2026-09")
    october = application.state(period="2026-10")
    november = application.state(period="2026-11")
    assert earlier["report"]["journals"] == september["report"]["journals"]
    assert earlier["policy"]["version"] == 1
    assert october["policy"]["version"] == 2
    assert october["report"]["policy_effective_period"] == "2026-10"
    assert {row["account"] for row in october["report"]["journals"] if row["role"] == "revenue"} == {"4100"}
    assert november["policy"]["accounts"]["revenue"] == "4100"


def test_obligation_revenue_override_splits_journal_without_splitting_contract_balance(tmp_path):
    application = app(tmp_path)
    application.execute("create_customer", {"id": "cus_1", "name": "Customer"})
    application.execute("create_contract", {
        "id": "con_1", "name": "Setup and service", "customer_id": "cus_1",
        "start_date": "2026-09-01", "end_date": "2026-09-30",
        "consideration": [{"id": "price", "kind": "fixed", "amount": "1000"}],
        "obligations": [
            {"id": "setup", "name": "Setup", "kind": "implementation", "ssp": "500", "method": "point_in_time", "start_date": "2026-09-01", "end_date": "2026-09-30"},
            {"id": "service", "name": "Service", "kind": "service", "ssp": "500", "method": "exact_days", "start_date": "2026-09-01", "end_date": "2026-09-30"},
        ],
    }, period="2026-09")
    application.execute("record_milestone", {"contract_id": "con_1", "obligation_id": "setup", "effective_date": "2026-09-01", "percentage": "100"}, period="2026-09")
    application.execute("record_billing", {"contract_id": "con_1", "effective_date": "2026-09-01", "amount": "1000"}, period="2026-09")
    application.execute("set_policy", {"effective_period": "2026-09", "account_overrides": {"contracts": {}, "obligations": {"con_1": {"setup": "4100"}}}}, period="2026-09")
    report = application.state(period="2026-09")["report"]
    revenue = {row["account"]: row for row in report["journals"] if row["role"] == "revenue"}
    assert {account: row["credit"] for account, row in revenue.items()} == {"4000": "500.00", "4100": "500.00"}
    assert revenue["4100"]["obligation_ids"] == ["setup"]
    assert report["contracts"][0]["deferred_revenue"] == "0.00"
    assert sum(row["debit_minor"] - row["credit_minor"] for row in report["journals"]) == 0


def test_deferred_account_change_requires_treatment_and_transfers_opening_balance(tmp_path):
    application = app(tmp_path)
    seed(application)
    application.execute("record_billing", {"contract_id": "con_1", "effective_date": "2026-09-01", "amount": "12000"}, period="2026-09")
    september = application.state(period="2026-09")["report"]
    assert september["contracts"][0]["contract_asset"] == "0.00"
    application.execute("set_policy", {"effective_period": "2026-10", "accounts": {"contract_asset": "1210"}}, period="2026-10")
    with pytest.raises(ValueError, match="existing balance-sheet balances"):
        application.execute("set_policy", {"effective_period": "2026-10", "accounts": {"deferred_revenue": "2310"}}, period="2026-10")
    application.execute("set_policy", {"effective_period": "2026-10", "accounts": {"deferred_revenue": "2310"}, "account_transition": "transfer"}, period="2026-10")
    with pytest.raises(ValueError, match="existing balance-sheet balances"):
        application.execute("set_policy", {"effective_period": "2026-10", "accounts": {"deferred_revenue": "2320"}}, period="2026-10")
    application.execute("set_policy", {"effective_period": "2026-10", "accounts": {"revenue": "4100"}}, period="2026-10")
    october = application.state(period="2026-10")["report"]
    assert application.state(period="2026-09")["report"]["journals"] == september["journals"]
    assert october["account_transitions"] == [{"contract_id": "con_1", "role": "deferred_revenue", "from_account": "2300", "to_account": "2310", "opening_balance": september["contracts"][0]["deferred_revenue"], "treatment": "transfer"}]
    assert october["account_positions"] == [{"contract_id": "con_1", "role": "deferred_revenue", "account": "2310", "dimensions": {}, "account_profile_id": None, "balance": october["contracts"][0]["deferred_revenue"]}]
    posted = {}
    for row in september["journals"] + october["journals"]:
        if row["role"] == "deferred_revenue":
            posted[row["account"]] = posted.get(row["account"], 0) + row["credit_minor"] - row["debit_minor"]
    assert posted["2300"] == 0
    assert posted["2310"] == int(Decimal(october["contracts"][0]["deferred_revenue"]) * 100)
    assert sum(row["debit_minor"] - row["credit_minor"] for row in october["journals"]) == 0
    book = load_workbook(io.BytesIO(export_bytes(application.state(period="2026-10"))), read_only=True)
    assert book["Account positions"]["C2"].value == "2310"
    assert Decimal(str(book["Account positions"]["F2"].value)) == Decimal(october["contracts"][0]["deferred_revenue"])
    book.close()


def test_balance_account_change_with_zero_opening_needs_no_transfer(tmp_path):
    application = app(tmp_path)
    application.execute("create_customer", {"id": "customer", "name": "Customer"})
    application.execute("create_contract", {
        "id": "contract", "name": "Quarterly service", "customer_id": "customer",
        "start_date": "2026-09-01", "end_date": "2026-12-31",
        "consideration": [{"id": "price", "kind": "fixed", "amount": "400.00"}],
        "obligations": [{"id": "service", "name": "Service", "kind": "service", "ssp": "400.00", "method": "monthly", "start_date": "2026-09-01", "end_date": "2026-12-31"}],
    }, period="2026-09")
    application.execute("record_billing", {"contract_id": "contract", "effective_date": "2026-09-01", "amount": "100.00"}, period="2026-09")
    september = application.state(period="2026-09")["report"]
    assert september["contracts"][0]["deferred_revenue"] == "0.00"
    assert september["contracts"][0]["contract_asset"] == "0.00"

    application.execute("set_policy", {"effective_period": "2026-10", "accounts": {"deferred_revenue": "2310"}}, period="2026-10")
    application.execute("set_policy", {"effective_period": "2026-10", "accounts": {"deferred_revenue": "2320"}}, period="2026-10")
    application.execute("record_billing", {"contract_id": "contract", "effective_date": "2026-10-01", "amount": "300.00"}, period="2026-10")
    october = application.state(period="2026-10")["report"]
    assert october["contracts"][0]["deferred_revenue"] == "200.00"
    assert october["account_transitions"] == []
    assert {row["account"] for row in october["journals"] if row["role"] == "deferred_revenue"} == {"2320"}
    assert sum(row["debit_minor"] - row["credit_minor"] for row in october["journals"]) == 0


def test_reviewed_deferred_account_runoff_allocates_each_month_without_transfer(tmp_path):
    application = app(tmp_path)
    application.execute("create_customer", {"id": "customer", "name": "Customer"})
    application.execute("create_contract", {
        "id": "contract", "name": "Quarterly service", "customer_id": "customer",
        "start_date": "2026-09-01", "end_date": "2026-12-31",
        "consideration": [{"id": "price", "kind": "fixed", "amount": "400.00"}],
        "obligations": [{"id": "service", "name": "Service", "kind": "service", "ssp": "400.00", "method": "monthly", "start_date": "2026-09-01", "end_date": "2026-12-31"}],
    }, period="2026-09")
    application.execute("record_billing", {"contract_id": "contract", "effective_date": "2026-09-01", "amount": "300.00"}, period="2026-09")
    application.execute("record_billing", {"contract_id": "contract", "effective_date": "2026-10-01", "amount": "100.00"}, period="2026-10")
    september_journals = application.state(period="2026-09")["report"]["journals"]
    application.execute("set_policy", {"effective_period": "2026-10", "accounts": {"deferred_revenue": "2310"},
                                       "account_transition": "runoff", "rationale": "Keep the existing liability in its historical account."}, period="2026-10")
    provisional = application.state(period="2026-10")["report"]
    assert provisional["account_transitions"][0]["treatment"] == "runoff"
    assert not any(":transfer" in row["id"] for row in provisional["journals"])
    assert len(provisional["runoff_pending"]) == 1
    assert next(check for check in application.reports(period="2026-10")["checks"] if check["id"] == "account_runoff")["status"] == "block"
    assert any(item["period"] == "2026-10" for item in application.state(period="2026-11")["report"]["runoff_unresolved"])
    assert next(check for check in application.reports(period="2026-11")["checks"] if check["id"] == "account_runoff")["status"] == "block"
    january = application.state(period="2027-01")["report"]
    assert january["runoff_pending"] == [] and any(item["period"] == "2026-10" for item in january["runoff_unresolved"])
    assert next(check for check in application.reports(period="2027-01")["checks"] if check["id"] == "account_runoff")["status"] == "block"
    draft_book = load_workbook(io.BytesIO(export_bytes(application.state(period="2026-10"))), read_only=True)
    assert draft_book["Account runoff"]["A2"].value == "Provisional - do not post"
    assert draft_book["Posting guide"]["A2"].value == "Draft account runoff"
    draft_book.close()
    routes = [
        {"account": "2300", "dimensions": {}, "balance": "100.00"},
        {"account": "2310", "dimensions": {}, "balance": "100.00"},
    ]
    with pytest.raises(ValueError, match="does not reconcile"):
        application.execute("record_account_runoff", {"contract_id": "contract", "role": "deferred_revenue", "period": "2026-10",
                                                      "positions": [{**routes[0], "balance": "200.00"}, routes[1]], "rationale": "Invalid split"}, period="2026-10")
    application.execute("record_account_runoff", {"contract_id": "contract", "role": "deferred_revenue", "period": "2026-10",
                                                  "positions": routes, "rationale": "Release half of the historical deferral this month."}, period="2026-10")
    october = application.state(period="2026-10")["report"]
    assert october["runoff_pending"] == []
    assert next(check for check in application.reports(period="2026-10")["checks"] if check["id"] == "account_runoff")["status"] == "pass"
    assert {(row["account"], row["balance"]) for row in october["account_positions"]} == {("2300", "100.00"), ("2310", "100.00")}
    assert {(row["account"], row["debit"], row["credit"]) for row in october["journals"] if row["role"] == "deferred_revenue"} == {("2300", "100.00", "0.00"), ("2310", "0.00", "100.00")}
    assert sum(row["debit_minor"] - row["credit_minor"] for row in october["journals"]) == 0
    accepted_book = load_workbook(io.BytesIO(export_bytes(application.state(period="2026-10"))), read_only=True)
    assert [row[0] for row in list(accepted_book["Account runoff"].values)[1:]] == ["Recorded", "Recorded"]
    assert {row[4] for row in list(accepted_book["Account runoff"].values)[1:]} == {"2300", "2310"}
    accepted_book.close()
    assert application.state(period="2026-09")["report"]["journals"] == september_journals
    assert len(application.state(period="2026-11")["report"]["runoff_pending"]) == 1
    application.execute("record_account_runoff", {"contract_id": "contract", "role": "deferred_revenue", "period": "2026-11",
                                                  "positions": [{**routes[0], "balance": "0.00"}, routes[1]],
                                                  "rationale": "Remaining historical obligation delivered in November."}, period="2026-11")
    november = application.state(period="2026-11")["report"]
    assert november["runoff_pending"] == []
    assert [(row["account"], row["balance"]) for row in november["account_positions"]] == [("2310", "100.00")]
    with pytest.raises(ValueError, match="No historical"):
        application.execute("record_account_runoff", {"contract_id": "contract", "role": "deferred_revenue", "period": "2026-10",
                                                      "positions": [{**routes[0], "balance": "0.00"}, {**routes[1], "balance": "200.00"}],
                                                      "rationale": "Would invalidate the accepted November allocation."}, period="2026-10")
    assert {(row["account"], row["balance"]) for row in application.state(period="2026-10")["report"]["account_positions"]} == {("2300", "100.00"), ("2310", "100.00")}
    assert application.state(period="2026-12")["report"]["runoff_pending"] == []
    assert application.state(period="2026-12")["report"]["account_positions"] == []
    application.execute("close_period", {"period": "2026-10", "review_dispositions": accept_review_items(application, "2026-10")}, period="2026-10")
    checkpoint = application.state(period="2026-10")["report"]
    assert checkpoint["runoff_pending"] == []
    assert {(row["account"], row["balance"]) for row in checkpoint["account_positions"]} == {("2300", "100.00"), ("2310", "100.00")}
    with pytest.raises(ValueError, match="Reopen closed periods"):
        application.execute("record_account_runoff", {"contract_id": "contract", "role": "deferred_revenue", "period": "2026-10",
                                                      "positions": routes, "rationale": "Cannot revise an accepted close."}, period="2026-10")


def test_asset_runoff_can_retain_old_dimension_and_later_transfer_every_route(tmp_path):
    application = app(tmp_path)
    application.execute("create_customer", {"id": "customer", "name": "Customer"})
    application.execute("create_contract", {
        "id": "contract", "name": "Quarterly service", "customer_id": "customer",
        "start_date": "2026-09-01", "end_date": "2026-12-31",
        "consideration": [{"id": "price", "kind": "fixed", "amount": "400.00"}],
        "obligations": [{"id": "service", "name": "Service", "kind": "service", "ssp": "400.00", "method": "monthly", "start_date": "2026-09-01", "end_date": "2026-12-31"}],
    }, period="2026-09")
    application.execute("set_policy", {"effective_period": "2026-09", "account_profiles": {"old": {"name": "Old", "accounts": {"contract_asset": "1200"}, "dimensions": {"Department": "Old"}}},
                                       "profile_assignments": {"contract": "old"}}, period="2026-09")
    application.execute("set_policy", {"effective_period": "2026-10", "account_profiles": {"old": {"name": "Old", "accounts": {"contract_asset": "1200"}, "dimensions": {"Department": "Old"}},
                                                                                       "new": {"name": "New", "accounts": {"contract_asset": "1210"}, "dimensions": {"Department": "New"}}},
                                       "profile_assignments": {"contract": "new"}, "account_transition": "runoff", "rationale": "Carry the old asset by department."}, period="2026-10")
    pending = application.state(period="2026-10")["report"]["runoff_pending"]
    assert len(pending) == 1 and pending[0]["role"] == "contract_asset"
    positions = [{"account": row["account"], "dimensions": row["dimensions"], "account_profile_id": row.get("account_profile_id"),
                  "balance": "50.00" if row["account"] == "1200" else "150.00"} for row in pending[0]["positions"]]
    with pytest.raises(ValueError, match="cannot increase"):
        application.execute("record_account_runoff", {"contract_id": "contract", "role": "contract_asset", "period": "2026-10",
                                                      "positions": [{**row, "balance": "150.00" if row["account"] == "1200" else "50.00"} for row in positions],
                                                      "rationale": "Invalid increase to an old account."}, period="2026-10")
    application.execute("record_account_runoff", {"contract_id": "contract", "role": "contract_asset", "period": "2026-10", "positions": positions,
                                                  "rationale": "Allocate the old asset to the remaining prior service."}, period="2026-10")
    october = application.state(period="2026-10")["report"]
    assert {(row["account"], row["dimensions"]["Department"], row["balance"]) for row in october["account_positions"]} == {
        ("1200", "Old", "50.00"), ("1210", "New", "150.00")}
    assert {(row["account"], row["debit"], row["credit"]) for row in october["journals"] if row["role"] == "contract_asset"} == {
        ("1200", "0.00", "50.00"), ("1210", "150.00", "0.00")}
    assert sum(row["debit_minor"] - row["credit_minor"] for row in october["journals"]) == 0
    application.execute("set_policy", {"effective_period": "2026-11", "account_profiles": {"old": {"name": "Old", "accounts": {"contract_asset": "1200"}, "dimensions": {"Department": "Old"}},
                                                                                       "new": {"name": "New", "accounts": {"contract_asset": "1220"}, "dimensions": {"Department": "New"}}},
                                       "account_transition": "transfer", "rationale": "Transfer the full asset to the replacement account."}, period="2026-11")
    november = application.state(period="2026-11")["report"]
    assert len([row for row in november["account_transitions"] if row["role"] == "contract_asset"]) == 2
    assert [(row["account"], row["balance"]) for row in november["account_positions"]] == [("1220", "300.00")]
    assert sum(row["debit_minor"] - row["credit_minor"] for row in november["journals"]) == 0


def test_runoff_policy_before_population_does_not_replay_empty_history(tmp_path):
    application = app(tmp_path)
    application.execute("create_customer", {"id": "customer", "name": "Customer"})
    application.execute("create_contract", {
        "id": "contract", "name": "New service", "customer_id": "customer",
        "start_date": "2026-09-01", "end_date": "2026-09-30",
        "consideration": [{"id": "price", "kind": "fixed", "amount": "100.00"}],
        "obligations": [{"id": "service", "name": "Service", "kind": "service", "ssp": "100.00", "method": "monthly", "start_date": "2026-09-01", "end_date": "2026-09-30"}],
    }, period="2026-09")
    application.execute("set_policy", {"effective_period": "1900-01", "accounts": {"contract_asset": "1210"},
                                       "account_transition": "runoff", "rationale": "Initial route predates this accounting population."}, period="2026-09")
    report = application.state(period="2026-09")["report"]
    assert report["runoff_active"] is False
    assert report["summary"]["revenue"] == "100.00"


def test_reusable_account_profile_maps_multiple_contracts_and_balances_dimension_change(tmp_path):
    application = app(tmp_path)
    seed(application)
    application.execute("create_contract", {
        "id": "con_2", "name": "Second service", "customer_id": "cus_1", "start_date": "2026-09-01", "end_date": "2027-08-31",
        "consideration": [{"id": "price_2", "kind": "fixed", "amount": "1200.00"}],
        "obligations": [{"id": "pob_2", "name": "Second service", "kind": "service", "ssp": "1200.00", "method": "exact_days", "start_date": "2026-09-01", "end_date": "2027-08-31"}],
    }, period="2026-09")
    application.execute("record_billing", {"contract_id": "con_1", "effective_date": "2026-09-01", "amount": "12000.00"}, period="2026-09")
    profile = {"name": "Services", "accounts": {"revenue": "4100", "deferred_revenue": "2310"}, "dimensions": {"Department": "Recurring", "Project": "P-17"}}
    application.execute("set_policy", {
        "effective_period": "2026-09", "account_profiles": {"services": profile},
        "profile_assignments": {"con_1": "services", "con_2": "services"}, "account_transition": "transfer",
        "account_overrides": {"contracts": {}, "obligations": {"con_1": {"pob_1": "4190"}}},
    }, period="2026-09")
    september = application.state(period="2026-09")["report"]
    assert {row["account"] for row in september["journals"] if row["contract_id"] == "con_1" and row["role"] == "revenue"} == {"4190"}
    assert {row["account"] for row in september["journals"] if row["contract_id"] == "con_2" and row["role"] == "revenue"} == {"4100"}
    assert all(row["dimensions"] == {"Department": "Recurring", "Project": "P-17"} for row in september["journals"])
    assert all(row["account_profile_id"] == "services" for row in september["journals"])
    with pytest.raises(ValueError, match="existing balance-sheet balances"):
        application.execute("set_policy", {"effective_period": "2026-10", "account_profiles": {"services": {**profile, "dimensions": {"Department": "New division", "Project": "P-17"}}}}, period="2026-10")
    application.execute("set_policy", {"effective_period": "2026-10", "account_profiles": {"services": {**profile, "dimensions": {"Department": "New division", "Project": "P-17"}}}, "account_transition": "transfer"}, period="2026-10")
    october = application.state(period="2026-10")["report"]
    assert application.state(period="2026-09")["report"]["journals"] == september["journals"]
    transition = next(item for item in october["account_transitions"] if item["contract_id"] == "con_1" and item["role"] == "deferred_revenue")
    assert transition["from_account"] == transition["to_account"] == "2310"
    assert transition["from_dimensions"]["Department"] == "Recurring"
    assert transition["to_dimensions"]["Department"] == "New division"
    transfers = [row for row in october["journals"] if row["contract_id"] == "con_1" and ":transfer" in row["id"]]
    assert len(transfers) == 2
    assert {row["dimensions"]["Department"] for row in transfers} == {"Recurring", "New division"}
    assert sum(row["debit_minor"] - row["credit_minor"] for row in october["journals"]) == 0
    book = load_workbook(io.BytesIO(export_bytes(application.state(period="2026-10"))), read_only=True)
    headings = [cell.value for cell in book["Journal entries"][1]]
    assert "dimension:Department" in headings and "dimension:Project" in headings
    rows = list(book["Journal entries"].values)
    assert any(row[headings.index("dimension:Department")] == "New division" and row[headings.index("account_profile_id")] == "services" for row in rows[1:])
    assert book["Account profiles"]["A2"].value == "services"
    assert book["Profile assignments"]["B2"].value == "services"
    book.close()
    next_profile = {**profile, "accounts": {"revenue": "4100", "deferred_revenue": "2320"}, "dimensions": {"Department": "New division", "Project": "P-17"}}
    with pytest.raises(ValueError, match="existing balance-sheet balances"):
        application.execute("set_policy", {"effective_period": "2026-11", "account_profiles": {"services": next_profile}}, period="2026-11")
    application.execute("set_policy", {"effective_period": "2026-11", "account_profiles": {"services": next_profile}, "account_transition": "transfer"}, period="2026-11")
    november = application.state(period="2026-11")["report"]
    assert any(item["contract_id"] == "con_1" and item["from_account"] == "2310" and item["to_account"] == "2320" for item in november["account_transitions"])
    with pytest.raises(ValueError, match="existing accounting profiles"):
        application.execute("set_policy", {"effective_period": "2026-12", "account_profiles": {}}, period="2026-12")
    with pytest.raises(ValueError, match="known journal roles"):
        application.execute("set_policy", {"effective_period": "2026-12", "account_profiles": {"services": {**profile, "accounts": {"unknown": "9999"}}}}, period="2026-12")


def test_obligation_profiles_route_revenue_dimensions_and_flag_segment_imbalance(tmp_path):
    application = app(tmp_path)
    application.execute("create_customer", {"id": "customer", "name": "Customer"})
    application.execute("create_contract", {
        "id": "contract", "name": "Two services", "customer_id": "customer", "start_date": "2026-01-01", "end_date": "2026-06-30",
        "consideration": [{"id": "fee", "kind": "fixed", "amount": "1200.00"}],
        "obligations": [
            {"id": "setup", "name": "Setup", "kind": "service", "ssp": "600.00", "method": "monthly", "start_date": "2026-01-01", "end_date": "2026-06-30"},
            {"id": "support", "name": "Support", "kind": "service", "ssp": "600.00", "method": "monthly", "start_date": "2026-01-01", "end_date": "2026-06-30"},
        ],
    }, period="2026-01")
    profiles = {
        "base": {"name": "Base", "accounts": {}, "dimensions": {"Department": "Head office", "Project": "P-17"}},
        "setup": {"name": "Setup", "accounts": {"revenue": "4100"}, "dimensions": {"Department": "Implementation"}},
        "support": {"name": "Support", "accounts": {"revenue": "4200"}, "dimensions": {"Department": "Recurring"}},
    }
    with pytest.raises(ValueError, match="existing accounting profiles"):
        application.execute("set_policy", {"effective_period": "2026-01", "account_profiles": profiles,
            "obligation_profile_assignments": {"contract": {"setup": "missing"}}}, period="2026-01")
    application.execute("set_policy", {"effective_period": "2026-01", "account_profiles": profiles,
        "profile_assignments": {"contract": "base"},
        "obligation_profile_assignments": {"contract": {"setup": "setup", "support": "support"}},
        "account_transition": "transfer"}, period="2026-01")
    application.execute("record_billing", {"contract_id": "contract", "effective_date": "2026-01-01", "amount": "1200.00"}, period="2026-01")
    state = application.state(period="2026-01")
    journals = state["report"]["journals"]
    revenue = {row["obligation_ids"][0]: row for row in journals if row["role"] == "revenue"}
    assert (revenue["setup"]["account"], revenue["setup"]["dimensions"], revenue["setup"]["account_profile_id"]) == ("4100", {"Department": "Implementation", "Project": "P-17"}, "setup")
    assert (revenue["support"]["account"], revenue["support"]["dimensions"], revenue["support"]["account_profile_id"]) == ("4200", {"Department": "Recurring", "Project": "P-17"}, "support")
    assert all(row["dimensions"] == {"Department": "Head office", "Project": "P-17"} for row in journals if row["role"] != "revenue")
    assert sum(row["debit_minor"] - row["credit_minor"] for row in journals) == 0
    assert any("not by dimension combination" in warning for warning in state["report"]["warnings"])
    assert {tuple(sorted(row["dimensions"].items())): row["net_debit"] for row in state["report"]["segment_imbalances"]} == {
        (("Department", "Head office"), ("Project", "P-17")): "200.00", (("Department", "Implementation"), ("Project", "P-17")): "-100.00", (("Department", "Recurring"), ("Project", "P-17")): "-100.00"}
    book = load_workbook(io.BytesIO(export_bytes(state)), read_only=True)
    assert list(book["Obligation profile assignments"].values)[1:] == [("contract", "setup", "setup"), ("contract", "support", "support")]
    assert book["Dimension balance"].max_row == 4
    book.close()
    application.execute("set_policy", {"effective_period": "2026-02", "account_profiles": {**profiles, "support": {**profiles["support"], "dimensions": {}}},
        "obligation_profile_assignments": {"contract": {"support": "support"}}}, period="2026-02")
    february = application.state(period="2026-02")
    assert february["policy"]["obligation_profile_assignments"] == {"contract": {"support": "support"}}
    assert next(row for row in february["report"]["journals"] if row["role"] == "revenue" and row["account"] == "4000")["dimensions"] == {"Department": "Head office", "Project": "P-17"}
    assert next(row for row in february["report"]["journals"] if row["role"] == "revenue" and row["account"] == "4200")["dimensions"] == {"Department": "Head office", "Project": "P-17"}
    assert application.state(period="2026-01")["report"]["journals"] == journals


def test_supplied_account_dimension_combinations_block_invalid_lines_and_expose_gaps(tmp_path):
    application = app(tmp_path)
    seed(application)
    application.execute("record_billing", {
        "contract_id": "con_1", "effective_date": "2026-09-01", "amount": "12000.00",
    }, period="2026-09")
    profile = {"recurring": {"name": "Recurring", "accounts": {}, "dimensions": {"Department": "Recurring"}}}
    application.execute("set_policy", {
        "effective_period": "2026-09", "account_profiles": profile,
        "obligation_profile_assignments": {"con_1": {"pob_1": "recurring"}},
    }, period="2026-09")
    wrong = [{"account": "4000", "dimensions": {"Department": "Other"}}]
    with pytest.raises(ValueError, match="Name the reviewed chart"):
        application.execute("set_policy", {"effective_period": "2026-09", "account_dimension_rules": wrong}, period="2026-09")
    with pytest.raises(ValueError, match="unique"):
        application.execute("set_policy", {"effective_period": "2026-09", "account_dimension_rules": wrong + wrong,
                                           "account_dimension_source": "Approved chart"}, period="2026-09")
    with pytest.raises(ValueError, match="unique after trimming"):
        application.execute("set_policy", {"effective_period": "2026-09", "account_dimension_rules": [
            {"account": "4000", "dimensions": {"Department": "Recurring", " Department ": "Other"}},
        ], "account_dimension_source": "Approved chart"}, period="2026-09")
    application.execute("set_policy", {"effective_period": "2026-09", "account_dimension_rules": wrong,
                                       "account_dimension_source": "Approved chart"}, period="2026-09")
    state = application.state(period="2026-09")
    exceptions = state["report"]["account_dimension_exceptions"]
    assert [(row["account"], row["dimensions"]) for row in exceptions] == [("4000", {"Department": "Recurring"})]
    assert state["report"]["account_dimension_unvalidated_accounts"] == ["1100", "2300"]
    book = load_workbook(io.BytesIO(export_bytes(state)), read_only=True)
    assert list(book["Approved combinations"].values)[1][2:] == ("4000", '{"Department":"Other"}')
    assert any(row[0] == "Invalid combination" and row[4] == "4000" for row in book["Combination preflight"].values)
    book.close()
    check = next(row for row in application.reports(period="2026-09")["checks"] if row["id"] == "account_dimensions")
    assert check["status"] == "block"
    assert check["target"] == {"view": "Journal entries"}
    with pytest.raises(ValueError, match="Resolve close blockers"):
        application.execute("close_period", {"period": "2026-09"}, period="2026-09")
    partial = [{"account": "4000", "dimensions": {"Department": "Recurring"}}]
    application.execute("set_policy", {"effective_period": "2026-09", "account_dimension_rules": partial,
                                       "account_dimension_source": "Approved chart"}, period="2026-09")
    assert next(row for row in application.reports(period="2026-09")["checks"] if row["id"] == "account_dimensions")["status"] == "review"
    complete = partial + [{"account": "1100", "dimensions": {}}, {"account": "2300", "dimensions": {}}]
    application.execute("set_policy", {"effective_period": "2026-09", "account_dimension_rules": complete,
                                       "account_dimension_source": "Approved chart"}, period="2026-09")
    report = application.state(period="2026-09")["report"]
    assert not report["account_dimension_exceptions"]
    assert not report["account_dimension_unvalidated_accounts"]
    assert next(row for row in application.reports(period="2026-09")["checks"] if row["id"] == "account_dimensions")["status"] == "pass"
    assert application.state(period="2026-10")["policy"]["account_dimension_rules"] == complete
    application.execute("close_period", {"period": "2026-09", "review_dispositions": accept_review_items(application, "2026-09")}, period="2026-09")
    with pytest.raises(ValueError, match="reopen"):
        application.execute("set_policy", {"effective_period": "2026-09", "account_dimension_rules": partial,
                                           "account_dimension_source": "Approved chart"}, period="2026-09")


def test_backdated_approved_combinations_flow_through_later_unrelated_policy(tmp_path):
    application = app(tmp_path)
    seed(application)
    application.execute("set_policy", {
        "effective_period": "2026-11", "accounts": {"billing_clearing": "1110"},
        "account_dimension_rules": [], "account_dimension_source": "",
    }, period="2026-11")
    rules = [{"account": "4000", "dimensions": {}}]
    application.execute("set_policy", {
        "effective_period": "2026-10", "account_dimension_rules": rules,
        "account_dimension_source": "Approved chart v2",
    }, period="2026-10")
    november = application.state(period="2026-11")
    assert november["policy"]["account_dimension_rules"] == rules
    assert november["policy"]["account_dimension_source"] == "Approved chart v2"
    assert november["policy"]["accounts"]["billing_clearing"] == "1110"


def test_export_batch_identity_changes_with_journal_not_descriptive_notes(tmp_path):
    application = app(tmp_path)
    seed(application)
    def batch_id():
        bundle = application.report_bundle(period="2026-09")
        book = load_workbook(io.BytesIO(export_bytes(bundle["state"], bundle["review"])), read_only=True)
        value = next(row[1] for row in book["Workspace"].values if row[0] == "Journal batch ID")
        book.close()
        return value
    first = batch_id()
    application.execute("add_note", {"entity_id": "con_1", "body": "Descriptive note"})
    assert batch_id() == first
    application.execute("record_billing", {"contract_id": "con_1", "effective_date": "2026-09-01", "amount": "100"}, period="2026-09")
    assert batch_id() != first


def test_external_posting_record_is_bound_to_closed_batch_and_exposes_stale_revisions(tmp_path):
    application = app(tmp_path)
    seed(application)
    with pytest.raises(ValueError, match="Close the period"):
        application.execute("record_export_posting", {"period": "2026-09", "batch_id": "wrong", "external_journal_reference": "ERP-1", "posted_date": "2026-10-01", "rationale": "Posted"}, period="2026-09")
    closed = application.execute("close_period", {"period": "2026-09", "review_dispositions": accept_review_items(application, "2026-09")}, period="2026-09")["state"]
    batch = closed["report"]["journal_batch_id"]
    with pytest.raises(ValueError, match="batch changed"):
        application.execute("record_export_posting", {"period": "2026-09", "batch_id": "wrong", "external_journal_reference": "ERP-1", "posted_date": "2026-10-01", "rationale": "Posted"}, period="2026-09")
    payload = {"period": "2026-09", "batch_id": batch, "external_journal_reference": "ERP-1", "posted_date": "2026-10-01", "rationale": "Posted to ERP and reconciled."}
    recorded = application.execute("record_export_posting", payload, period="2026-09")["state"]
    assert recorded["postings"][0]["batch_id"] == batch
    assert recorded["report"]["journal_batch_id"] == batch
    with pytest.raises(ValueError, match="already recorded"):
        application.execute("record_export_posting", payload, period="2026-09")
    application.execute("reopen_period", {"period": "2026-09", "rationale": "Correct late source invoice."}, period="2026-09")
    changed = application.execute("record_billing", {"contract_id": "con_1", "effective_date": "2026-09-30", "amount": "50.00"}, period="2026-09")["state"]
    assert changed["report"]["journal_batch_id"] != batch
    assert changed["postings"][0]["batch_id"] == batch
    comparison = changed["posting_comparisons"][0]
    assert comparison["source_batch_id"] == batch
    assert comparison["source_close_id"] == closed["report"]["close_id"]
    assert comparison["posting_references"] == ["ERP-1"]
    assert comparison["target_closed"] is False
    assert sum(Decimal(row["debit"]) - Decimal(row["credit"]) for row in comparison["lines"]) == 0
    assert any(row["account"] == "1100" and row["debit"] == "50.00" for row in comparison["lines"])
    application.execute("close_period", {"period": "2026-09", "review_dispositions": accept_review_items(application, "2026-09")}, period="2026-09")
    accepted = application.state(period="2026-09")
    assert accepted["posting_comparisons"][0]["target_closed"] is True
    book = load_workbook(io.BytesIO(export_bytes(accepted)), read_only=True)
    replacement = list(book["Replacement journal"].values)
    assert any(row[0] == batch and row[2] == "Accepted replacement" and row[5] == "1100" and row[10] == 50 for row in replacement)
    book.close()


def test_replacement_journal_preserves_account_and_dimension_changes():
    def line(account, debit, credit, department):
        return {"contract_id": "con_1", "account": account, "role": "deferred_revenue" if account != "1100" else "billing_clearing",
                "debit_minor": debit, "credit_minor": credit, "dimensions": {"Department": department}}
    posted = [line("1100", 10000, 0, "Shared"), line("2300", 0, 10000, "Legacy")]
    revised = [line("1100", 10000, 0, "Shared"), line("2310", 0, 10000, "Current")]
    comparison = compare_journals(posted, revised)
    assert [(row["account"], row["dimensions"], row["debit"], row["credit"]) for row in comparison] == [
        ("2300", {"Department": "Legacy"}, "100.00", "0.00"),
        ("2310", {"Department": "Current"}, "0.00", "100.00"),
    ]


def test_future_mapping_is_allowed_after_close_but_backdated_mapping_requires_reopen(tmp_path):
    application = app(tmp_path)
    seed(application)
    application.execute("close_period", {"period": "2026-09", "review_dispositions": accept_review_items(application, "2026-09")}, period="2026-09")
    application.execute("set_policy", {"effective_period": "2026-10", "accounts": {"revenue": "4100"}}, period="2026-10")
    assert application.state(period="2026-09")["report"]["closed"] is True
    with pytest.raises(ValueError, match="closed period"):
        application.execute("set_policy", {"effective_period": "2026-09", "accounts": {"revenue": "4200"}}, period="2026-09")
    assert application.state(period="2026-09")["policy"]["accounts"]["revenue"] == "4000"


def test_same_month_mapping_uses_last_recorded_policy_and_keeps_versions(tmp_path):
    application = app(tmp_path)
    seed(application)
    application.execute("set_policy", {"effective_period": "2026-10", "accounts": {"revenue": "4100"}}, period="2026-10")
    application.execute("set_policy", {"effective_period": "2026-10", "accounts": {"billing_clearing": "1110"}}, period="2026-10")
    state = application.state(period="2026-10")
    assert [row["version"] for row in state["policy_versions"]] == [1, 2, 3]
    assert state["policy"]["accounts"]["revenue"] == "4100"
    assert state["policy"]["accounts"]["billing_clearing"] == "1110"


def test_backdated_role_change_flows_through_later_policy_without_overwriting_its_role(tmp_path):
    application = app(tmp_path)
    seed(application)
    application.execute("set_policy", {"effective_period": "2026-10", "accounts": {"revenue": "4100"}}, period="2026-10")
    application.execute("set_policy", {"effective_period": "2026-09", "accounts": {"billing_clearing": "1110"}}, period="2026-09")
    application.execute("record_billing", {"contract_id": "con_1", "effective_date": "2026-10-01", "amount": "100.00"}, period="2026-10")
    september = application.state(period="2026-09")
    october = application.state(period="2026-10")
    assert september["policy"]["accounts"]["billing_clearing"] == "1110"
    assert september["policy"]["accounts"]["revenue"] == "4000"
    assert october["policy"]["accounts"]["billing_clearing"] == "1110"
    assert october["policy"]["accounts"]["revenue"] == "4100"
    assert {row["account"] for row in october["report"]["journals"] if row["role"] == "billing_clearing"} == {"1110"}


def test_change_detail_replays_the_exact_financial_transition(tmp_path):
    application = app(tmp_path)
    seed(application)
    change = application.execute("record_billing", {
        "contract_id": "con_1", "effective_date": "2026-09-01", "amount": "12000.00", "reference": "INV-1",
    }, period="2026-09")["result"]["change_set_id"]
    application.execute("record_billing", {
        "contract_id": "con_1", "effective_date": "2026-10-01", "amount": "100.00", "reference": "INV-2",
    }, period="2026-10")
    detail = application.change_detail(change, "2026-09")
    assert detail["change"]["entity_name"] == "Annual service"
    assert detail["before_report"]["summary"]["billings"] == "0.00"
    assert detail["after_report"]["summary"]["billings"] == "12000.00"
    assert detail["comparison"]["summary"]["billings"] == "12000.00"
    assert detail["after_report"]["journals"] != detail["before_report"]["journals"]


def test_policy_change_detail_shows_journal_mapping_even_without_revenue_delta(tmp_path):
    application = app(tmp_path)
    seed(application)
    change = application.execute("set_policy", {"effective_period": "2026-10", "accounts": {"revenue": "4100"}}, period="2026-10")["result"]["change_set_id"]
    detail = application.change_detail(change, "2026-10")
    assert detail["before_policy"]["version"] == 1
    assert detail["after_policy"]["version"] == 2
    assert detail["comparison"]["summary"]["revenue"] == "0.00"
    assert "2026-10" in detail["comparison"]["affected_periods"]


def test_judgment_review_is_separate_from_file_linkage_and_latest_review_wins(tmp_path):
    application = app(tmp_path)
    seed(application)
    first = application.execute("record_adjustment", {"contract_id": "con_1", "obligation_id": "pob_1", "effective_date": "2026-09-30", "amount": "10.00", "rationale": "Reviewed true-up"}, period="2026-09")["result"]["change_set_id"]
    second = application.execute("record_adjustment", {"contract_id": "con_1", "obligation_id": "pob_1", "effective_date": "2026-09-30", "amount": "5.00", "rationale": "Separate estimate"}, period="2026-09")["result"]["change_set_id"]
    folder = application.workspace.path / "attachments"
    (folder / "support.txt").write_text("Reviewed calculation")
    application.execute("attach_evidence", {"name": "support.txt", "path": "attachments/support.txt", "entity_id": "con_1"}, period="2026-09")
    assert {first, second}.issubset({row["change_set_id"] for row in application.reports(period="2026-09")["exceptions"]["evidence"]})
    application.execute("attach_evidence", {"name": "support.txt", "path": "attachments/support.txt", "entity_id": "con_1", "target_change_set_id": first, "obligation_id": "pob_1"}, period="2026-09")
    review = application.reports(period="2026-09")
    assert first in {row["change_set_id"] for row in review["exceptions"]["evidence"]}
    assert application.change_detail(first, "2026-09")["evidence"][0]["target_change_set_id"] == first
    payload = {"target_change_set_id": first, "reviewer": "Accountant", "conclusion": "True-up is appropriate", "support_memo": "Agreed calculation to signed contract"}
    application.execute("record_judgment_review", {**payload, "disposition": "supported"}, period="2026-09")
    assert first not in {row["change_set_id"] for row in application.reports(period="2026-09")["exceptions"]["evidence"]}
    assert application.change_detail(first, "2026-09")["judgment_reviews"][-1]["conclusion"] == "True-up is appropriate"
    application.execute("record_judgment_review", {**payload, "disposition": "exception", "exception_reason": "Missing approval"}, period="2026-09")
    assert next(row for row in application.reports(period="2026-09")["exceptions"]["evidence"] if row["change_set_id"] == first)["review_status"] == "exception"
    application.execute("record_judgment_review", {**payload, "target_change_set_id": second, "disposition": "supported"}, period="2026-09")
    assert second not in {row["change_set_id"] for row in application.reports(period="2026-09")["exceptions"]["evidence"]}
    with pytest.raises(ValueError, match="Judgment change set"):
        application.execute("record_judgment_review", {**payload, "target_change_set_id": "missing", "disposition": "supported"}, period="2026-09")


def test_change_detail_covers_modification_and_scenario_application(tmp_path):
    application = app(tmp_path)
    seed(application)
    scenario = application.execute("create_scenario", {"name": "Expanded scope"})["result"]["id"]
    modification = application.execute("modify_contract", {
        "contract_id": "con_1", "effective_date": "2026-10-01", "treatment": "prospective",
        "consideration": [{"id": "price_1", "label": "Fixed fee", "kind": "fixed", "amount": "18000.00"}],
        "rationale": "Expanded service",
    }, scenario_id=scenario, period="2026-10")["result"]["change_set_id"]
    proposed = application.change_detail(modification, "2026-10")
    assert proposed["comparison"]["summary"]["transaction_price"] == "6000.00"
    applied = application.execute("apply_scenario", {"scenario_id": scenario}, scenario_id=scenario, period="2026-10")["result"]["change_set_id"]
    accepted = application.change_detail(applied, "2026-10")
    assert accepted["before_report"]["summary"]["transaction_price"] == "12000.00"
    assert accepted["after_report"]["summary"]["transaction_price"] == "18000.00"
    assert accepted["comparison"]["summary"]["transaction_price"] == "6000.00"


def test_scenario_evidence_follows_the_accepted_judgment_without_copying_the_file(tmp_path):
    application = app(tmp_path)
    seed(application)
    scenario = application.execute("create_scenario", {"name": "Expanded scope"})["result"]["id"]
    proposal = application.execute("modify_contract", {
        "contract_id": "con_1", "effective_date": "2026-10-01", "treatment": "prospective",
        "consideration": [{"id": "price_1", "label": "Fixed fee", "kind": "fixed", "amount": "18000.00"}],
        "rationale": "Expanded service",
    }, scenario_id=scenario, period="2026-10")["result"]["change_set_id"]
    (application.workspace.path / "attachments" / "amendment.txt").write_text("Signed amendment")
    application.execute("attach_evidence", {"name": "amendment.txt", "path": "attachments/amendment.txt", "entity_id": "con_1", "target_change_set_id": proposal}, scenario_id=scenario, period="2026-10")
    application.execute("record_judgment_review", {"target_change_set_id": proposal, "reviewer": "Accountant", "disposition": "supported", "conclusion": "Prospective treatment applies", "support_memo": "Reviewed amended scope and price"}, scenario_id=scenario, period="2026-10")
    application.execute("apply_scenario", {"scenario_id": scenario}, scenario_id=scenario, period="2026-10")
    accepted = next(row for row in application.state(period="2026-10")["change_sets"] if row.get("originating_change_set_id") == proposal)
    detail = application.change_detail(accepted["id"], "2026-10")
    assert detail["evidence"][0]["target_change_set_id"] == accepted["id"]
    assert detail["evidence"][0]["path"] == "attachments/amendment.txt"
    assert accepted["id"] not in {row["change_set_id"] for row in application.reports(period="2026-10")["exceptions"]["evidence"]}
    assert next(row for row in application.state(period="2026-10")["judgment_reviews"] if row["target_change_set_id"] == accepted["id"])["conclusion"] == "Prospective treatment applies"


def test_close_export_indexes_change_support_and_accepted_policy(tmp_path):
    application = app(tmp_path)
    seed(application)
    policy_change = application.execute("set_policy", {"effective_period": "2026-09", "accounts": {"revenue": "4100"}, "rationale": "Reviewed mapping"}, period="2026-09")["result"]["change_set_id"]
    application.execute("close_period", {"period": "2026-09", "rationale": "Reviewed close", "review_dispositions": accept_review_items(application, "2026-09")}, period="2026-09")
    (application.workspace.path / "attachments" / "mapping.txt").write_text("Account memo")
    application.execute("attach_evidence", {"name": "mapping.txt", "path": "attachments/mapping.txt", "target_change_set_id": policy_change, "rationale": "Mapping memo"}, period="2026-09")
    application.execute("record_judgment_review", {"target_change_set_id": policy_change, "reviewer": "Accountant", "disposition": "supported", "conclusion": "4100 is the appropriate revenue account", "support_memo": "Reviewed mapping memo and GL chart"}, period="2026-09")
    bundle = application.report_bundle(period="2026-09")
    book = load_workbook(io.BytesIO(export_bytes(bundle["state"], bundle["review"])), read_only=True)
    assert book["Workspace"]["B13"].value is True
    assert book["Account policy"]["B2"].value == "4100" or any(row[1] == "4100" for row in book["Account policy"].values)
    support = list(book["Judgment support"].values)
    assert any(row[0] == policy_change and row[6] == "supported" and row[7] == "Accountant" and "mapping.txt" in row[14] for row in support)
    assert any(row[4] == policy_change and row[8] == "attachments/mapping.txt" for row in book["Evidence index"].values)
    assert any(row[8] == 2 for row in book["Journal entries"].values if isinstance(row[8], int))
    book.close()


def test_legacy_policy_command_uses_its_recorded_effective_month(tmp_path):
    application = app(tmp_path)
    seed(application)
    with application.workspace.connect() as db:
        application._append(db, "set_policy", {"effective_date": "2026-10-01", "accounts": {"revenue": "4100"}}, "main", "legacy-test")
    assert application.state(period="2026-09")["policy"]["accounts"]["revenue"] == "4000"
    assert application.state(period="2026-10")["policy"]["accounts"]["revenue"] == "4100"


def test_historical_scenario_detail_keeps_its_original_main_base_after_rebase(tmp_path):
    application = app(tmp_path)
    seed(application)
    scenario = application.execute("create_scenario", {"name": "Billing proposal"})["result"]["id"]
    change = application.execute("record_billing", {"contract_id": "con_1", "effective_date": "2026-10-01", "amount": "100.00"}, scenario_id=scenario, period="2026-10")["result"]["change_set_id"]
    application.execute("create_customer", {"id": "cus_later", "name": "Later Main customer"})
    application.execute("rebase_scenario", {"scenario_id": scenario})
    detail = application.change_detail(change, "2026-10")
    assert {row["id"] for row in detail["before_state"]["customers"]} == {"cus_1"}
    assert detail["comparison"]["summary"]["billings"] == "100.00"
    assert {row["id"] for row in application.state(scenario, "2026-10")["customers"]} == {"cus_1", "cus_later"}


def test_reassessment_and_close_change_details_rebuild_their_recorded_transition(tmp_path):
    application = app(tmp_path)
    load_demo(application)
    reassessment = next(row for row in application.state(period="2026-09")["change_sets"] if row["command"] == "reassess_variable_consideration")
    detail = application.change_detail(reassessment["id"], "2026-09")
    assert detail["comparison"]["summary"]["transaction_price"] == "1000.00"
    assert detail["comparison"]["summary"]["revenue"] == "300.00"
    close = application.execute("close_period", {"period": "2026-09", "rationale": "Reviewed", "review_dispositions": accept_review_items(application, "2026-09")}, period="2026-09")["result"]["change_set_id"]
    closed_detail = application.change_detail(close, "2026-09")
    assert closed_detail["before_report"].get("closed") is not True
    assert closed_detail["after_report"]["closed"] is True
