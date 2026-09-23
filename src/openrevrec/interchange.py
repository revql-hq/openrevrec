"""Spreadsheet interchange is an adapter over the same accounting commands."""

from __future__ import annotations

import io
import hashlib
import json
import zipfile
from datetime import date, datetime
from decimal import Decimal

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment

from .workspace import dumps, now
from .application import JUDGMENT_COMMANDS

TEMPLATE_VERSION = "2"
SHEETS = {
    "Customers": ["id", "name", "email", "reference", "source_system"],
    "Contracts": ["id", "customer_id", "name", "start_date", "end_date", "rationale", "cutover_date", "term_basis", "term_assessment_rationale", "term_reassessment_trigger", "term_review_date", "reference", "combination_basis", "combination_rationale"],
    "Contract Sources": ["contract_id", "reference", "agreement_date"],
    "Consideration": ["contract_id", "id", "label", "kind", "amount", "included_amount", "potential_amount", "estimated_amount", "estimation_method", "rationale", "allocation_scope", "target_obligation_ids", "allocation_rationale", "target_period", "unit_rate", "pricing_basis", "rounding_period", "metered_value_mode"],
    "Obligations": ["contract_id", "id", "name", "kind", "ssp", "method", "start_date", "end_date", "total_units", "exercise_start", "exercise_end", "rationale"],
    "Opening Positions": ["source_id", "contract_id", "effective_date", "billed_to_date", "contract_asset", "deferred_revenue", "source_name", "rationale"],
    "Opening Obligations": ["opening_source_id", "obligation_id", "recognized_to_date", "measure"],
    "Amendments": ["source_id", "contract_id", "effective_date", "treatment", "replace_consideration", "replace_obligations", "rationale", "term_basis", "term_assessment_rationale", "term_reassessment_trigger", "term_review_date"],
    "Amendment Consideration": ["amendment_source_id", "id", "label", "kind", "amount", "included_amount", "potential_amount", "estimated_amount", "estimation_method", "rationale", "allocation_scope", "target_obligation_ids", "allocation_rationale", "target_period", "unit_rate", "pricing_basis", "rounding_period", "metered_value_mode"],
    "Amendment Obligations": ["amendment_source_id", "id", "name", "kind", "ssp", "method", "start_date", "end_date", "total_units", "exercise_start", "exercise_end", "rationale"],
    "Mixed Allocations": ["amendment_source_id", "obligation_id", "treatment", "amount", "revised_progress"],
    "Billing": ["contract_id", "effective_date", "amount", "reference", "applies_to_change_set_id", "applies_to_reference", "rationale", "source_id", "source_contract_reference"],
    "Rate Changes": ["contract_id", "component_id", "effective_date", "unit_rate", "rationale", "source_id"],
    "Progress": ["contract_id", "obligation_id", "effective_date", "percentage", "rationale", "source_id"],
    "Usage": ["contract_id", "obligation_id", "effective_date", "quantity", "reference", "rationale", "source_id", "invoice_value", "source_contract_reference"],
    "Right Exercises": ["contract_id", "obligation_id", "effective_date", "delivery_method", "delivery_start", "delivery_end", "rationale", "source_id"],
    "Renewal Links": ["contract_id", "obligation_id", "renewal_contract_id", "additional_consideration", "price_basis", "rationale", "source_id"],
    "Modification Links": ["contract_id", "added_contract_id", "effective_date", "additional_consideration", "price_basis", "original_terms_effect", "rationale", "source_id"],
    "Milestones": ["contract_id", "obligation_id", "effective_date", "percentage", "rationale", "source_id"],
    "Adjustments": ["contract_id", "obligation_id", "effective_date", "amount", "rationale", "source_id"],
    "Reassessments": ["contract_id", "component_id", "effective_date", "included_amount", "rationale", "source_id"],
    "Corrections": ["source_id", "target_change_set_id", "target_source_id", "activity_type", "contract_id", "obligation_id", "effective_date", "amount", "percentage", "quantity", "reference", "rationale", "component_id", "unit_rate", "invoice_value", "source_contract_reference"],
    "Notes": ["entity_id", "kind", "body", "due_date", "source_id"],
    "Commands": ["command", "payload_json", "source_id"],
}
COMMAND_BY_SHEET = {"Customers": "create_customer", "Billing": "record_billing", "Rate Changes": "record_rate_change", "Progress": "record_progress", "Usage": "record_usage", "Right Exercises": "record_right_exercise", "Renewal Links": "link_renewal_contract", "Modification Links": "link_modification_contract", "Milestones": "record_milestone", "Adjustments": "record_adjustment", "Reassessments": "reassess_variable_consideration", "Notes": "add_note"}


def safe_cell(value):
    if isinstance(value, (dict, list)):
        value = dumps(value)
    if isinstance(value, str) and value.startswith(("=", "+", "-", "@")):
        return "'" + value
    return value


def _sheet(book, name, headers, rows=()):
    sheet = book.create_sheet(name)
    sheet.append(headers)
    for row in rows:
        sheet.append([safe_cell(v) for v in row])
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for cell in sheet[1]:
        cell.font = Font(name="Aptos", bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="303532")
        cell.alignment = Alignment(vertical="center")
    sheet.row_dimensions[1].height = 24
    for column in sheet.columns:
        letter = column[0].column_letter
        length = max((len(str(c.value or "")) for c in column), default=10)
        sheet.column_dimensions[letter].width = min(55, max(16, length + 2))
    return sheet


def template_bytes():
    book = Workbook()
    book.remove(book.active)
    _sheet(book, "Instructions", ["OpenRevRec import template", "Value"], [
        ["Template version", TEMPLATE_VERSION], ["Currency", "Uses workspace currency; decimal amounts, no currency symbols."],
        ["Dates", "YYYY-MM-DD. Service start and end dates are inclusive."],
        ["Relationships", "Supply IDs for customers, contracts, consideration, and obligations; reference these IDs on related sheets."],
        ["Contract setup", "Contracts, Consideration and Obligations are combined into one create_contract command."],
        ["Combined contracts", "When two or more source agreements meet a reviewed AASB 15 paragraph 17 criterion, represent them as one accounting contract. Enter the primary agreement reference in Contracts.reference, choose combination_basis (package, interdependent_price, or single_obligation), and explain the conclusion. Put every source agreement, including the primary, on Contract Sources with its agreement date. Billing and Usage rows for that accounting contract must identify the source_contract_reference. The software does not infer the combination conclusion or apply a fixed near-time threshold."],
        ["Source population", "Set reference on each contract to its stable register ID. Record independent contract, invoice/credit, and externally priced usage lists for each close period in Reports > Source population. For priced usage, supply each source record ID, delivered units, and invoice value so both population and amounts can be compared. Older ID-only lists remain review items. An imported billing source_id can identify a billing transaction when invoice reference is blank."],
        ["Cancellable or evergreen terms", "The contract end is the assessed accounting end, not an unlimited legal end. Enter term_basis as cancellable or evergreen, explain the assessment and reassessment trigger, and optionally enter term_review_date. Revisit the obligation dates through a reviewed amendment when the assessment changes."],
        ["Recognition methods", "exact_days, monthly, prorated_monthly, point_in_time, progress, usage, metered, milestone"],
        ["Opening positions", "For a migrated contract, supply current terms, then one Opening Positions row dated the first day of the cutover month and one Opening Obligations row per obligation. A matching opening row sets the new contract's cutover date automatically. Enter cumulative legacy recognition, billing, net asset/deferred balance, and measures before post-cutover activity. Pre-cutover periods are excluded."],
        ["Progress", "Cumulative percentage 0–100. Usage quantity is incremental; set total_units on finite usage obligations. Metered obligations have no unit cap."],
        ["Right exercises", "For a material right exercised before delivery, enter the exercise date, future delivery method and delivery dates on Right Exercises. A point-in-time delivery also needs a 100% Milestones row on its delivery date. Unexercised rights recognize at expiry."],
        ["Renewal links", "After entering both contracts and the right exercise, use Renewal Links to pair them. The renewal contract must cover the same delivery dates and contain only the NEW consideration; enter its initial transaction price in additional_consideration and new_consideration_only in price_basis. The old right allocation remains on the original contract. Review offsetting balances before posting."],
        ["Separate-contract amendments", "After entering the original and added-service contracts, use Modification Links only when the added services are distinct, the amendment's price increase reflects their standalone selling prices, and the original remaining promises and price are unchanged. Enter distinct_at_standalone_price in price_basis, unchanged in original_terms_effect, and the added contract's initial transaction price in additional_consideration. A combined change to the original service needs a separate mixed-treatment analysis; do not force it into this link."],
        ["Consideration kinds", "fixed, variable, usage, metered, credit. Use included_amount for constrained consideration. Metered requires amount 0, pricing_basis right_to_invoice, rounding_period calendar_month, and an accountant rationale. Set metered_value_mode to unit_rate with a positive unit_rate, or invoice_value with no unit_rate. Invoice-value mode requires a priced source amount and reference on each usage row; it does not calculate tiers or permit rate changes. One service obligation; no opening position or other amendment."],
        ["Rate changes", "Use Rate Changes for a new positive unit_rate effective before that day's usage. Enter the metered component ID and explain why invoice value at the revised rate still corresponds to delivered value. One rate change per effective date; each calendar month rounds once after valuing its usage at the applicable dates' rates."],
        ["Specific allocation", "For an eligible variable, usage, or credit component, set allocation_scope to specific, list target_obligation_ids separated by commas, and record allocation_rationale. For a variable or usage amount targeting one month of one time-based series obligation, also enter target_period as YYYY-MM and only one obligation ID. Otherwise leave these fields blank for relative SSP allocation."],
        ["Amendments", "Each amendment needs a stable source_id, effective date, treatment, and rationale. Set replace_consideration and/or replace_obligations to yes. Related rows must list the COMPLETE replacement set for that section. For mixed treatment, add one reviewed lifetime amount for every revised obligation in Mixed Allocations; the amounts must sum to revised consideration. If an integrated progress-method service has a revised completion measure, enter revised_progress as a percentage on its catch-up row."],
        ["Corrections", "Supply the original change set ID or its prior source_id, the activity type, replacement facts, and a new source_id. The original remains in history."],
        ["Other commands", "Use Commands with command, payload_json, and source_id for policies or advanced actions. Close/reopen and scenario lifecycle require separate review."],
        ["Import behavior", "Preview first. All rows then commit together; a failed import commits no accounting changes."],
        ["Source identity", "Use a stable, system-prefixed source_id for every activity, amendment, correction, or command. A billing invoice reference can substitute for source_id. Reimporting the same source fact is rejected."],
        ["Review", "Always review totals and accounting judgments after import. Amount cells may be numbers or decimal text."],
    ])
    for name, headers in SHEETS.items():
        _sheet(book, name, headers)
    output = io.BytesIO()
    book.save(output)
    return output.getvalue()


def journal_batch_id(state):
    report = state["report"]
    return "orr-" + hashlib.sha256(dumps({"workspace_id": state["workspace"]["id"], "scenario_id": state["scenario_id"], "period": report["period"], "policy_version": report.get("policy_version"), "journals": report["journals"]}).encode()).hexdigest()[:20]


def export_bytes(state, review=None):
    report = state["report"]
    batch_id = journal_batch_id(state)
    book = Workbook()
    book.remove(book.active)
    _sheet(book, "Workspace", ["Field", "Value"], [["Company", state["workspace"]["name"]], ["Workspace ID", state["workspace"]["id"]], ["Currency", state["policy"]["currency"]], ["Period", report["period"]], ["Scenario", state["scenario_id"]], ["Version", state["frontier"]], ["Policy version", report.get("policy_version", state["policy"]["version"])], ["Policy effective period", report.get("policy_effective_period", state["policy"]["effective_period"])], ["Posting preflight coverage", report.get("policy_account_dimension_coverage", "listed")], ["Ruleset", "orr-0.1"], ["Engine", "0.1.0"], ["Exported at", now()], ["Accepted close", bool(report.get("closed"))], ["Support metadata", "Evidence recorded after close is shown with its own recorded date; financial values use the accepted checkpoint." if report.get("closed") else "Current accepted workspace state."], ["Output status", "Hypothetical scenario support" if state["scenario_id"] != "main" else "Accepted Main close" if report.get("closed") else "Open Main support; subject to change"], ["Journal batch ID", batch_id]])
    _sheet(book, "Summary", ["Measure", "Amount"], [[k.replace("_", " ").capitalize(), Decimal(v)] for k, v in report["summary"].items() if isinstance(v, (str, int))])
    if review:
        _sheet(book, "Close readiness", ["Check", "Status", "Detail", "Exceptions"], [[row["label"], row["status"], row["detail"], row["count"]] for row in review["checks"]])
        _sheet(book, "External controls", ["Measure", "Model", "External", "Difference", "Source", "Review explanation", "Close cutoff date"],
               [[field, Decimal(values["derived"]), Decimal(values["external"]), Decimal(values["difference"]), review["external_control"]["source_name"], review["external_control"]["rationale"], review["external_control"].get("close_cutoff_date", "")]
                for field, values in review.get("external_control_comparison", {}).items()])
        population = review.get("population_manifest")
        comparison = review.get("population_comparison", {})
        if population:
            missing_contracts = set(comparison["missing_contracts"])
            missing_billings = {tuple(item) for item in comparison["missing_billings"]}
            duplicate_contracts = set(comparison["duplicate_contracts"])
            _sheet(book, "Source contracts", ["Contract reference", "Comparison", "Source", "Population basis"],
                   [[reference, "Missing in workspace" if reference in missing_contracts else "Duplicate in workspace" if reference in duplicate_contracts else "Matched", population["source_name"], population["rationale"]]
                    for reference in population["contract_references"]] +
                   [[reference, "Unexpected in workspace", population["source_name"], population["rationale"]] for reference in comparison["unexpected_contracts"]] +
                   [[item["id"], "Workspace contract lacks reference", population["source_name"], population["rationale"]] for item in comparison["unidentified_contracts"]])
            billing_rows = comparison.get("billing_value_rows")
            if billing_rows is None:  # Older closed checkpoints compared invoice identities only.
                duplicate_billings = {tuple(item) for item in comparison["duplicate_billings"]}
                billing_rows = [{"contract_reference": item["contract_reference"], "invoice_reference": item["invoice_reference"],
                                 "status": "Missing in workspace" if (item["contract_reference"], item["invoice_reference"]) in missing_billings else "Duplicate in workspace" if (item["contract_reference"], item["invoice_reference"]) in duplicate_billings else "Identity matched; amount not checked"}
                                for item in population["billing_references"]]
            _sheet(book, "Source invoices", ["Contract reference", "Invoice reference", "Comparison", "Source", "Population basis", "Source amount", "Workspace amount", "Activity ID"],
                   [[item["contract_reference"], item["invoice_reference"], item["status"], population["source_name"], population["rationale"],
                     Decimal(item["source_amount"]) if item.get("source_amount") not in (None, "") else "",
                     Decimal(item["workspace_amount"]) if item.get("workspace_amount") not in (None, "") else "", item.get("activity_id", "")]
                    for item in billing_rows] +
                   [[contract_ref, invoice_ref, "Unexpected in workspace", population["source_name"], population["rationale"], "", "", ""] for contract_ref, invoice_ref in comparison["unexpected_billings"]] +
                   [[item["contract_id"], item["activity_id"], "Workspace billing lacks reference", population["source_name"], population["rationale"], "", "", item["activity_id"]] for item in comparison["unidentified_billings"]])
            usage_rows = comparison.get("usage_value_rows")
            if usage_rows is None:  # Closed checkpoints from before value comparison retain their original identity result.
                missing_usage = {tuple(item) for item in comparison.get("missing_usage", [])}
                duplicate_usage = {tuple(item) for item in comparison.get("duplicate_usage", [])}
                usage_rows = [{"contract_reference": item["contract_reference"], "usage_reference": item["usage_reference"],
                               "status": "Missing in workspace" if (item["contract_reference"], item["usage_reference"]) in missing_usage else "Duplicate in workspace" if (item["contract_reference"], item["usage_reference"]) in duplicate_usage else "Identity matched; values not checked"}
                              for item in population.get("usage_references", [])]
            _sheet(book, "Source priced usage", ["Contract reference", "Usage source reference", "Source units", "Workspace units", "Source invoice value", "Workspace invoice value", "Comparison", "Source", "Population basis"],
                   [[item["contract_reference"], item["usage_reference"], Decimal(item["source_quantity"]) if item.get("source_quantity") not in (None, "") else "", Decimal(item["workspace_quantity"]) if item.get("workspace_quantity") not in (None, "") else "", Decimal(item["source_invoice_value"]) if item.get("source_invoice_value") not in (None, "") else "", Decimal(item["workspace_invoice_value"]) if item.get("workspace_invoice_value") not in (None, "") else "", item["status"], population["source_name"], population["rationale"]]
                    for item in usage_rows] +
                   [[contract_ref, usage_ref, "", "", "", "", "Unexpected in workspace", population["source_name"], population["rationale"]] for contract_ref, usage_ref in comparison.get("unexpected_usage", [])] +
                   [[item["contract_id"], item["activity_id"], "", "", "", "", "Workspace priced usage lacks reference", population["source_name"], population["rationale"]] for item in comparison.get("unidentified_usage", [])])
            opening_fields = ("billed_to_date", "contract_asset", "deferred_revenue", "recognized_to_date")
            _sheet(book, "Source openings", ["Contract reference", "Cutover date", "Comparison", "Source", "Population basis",
                                             *[f"Source {field.replace('_', ' ')}" for field in opening_fields],
                                             *[f"Workspace {field.replace('_', ' ')}" for field in opening_fields], "Activity ID"],
                   [[item["contract_reference"], item["cutover_date"], item["status"], population["source_name"], population["rationale"],
                     *[Decimal(item["source_values"][field]) for field in opening_fields],
                     *[Decimal(item["workspace_values"][field]) if item["workspace_values"][field] != "" else "" for field in opening_fields], item["activity_id"]]
                    for item in comparison.get("opening_value_rows", [])] +
                   [[reference, cutover, "Unexpected in workspace", population["source_name"], population["rationale"], *([""] * 8), ""]
                    for reference, cutover in comparison.get("unexpected_openings", [])] +
                   [[reference, cutover, "Source opening balances not supplied", population["source_name"], population["rationale"], *([""] * 8), ""]
                    for reference, cutover in comparison.get("unverified_openings", [])] +
                   [[item["contract_id"], "", "Workspace opening lacks contract reference", population["source_name"], population["rationale"], *([""] * 8), item["activity_id"]]
                    for item in comparison.get("unidentified_openings", [])])
        money_keys = {"opening_contract_asset", "opening_deferred_revenue", "revenue", "billings", "closing_contract_asset", "closing_deferred_revenue"}
        rollforward_keys = ["contract_id", "contract_name", "opening_contract_asset", "opening_deferred_revenue", "revenue", "billings", "closing_contract_asset", "closing_deferred_revenue"]
        _sheet(book, "Contract rollforward", rollforward_keys, [[Decimal(row[k]) if k in money_keys else row[k] for k in rollforward_keys] for row in review["rollforward"]])
        billing_keys = ["contract_id", "contract_name", "revenue", "billings", "net_movement", "closing_contract_asset", "closing_deferred_revenue"]
        _sheet(book, "Billing vs revenue", billing_keys, [[Decimal(row[k]) if k not in {"contract_id", "contract_name"} else row[k] for k in billing_keys] for row in review["billing_vs_revenue"]])
        coverage_keys = ["contract_id", "contract_name", "remaining_revenue", "future_scheduled", "unscheduled", "status"]
        _sheet(book, "Recognition coverage", coverage_keys, [[Decimal(row[k]) if k in {"remaining_revenue", "future_scheduled", "unscheduled"} else row[k] for k in coverage_keys] for row in review["recognition_coverage"]])
        scenario_keys = ["scenario_id", "name", "base_version", "main_version", "behind", "period_revenue_delta", "transaction_price_delta", "remaining_revenue_delta", "affected_periods"]
        _sheet(book, "Scenario impact", scenario_keys, [[Decimal(row[k]) if k.endswith("_delta") else ", ".join(row[k]) if k == "affected_periods" else row[k] for k in scenario_keys] for row in review["scenario_impacts"]])
    contract_keys = ["id", "name", "customer_id", "transaction_price", "revenue", "recognized_to_date", "billings", "billed_to_date", "deferred_revenue", "contract_asset", "remaining_revenue"]
    sheet = _sheet(book, "Contract balances", contract_keys, [[c.get(k, "") if k in contract_keys[:3] else Decimal(c.get(k, "0")) for k in contract_keys] for c in report["contracts"]])
    for row in sheet.iter_rows(min_row=2, min_col=4):
        for cell in row:
            cell.number_format = '#,##0.00;[Red](#,##0.00);–'
    term_fields = ("term_basis", "term_assessment_rationale", "term_reassessment_trigger", "term_review_date")
    term_rows = []
    for contract in state["contracts"]:
        assessment = {field: contract.get(field, "") for field in term_fields}
        assessment["term_basis"] = assessment["term_basis"] or "fixed"
        assessed_end = contract["end_date"]
        term_rows.append([contract["id"], contract["start_date"], assessment["term_basis"], assessed_end, assessment["term_assessment_rationale"], assessment["term_reassessment_trigger"], assessment["term_review_date"], contract.get("change_set_id", "")])
        for activity in sorted(contract["activities"], key=lambda row: row["effective_date"]):
            if activity["type"] != "modification":
                continue
            if "obligations" in activity:
                assessed_end = max((item["end_date"] for item in activity["obligations"]), default="")
            if not any(field in activity for field in term_fields):
                continue
            if activity.get("term_basis") == "fixed":
                assessment = {field: "" for field in term_fields}
            assessment.update({field: activity[field] for field in term_fields if field in activity})
            term_rows.append([contract["id"], activity["effective_date"], assessment["term_basis"], assessed_end, assessment["term_assessment_rationale"], assessment["term_reassessment_trigger"], assessment["term_review_date"], activity["id"]])
    _sheet(book, "Term assessments", ["Contract ID", "Effective date", "Term basis", "Assessed service end", "Assessment rationale", "Reassessment trigger", "Planned review date", "Change set ID"], term_rows)
    _sheet(book, "Term reviews", ["Contract ID", "Review date", "Reviewer", "Unchanged-term conclusion", "Supporting basis", "Next review date", "Change set ID"],
           [[item["contract_id"], item["effective_date"], item["reviewer"], item["conclusion"], item["support_memo"], item["next_review_date"], item["change_set_id"]] for item in state.get("term_reviews", [])])
    openings = [change for change in state["change_sets"] if change["command"] == "record_opening_position"]
    _sheet(book, "Opening positions", ["Contract ID", "Cutover date", "Legacy billed to date", "Legacy contract asset", "Legacy deferred revenue", "Legacy source", "Reconciliation rationale", "Change set ID"],
           [[item["payload"].get(key, "") for key in ("contract_id", "effective_date", "billed_to_date", "contract_asset", "deferred_revenue", "source_name", "rationale")] + [item["id"]] for item in openings])
    _sheet(book, "Opening obligation balances", ["Contract ID", "Obligation ID", "Recognized to date", "Cumulative measure", "Cutover date", "Change set ID"],
           [[item["payload"]["contract_id"], row["obligation_id"], row["recognized_to_date"], row.get("measure", ""), item["payload"]["effective_date"], item["id"]]
            for item in openings for row in item["payload"]["opening_obligations"]])
    _sheet(book, "Right exercises", ["Contract ID", "Obligation ID", "Exercise date", "Delivery method", "Delivery start", "Delivery end", "Rationale", "Change set ID"],
           [[item["payload"].get(key, "") for key in ("contract_id", "obligation_id", "effective_date", "delivery_method", "delivery_start", "delivery_end", "rationale")] + [item["id"]]
            for item in reversed(state["change_sets"]) if item["command"] == "record_right_exercise"])
    renewal_keys = ("contract_id", "obligation_id", "renewal_contract_id", "exercise_date", "delivery_start", "delivery_end", "original_right_allocation", "initial_new_consideration", "current_renewal_price", "combined_consideration", "right_revenue", "renewal_revenue", "combined_revenue", "rationale", "recorded_at", "change_set_id")
    renewal_money = {"original_right_allocation", "initial_new_consideration", "current_renewal_price", "combined_consideration", "right_revenue", "renewal_revenue", "combined_revenue"}
    _sheet(book, "Renewal links", list(renewal_keys), [[Decimal(item[key]) if key in renewal_money else item.get(key, "") for key in renewal_keys] for item in report.get("renewal_links", [])])
    modification_keys = ("contract_id", "contract_name", "added_contract_id", "effective_date", "added_contract_name", "price_basis", "original_terms_effect", "initial_additional_consideration", "current_added_price", "original_revenue", "added_revenue", "combined_revenue", "original_contract_asset", "original_deferred_revenue", "added_contract_asset", "added_deferred_revenue", "rationale", "recorded_at", "change_set_id")
    modification_money = {"initial_additional_consideration", "current_added_price", "original_revenue", "added_revenue", "combined_revenue", "original_contract_asset", "original_deferred_revenue", "added_contract_asset", "added_deferred_revenue"}
    _sheet(book, "Modification links", list(modification_keys), [[Decimal(item[key]) if key in modification_money else item.get(key, "") for key in modification_keys] for item in report.get("modification_links", [])])
    _sheet(book, "Contract combinations", ["Accounting contract ID", "Source agreement reference", "Agreement date", "Paragraph 17 basis", "Accounting rationale", "Change set ID"],
           [[contract["id"], source["reference"], source["agreement_date"], contract["combination_basis"], contract["combination_rationale"], contract.get("change_set_id", "")]
            for contract in state["contracts"] for source in contract.get("source_contracts", [])])
    _sheet(book, "Mixed modifications", ["Contract ID", "Effective date", "Obligation ID", "Treatment", "Revised lifetime allocation", "Revised completion (%)", "Rationale", "Change set ID"],
           [[change["payload"]["contract_id"], change["payload"]["effective_date"], row["obligation_id"], row["treatment"], Decimal(row["amount"]), Decimal(row["revised_progress"]) if "revised_progress" in row else "", change["payload"]["rationale"], change["id"]]
            for change in reversed(state["change_sets"]) if change["command"] == "modify_contract" and change["payload"].get("treatment") == "mixed"
            for row in change["payload"]["mixed_allocation"]])
    _sheet(book, "Recognition schedule", ["period", "contract_id", "obligation_id", "revenue"], [[r["period"], r["contract_id"], r["obligation_id"], Decimal(r["revenue"])] for r in report["schedule"]])
    dimension_keys = sorted({key for row in report["journals"] for key in row.get("dimensions", {})})
    journal_keys = ["period", "contract_id", "account", "account_name", "role", "debit", "credit", "description", "policy_version", "policy_effective_period", "obligation_ids", "batch_id", "account_profile_id"] + [f"dimension:{key}" for key in dimension_keys]
    def journal_cell(row, key):
        if key in {"debit", "credit"}:
            return Decimal(row.get(key, "0"))
        if key in {"policy_version", "policy_effective_period"}:
            return report.get(key, "")
        if key == "obligation_ids":
            return ", ".join(row.get(key, []))
        if key == "batch_id":
            return batch_id
        if key.startswith("dimension:"):
            return row.get("dimensions", {}).get(key.removeprefix("dimension:"), "")
        return row.get(key, "")
    _sheet(book, "Journal entries", journal_keys, [[journal_cell(row, key) for key in journal_keys] for row in report["journals"]])
    _sheet(book, "Dimension balance", ["Contract ID", "Dimensions", "Net debit / (credit)"],
           [[row["contract_id"], row["dimensions"], Decimal(row["net_debit"])] for row in report.get("segment_imbalances", [])])
    approved = report.get("policy_account_dimension_rules", [])
    chart_source = report.get("policy_account_dimension_source", "")
    _sheet(book, "Approved combinations", ["Chart source", "Effective period", "Account", "Dimensions"],
           [[chart_source, report.get("policy_effective_period", ""), item["account"], item["dimensions"]] for item in approved])
    preflight = [["Invalid combination", item["journal_id"], item["contract_id"], item["role"], item["account"], item["dimensions"]]
                 for item in report.get("account_dimension_exceptions", [])]
    uncovered_status = "Unlisted account blocks close" if report.get("policy_account_dimension_coverage") == "complete" else "Account not covered; review"
    preflight += [[uncovered_status, "", "", "", account, ""] for account in report.get("account_dimension_unvalidated_accounts", [])]
    if not approved:
        preflight.append(["No approved list configured", "", "", "", "", ""])
    _sheet(book, "Combination preflight", ["Status", "Journal ID", "Contract ID", "Role", "Account", "Dimensions"], preflight)
    posting_steps = [["1", "Confirm how your billing system posts invoices before using journal entries."], ["2", "Example for a 1,200 invoice and 100 recognized revenue: OpenRevRec derives debit billing clearing 1,200; credit deferred revenue 1,100; credit revenue 100."], ["3", "If the external invoice posting is debit accounts receivable 1,200 and credit the mapped billing clearing account 1,200, the clearing account nets to zero."], ["4", "If external invoices credit revenue or deferred revenue instead, map or transform the offset and reconcile before posting. Do not post both full outputs unchanged."], ["5", "Review profile dimensions, supplied account-combination preflight, and cross-dimension opening transfers against the destination ledger's actual posting rules."], ["6", "Record the external journal reference against this batch ID after posting. Reexporting the same batch is not a new posting instruction."], ["7", "If a posted close is revised, the replacement sheet compares each recorded source batch with this version. Confirm which batch represents the ledger before using its delta."], ["Boundary", "OpenRevRec does not post to the ledger or independently verify a user-entered posting reference."]]
    if report.get("runoff_unresolved", report.get("runoff_pending", [])):
        posting_steps.insert(0, ["Draft account runoff", "This journal depends on provisional historical-account allocations in this or an earlier month. Record and review each closing account split, then pass close checks before posting."])
    _sheet(book, "Posting guide", ["Step", "Treatment"], posting_steps)
    _sheet(book, "External posting records", ["Period", "Batch ID", "Close ID", "External journal reference", "Posted date", "Rationale", "Recorded at", "Current batch"],
           [[item.get("period", ""), item.get("batch_id", ""), item.get("close_id", ""), item.get("external_journal_reference", ""), item.get("posted_date", ""), item.get("rationale", ""), item.get("recorded_at", ""), item.get("batch_id") == batch_id if item.get("period") == report["period"] else ""] for item in state.get("postings", [])])
    comparisons = state.get("posting_comparisons", [])
    replacement_dimensions = sorted({key for item in comparisons for line in item["lines"] for key in line["dimensions"]})
    replacement_headers = ["Source batch", "Target batch", "Status", "Source references", "Contract ID", "Account", "Roles", "Obligation IDs", "Posted net", "Revised net", "Adjustment debit", "Adjustment credit"] + [f"dimension:{key}" for key in replacement_dimensions]
    replacement_rows = []
    for item in comparisons:
        status = "Source checkpoint unavailable" if not item["available"] else "Current batch already posted" if item["current_batch_posted"] else "Accepted replacement" if item["target_closed"] else "Draft comparison"
        for line in item["lines"] or [{}]:
            replacement_rows.append([item["source_batch_id"], item["target_batch_id"], status, ", ".join(item["posting_references"]), line.get("contract_id", ""), line.get("account", ""), ", ".join(line.get("roles", [])), ", ".join(line.get("obligation_ids", [])), Decimal(line["posted_net"]) if line else "", Decimal(line["revised_net"]) if line else "", Decimal(line["debit"]) if line else "", Decimal(line["credit"]) if line else ""] + [line.get("dimensions", {}).get(key, "") for key in replacement_dimensions])
    _sheet(book, "Replacement journal", replacement_headers, replacement_rows)
    _sheet(book, "Account policy", ["Role", "Account", "Policy version", "Effective period"], [[role, account, report.get("policy_version", state["policy"]["version"]), report.get("policy_effective_period", state["policy"]["effective_period"])] for role, account in report.get("policy_accounts", state["policy"]["accounts"]).items()])
    overrides = report.get("policy_account_overrides", {"contracts": {}, "obligations": {}})
    profiles = report.get("policy_account_profiles", {})
    assignments = report.get("policy_profile_assignments", {})
    _sheet(book, "Account profiles", ["Profile ID", "Profile name", "Role", "Account"],
           [[profile_id, profile["name"], role, account] for profile_id, profile in sorted(profiles.items()) for role, account in sorted(profile["accounts"].items())])
    _sheet(book, "Profile dimensions", ["Profile ID", "Dimension", "Value"],
           [[profile_id, key, value] for profile_id, profile in sorted(profiles.items()) for key, value in sorted(profile["dimensions"].items())])
    _sheet(book, "Profile assignments", ["Contract ID", "Profile ID"], [[contract_id, profile_id] for contract_id, profile_id in sorted(assignments.items())])
    obligation_assignments = report.get("policy_obligation_profile_assignments", {})
    _sheet(book, "Obligation profile assignments", ["Contract ID", "Obligation ID", "Profile ID"],
           [[contract_id, obligation_id, profile_id] for contract_id, items in sorted(obligation_assignments.items()) for obligation_id, profile_id in sorted(items.items())])
    _sheet(book, "Account overrides", ["Scope", "Contract ID", "Obligation ID", "Role", "Account"],
           [["Contract", contract_id, "", role, account] for contract_id, roles in overrides.get("contracts", {}).items() for role, account in roles.items()] +
           [["Obligation", contract_id, obligation_id, "revenue", account] for contract_id, items in overrides.get("obligations", {}).items() for obligation_id, account in items.items()])
    _sheet(book, "Account transitions", ["Contract ID", "Role", "From account", "To account", "Opening balance", "Treatment", "From dimensions", "To dimensions"],
           [[row[k] for k in ("contract_id", "role", "from_account", "to_account", "opening_balance", "treatment")] + [row.get("from_dimensions", {}), row.get("to_dimensions", {})] for row in report.get("account_transitions", [])])
    _sheet(book, "Account positions", ["Contract ID", "Role", "Account", "Dimensions", "Profile ID", "Projected closing balance"],
           [[row["contract_id"], row["role"], row["account"], row.get("dimensions", {}), row.get("account_profile_id") or "", Decimal(row["balance"])]
            for row in report.get("account_positions", [])])
    runoff_rows = []
    for item in state.get("runoff_allocations", []):
        if item["period"] == report["period"]:
            runoff_rows.extend([["Recorded", item["contract_id"], item["role"], item["period"], row["account"], row.get("dimensions", {}), Decimal(row["balance"]), item.get("rationale", ""), item.get("change_set_id", "")]
                                for row in item["positions"]])
    for item in report.get("runoff_unresolved", report.get("runoff_pending", [])):
        runoff_rows.extend([["Provisional - do not post", item["contract_id"], item["role"], item["period"], row["account"], row.get("dimensions", {}), Decimal(row["balance"]), "Requires reviewed allocation", ""]
                            for row in item["positions"]])
    _sheet(book, "Account runoff", ["Status", "Contract ID", "Role", "Period", "Account", "Dimensions", "Closing balance", "Rationale", "Change set ID"], runoff_rows)
    _sheet(book, "Allocation support", ["contract_id", "obligation_id", "name", "ssp", "amount"], [[c["id"], a["obligation_id"], a["name"], Decimal(a["ssp"]), Decimal(a["amount"])] for c in report["contracts"] for a in c["allocation"]])
    _sheet(book, "Component allocation", ["Contract ID", "Component ID", "Component", "Kind", "Included amount", "Scope", "Target obligation IDs", "Accounting rationale", "Target service month", "Targeted amount recognized to date", "Unit rate", "Pricing basis", "Rounding period", "Metered value mode"],
           [[contract["id"], item["component_id"], item["label"], item["kind"], Decimal(item["included_amount"]), item["scope"], ", ".join(item["target_obligation_ids"]), item["rationale"], item.get("target_period", ""), Decimal(item["recognized_to_date"]) if item.get("recognized_to_date") is not None else "", item.get("unit_rate", ""), item.get("pricing_basis", ""), item.get("rounding_period", ""), item.get("metered_value_mode", "")]
            for contract in report["contracts"] for item in contract.get("allocation_components", [])])
    _sheet(book, "Metered rate history", ["Contract ID", "Effective date", "Unit rate", "Accounting conclusion", "Activity ID"],
           [[contract["id"], item["effective_date"], item["unit_rate"], item["rationale"], item["activity_id"]]
            for contract in report["contracts"] for item in contract.get("metered_rate_history", [])])
    _sheet(book, "Metered usage valuation", ["Contract ID", "Activity ID", "Delivery date", "Calendar month", "Units", "Applied rate", "Unrounded value", "Value source", "Source reference", "Source agreement"],
           [[contract["id"], item["activity_id"], item["effective_date"], item["period"], item["quantity"], item.get("unit_rate", ""), item["unrounded_value"], item.get("value_source", "unit_rate"), item.get("reference", ""), item.get("source_contract_reference", "")]
            for contract in report["contracts"] for item in contract.get("metered_usage_valuation", [])])
    _sheet(book, "Metered monthly revenue", ["Contract ID", "Calendar month", "Units", "Unrounded value", "Recognized revenue"],
           [[contract["id"], item["period"], item["quantity"], item["unrounded_value"], Decimal(item["revenue"])]
            for contract in report["contracts"] for item in contract.get("metered_monthly_values", [])])
    _sheet(book, "Original promise changes", ["Contract ID", "Activity ID", "Effective date", "Component ID", "Component", "Original obligation ID", "Original obligation", "Allocated change", "Recognized to date", "Accounting rationale"],
           [[contract["id"], item.get("activity_id", ""), item["effective_date"], item["component_id"], item["component"], item["obligation_id"], item["obligation"], Decimal(item["allocated_change"]), Decimal(item["recognized_to_date"]), item["rationale"]]
            for contract in report["contracts"] for item in contract.get("original_promise_changes", [])])
    _sheet(book, "Activity", ["version", "change_set_id", "scenario_id", "command", "entity_id", "effective_date", "recorded_at", "rationale", "source", "payload"], [[r.get(k, "") for k in ["version", "id", "scenario_id", "command", "entity_id", "effective_date", "recorded_at", "rationale", "source", "payload"]] for r in reversed(state["change_sets"])])
    entity_names = {row["id"]: row["name"] for row in state["contracts"] + state["customers"]}
    period_end = report["period"] + "-31"
    changes = [row for row in reversed(state["change_sets"]) if row["effective_date"] <= period_end and row["command"] not in {"add_note", "edit_note", "edit_details", "attach_evidence", "record_judgment_review"}]
    evidence_by_change = {}
    for item in state["evidence"]:
        evidence_by_change.setdefault(item.get("target_change_set_id"), []).append(item)
    reviews_by_change = {}
    for item in state.get("judgment_reviews", []):
        reviews_by_change[item["target_change_set_id"]] = item
    _sheet(book, "Change lineage", ["Version", "Change set ID", "Scenario", "Command", "Entity", "Entity ID", "Effective date", "Recorded at", "Rationale", "Treatment", "Originating change ID", "Evidence files"],
           [[row["version"], row["id"], row["scenario_id"], row["command"], entity_names.get(row["entity_id"], row["entity_id"]), row["entity_id"], row["effective_date"], row["recorded_at"], row.get("rationale", ""), row["payload"].get("treatment", ""), row.get("originating_change_set_id", ""), ", ".join(item["name"] for item in evidence_by_change.get(row["id"], []))] for row in changes])
    _sheet(book, "Judgment support", ["Change set ID", "Command", "Entity", "Effective date", "Rationale", "Treatment", "Review disposition", "Reviewer", "Conclusion", "Support memo", "Exception reason", "Reviewed at", "Review change set ID", "Evidence IDs", "Evidence files"],
           [[row["id"], row["command"], entity_names.get(row["entity_id"], row["entity_id"]), row["effective_date"], row.get("rationale", ""), row["payload"].get("treatment", ""), reviews_by_change.get(row["id"], {}).get("disposition", "missing"), reviews_by_change.get(row["id"], {}).get("reviewer", ""), reviews_by_change.get(row["id"], {}).get("conclusion", ""), reviews_by_change.get(row["id"], {}).get("support_memo", ""), reviews_by_change.get(row["id"], {}).get("exception_reason", ""), reviews_by_change.get(row["id"], {}).get("recorded_at", ""), reviews_by_change.get(row["id"], {}).get("change_set_id", ""), ", ".join(item["id"] for item in evidence_by_change.get(row["id"], [])), ", ".join(item["name"] for item in evidence_by_change.get(row["id"], []))] for row in changes if row["command"] in JUDGMENT_COMMANDS])
    _sheet(book, "Evidence index", ["Evidence ID", "File name", "Entity", "Entity ID", "Change set ID", "Obligation ID", "Period close ID", "Recorded at", "Workspace path", "Rationale", "Scenario"],
           [[item["id"], item["name"], entity_names.get(item.get("entity_id"), item.get("entity_id", "")), item.get("entity_id", ""), item.get("target_change_set_id", ""), item.get("obligation_id", ""), item.get("period_close_id", ""), item.get("recorded_at", ""), item["path"], item.get("rationale", ""), item.get("scenario_id", "")] for item in state["evidence"]])
    _sheet(book, "Notes", ["id", "entity_id", "kind", "body", "due_date", "period", "completed"], [[n.get(k, "") for k in ["id", "entity_id", "kind", "body", "due_date", "period", "completed"]] for n in state["notes"]])
    _sheet(book, "Closes", ["period", "status", "frontier", "recorded_at", "rationale"], [[c.get(k, "") for k in ["period", "status", "frontier", "recorded_at", "rationale"]] for c in state["closes"]])
    _sheet(book, "Close dispositions", ["Period", "Change set ID", "Check", "Disposition", "Reason"],
           [[change["payload"].get("period", ""), change["id"], check_id, item.get("disposition", ""), item.get("reason", "")]
            for change in state["change_sets"] if change["command"] == "close_period"
            for check_id, item in change["payload"].get("review_dispositions", {}).items()])
    warning_details = report.get("warning_details", [])
    if len(warning_details) != len(report["warnings"]) or any(item["message"] != warning for item, warning in zip(warning_details, report["warnings"])):
        warning_details = [{"message": warning} for warning in report["warnings"]]
    _sheet(book, "Warnings", ["Contract ID", "Obligation ID", "Related contract ID", "Accounting review"],
           [[item.get("contract_id", ""), item.get("obligation_id", ""), item.get("related_contract_id", ""), item["message"]] for item in warning_details])
    output = io.BytesIO()
    book.save(output)
    return output.getvalue()


def export_package_bytes(state, review, workspace):
    """Bundle the workbook and every file named in its evidence index."""
    output = io.BytesIO()
    evidence = state.get("evidence", [])
    with zipfile.ZipFile(output, "w", compression=zipfile.ZIP_DEFLATED) as package:
        workbook = export_bytes(state, review)
        package.writestr("OpenRevRec-support.xlsx", workbook)
        files = {}
        archive_paths = {}
        for item in evidence:
            relative = item["path"]
            attachment = workspace.evidence_path(relative)
            archive_name = attachment.relative_to(workspace.path).as_posix()
            archive_paths[relative] = archive_name
            if archive_name not in files:
                content = attachment.read_bytes()
                package.writestr(archive_name, content)
                files[archive_name] = hashlib.sha256(content).hexdigest()
        package.writestr("manifest.json", json.dumps({
            "workspace_id": state["workspace"]["id"], "scenario_id": state["scenario_id"],
            "period": state["report"]["period"], "frontier": state["frontier"],
            "workbook_sha256": hashlib.sha256(workbook).hexdigest(),
            "files_sha256": files,
            "evidence": [{**{key: item.get(key) for key in ("id", "name", "path", "entity_id", "target_change_set_id", "recorded_at")}, "archive_path": archive_paths[item["path"]]} for item in evidence],
        }, indent=2))
    return output.getvalue()


def _value(value):
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, (int, float)):
        return str(value)
    return str(value).strip() if value is not None else None


def _component_row(values):
    result = dict(values)
    if "target_obligation_ids" in result:
        result["target_obligation_ids"] = [item.strip() for item in str(result["target_obligation_ids"]).split(",") if item.strip()]
    return result


def parse_workbook(data):
    try:
        book = load_workbook(io.BytesIO(data), read_only=True, data_only=False)
    except Exception as exc:
        raise ValueError("This file is not a readable .xlsx workbook.") from exc
    try:
        parsed = {}
        version = "1"
        if "Instructions" in book.sheetnames:
            version = str(book["Instructions"]["B2"].value)
            if str(version) not in {"1", TEMPLATE_VERSION}:
                raise ValueError(f"Unsupported template version {version}.")
        for name, headers in SHEETS.items():
            parsed[name] = []
            if name not in book.sheetnames:
                continue
            rows = iter(book[name].iter_rows())
            heading = next(rows, ())
            columns = [str(c.value or "").strip() for c in heading]
            named_columns = [c for c in columns if c]
            if len(named_columns) != len(set(named_columns)):
                raise ValueError(f"{name}: duplicate column headings are not allowed.")
            if any(h not in headers for h in named_columns):
                raise ValueError(f"{name}: unrecognized column. Download the current template.")
            for number, cells in enumerate(rows, 2):
                if all(c.value is None for c in cells):
                    continue
                if any(c.data_type == "f" for c in cells):
                    raise ValueError(f"{name} row {number}: formulas must be replaced with their reviewed values.")
                if any(not key and cell.value is not None for key, cell in zip(columns, cells)):
                    raise ValueError(f"{name} row {number}: every populated column requires a heading.")
                payload = {key: _value(cell.value) for key, cell in zip(columns, cells) if key and cell.value is not None}
                parsed[name].append((number, payload))
        commands = []
        for row, payload in parsed["Customers"]:
            commands.append(("Customers", row, "create_customer", payload))
        contract_ids = {c.get("id") for _, c in parsed["Contracts"]}
        for sheet in ("Consideration", "Obligations"):
            for row, item in parsed[sheet]:
                if item.get("contract_id") not in contract_ids:
                    raise ValueError(f"{sheet} row {row}: terms must reference a contract on the Contracts sheet; use a modification for an existing contract.")
        opening_dates = {}
        for row, opening in parsed["Opening Positions"]:
            contract_id = opening.get("contract_id")
            if contract_id in opening_dates:
                raise ValueError(f"Opening Positions row {row}: one opening position is allowed per contract.")
            opening_dates[contract_id] = opening.get("effective_date")
        for row, contract in parsed["Contracts"]:
            if contract.get("id") in opening_dates and not contract.get("cutover_date"):
                contract["cutover_date"] = opening_dates[contract["id"]]
            for field, sheet in (("consideration", "Consideration"), ("obligations", "Obligations")):
                contract[field] = [_component_row({k: v for k, v in p.items() if k != "contract_id"}) if field == "consideration" else {k: v for k, v in p.items() if k != "contract_id"}
                                   for _, p in parsed[sheet] if p.get("contract_id") == contract.get("id")]
            sources = [{k: v for k, v in item.items() if k != "contract_id"}
                       for _, item in parsed["Contract Sources"] if item.get("contract_id") == contract.get("id")]
            if sources:
                contract["source_contracts"] = sources
            commands.append(("Contracts", row, "create_contract", contract))
        for row, item in parsed["Contract Sources"]:
            if item.get("contract_id") not in contract_ids:
                raise ValueError(f"Contract Sources row {row}: contract_id must reference a Contracts row.")
        opening_ids = set()
        for row, opening in parsed["Opening Positions"]:
            source_id = opening.get("source_id")
            if not source_id or source_id in opening_ids:
                raise ValueError(f"Opening Positions row {row}: provide a unique, stable source_id.")
            opening_ids.add(source_id)
            opening["opening_obligations"] = [{key: value for key, value in item.items() if key != "opening_source_id"}
                                      for _, item in parsed["Opening Obligations"] if item.get("opening_source_id") == source_id]
            commands.append(("Opening Positions", row, "record_opening_position", opening))
        for row, item in parsed["Opening Obligations"]:
            if item.get("opening_source_id") not in opening_ids:
                raise ValueError(f"Opening Obligations row {row}: opening_source_id must reference an Opening Positions row.")
        amendment_ids = set()
        for row, amendment in parsed["Amendments"]:
            source_id = amendment.get("source_id")
            if not source_id or source_id in amendment_ids:
                raise ValueError(f"Amendments row {row}: provide a unique, stable source_id.")
            amendment_ids.add(source_id)
            payload = {key: amendment[key] for key in ("contract_id", "effective_date", "treatment", "rationale", "source_id", "term_basis", "term_assessment_rationale", "term_reassessment_trigger", "term_review_date") if key in amendment}
            replaced = False
            for flag, sheet, target in (("replace_consideration", "Amendment Consideration", "consideration"), ("replace_obligations", "Amendment Obligations", "obligations")):
                value = (amendment.get(flag) or "").casefold()
                if value not in {"", "no", "yes"}:
                    raise ValueError(f"Amendments row {row}: {flag} must be yes or no.")
                rows = [(number, item) for number, item in parsed[sheet] if item.get("amendment_source_id") == source_id]
                if rows and value != "yes":
                    raise ValueError(f"Amendments row {row}: set {flag} to yes when {sheet} contains replacement rows.")
                if value == "yes":
                    payload[target] = [_component_row({key: item for key, item in values.items() if key != "amendment_source_id"}) if target == "consideration" else {key: item for key, item in values.items() if key != "amendment_source_id"} for _, values in rows]
                    replaced = True
            mixed_rows = [(number, item) for number, item in parsed["Mixed Allocations"] if item.get("amendment_source_id") == source_id]
            if mixed_rows and payload.get("treatment") != "mixed":
                raise ValueError(f"Amendments row {row}: Mixed Allocations require mixed treatment.")
            if payload.get("treatment") == "mixed":
                payload["mixed_allocation"] = [{key: value for key, value in item.items() if key != "amendment_source_id"} for _, item in mixed_rows]
            if not replaced:
                raise ValueError(f"Amendments row {row}: choose at least one section to replace.")
            commands.append(("Amendments", row, "modify_contract", payload))
        for sheet in ("Amendment Consideration", "Amendment Obligations", "Mixed Allocations"):
            for row, item in parsed[sheet]:
                if item.get("amendment_source_id") not in amendment_ids:
                    raise ValueError(f"{sheet} row {row}: amendment_source_id must reference an Amendments row.")
        for row, p in parsed["Commands"]:
            try:
                payload = json.loads(p.get("payload_json", "{}"))
            except json.JSONDecodeError as exc:
                raise ValueError(f"Commands row {row}: payload_json is not valid JSON.") from exc
            if not isinstance(payload, dict):
                raise ValueError(f"Commands row {row}: payload_json must be an object.")
            if p.get("command") in {"close_period", "reopen_period", "create_scenario", "apply_scenario", "rebase_scenario", "archive_scenario", "restore_scenario"}:
                raise ValueError(f"Commands row {row}: review this workspace operation separately from an import.")
            if p.get("source_id"):
                payload["source_id"] = p["source_id"]
            commands.append(("Commands", row, p.get("command"), payload))
        for sheet, command in COMMAND_BY_SHEET.items():
            if sheet != "Customers":
                commands.extend((sheet, row, command, p) for row, p in parsed[sheet])
        for row, item in parsed["Corrections"]:
            activity_type = item.get("activity_type")
            if activity_type not in {"billing", "progress", "usage", "milestone", "rate_change"}:
                raise ValueError(f"Corrections row {row}: choose billing, progress, usage, milestone, or rate_change activity_type.")
            replacement = {key: item[key] for key in ("contract_id", "obligation_id", "component_id", "effective_date", "amount", "percentage", "quantity", "unit_rate", "invoice_value", "reference", "source_contract_reference") if key in item}
            payload = {"replacement": replacement, "rationale": item.get("rationale", ""), "source_id": item.get("source_id", "")}
            if item.get("target_change_set_id"):
                payload["target_change_set_id"] = item["target_change_set_id"]
            if item.get("target_source_id"):
                payload["target_source_id"] = item["target_source_id"]
                payload["target_command"] = f"record_{activity_type}"
            if bool(payload.get("target_change_set_id")) == bool(payload.get("target_source_id")):
                raise ValueError(f"Corrections row {row}: choose exactly one original change set ID or source_id.")
            commands.append(("Corrections", row, "correct_activity", payload))
        if not commands:
            raise ValueError("The workbook contains no rows to import.")
        return commands, version
    finally:
        book.close()


def _source_key(scenario_id, command, payload):
    source_id = payload.pop("source_id", None)
    if source_id is not None and not isinstance(source_id, str):
        raise ValueError("source_id must be text.")
    source_id = (source_id or "").strip()
    if source_id:
        identity = ["source_id", source_id]
    elif command == "record_billing" and payload.get("reference"):
        identity = ["invoice_reference", payload.get("contract_id"), payload.get("source_contract_reference"), payload["reference"]]
    else:
        identity = ["contents", dumps(payload)]
    return "import:" + hashlib.sha256(dumps([scenario_id, command, identity]).encode()).hexdigest(), source_id, identity


def _import_bytes(app, data, filename, scenario_id, period, preview, expected_frontier=None, expected_hash=None):
    from .application import compare_reports, valid_period

    commands, template_version = parse_workbook(data)
    results = []
    counts = {}
    input_billing = Decimal(0)
    input_units = Decimal(0)
    opening_positions = []
    selected_period = valid_period(period or date.today().strftime("%Y-%m"))
    file_hash = hashlib.sha256(data).hexdigest()
    if expected_hash is not None and expected_hash != file_hash:
        raise ValueError("The workbook changed since review. Preview it again.")
    with app.workspace.connect() as db:
        db.execute("BEGIN IMMEDIATE")
        try:
            before = app._state(db, scenario_id, selected_period)
            if expected_frontier is not None and str(expected_frontier) != str(before["frontier"]):
                raise ValueError("The workspace changed since review. Preview the workbook again.")
            backup = None if preview else app.workspace.backup("import")
            for sheet, row, command, payload in commands:
                try:
                    payload = dict(payload)
                    if sheet == "Corrections" and payload.get("target_source_id"):
                        target_source_id = payload.pop("target_source_id")
                        target_command = payload.pop("target_command")
                        target_key, _, _ = _source_key(scenario_id, target_command, {"source_id": target_source_id})
                        target = db.execute("SELECT id FROM change_sets WHERE idempotency_key=?", (target_key,)).fetchone()
                        if target is None:
                            raise ValueError("Original source_id was not found in this scenario; import the original activity first")
                        payload["target_change_set_id"] = target["id"]
                    if template_version == TEMPLATE_VERSION and sheet in {"Opening Positions", "Rate Changes", "Progress", "Usage", "Right Exercises", "Renewal Links", "Modification Links", "Milestones", "Adjustments", "Reassessments", "Amendments", "Corrections", "Commands"} and not payload.get("source_id"):
                        raise ValueError("This activity requires a stable source_id")
                    if template_version == TEMPLATE_VERSION and sheet == "Billing" and not (payload.get("source_id") or payload.get("reference")):
                        raise ValueError("Billing requires a source_id or invoice reference")
                    key, source_id, identity = _source_key(scenario_id, command, payload)
                    if db.execute("SELECT 1 FROM change_sets WHERE idempotency_key=?", (key,)).fetchone():
                        raise ValueError("This source row was already imported. Use a source correction for changed facts; give genuinely distinct identical rows unique source_id values")
                    if scenario_id != "main":
                        main_key = "import:" + hashlib.sha256(dumps(["main", command, identity]).encode()).hexdigest()
                        if db.execute("SELECT 1 FROM change_sets WHERE idempotency_key=?", (main_key,)).fetchone():
                            raise ValueError("This source row is already accepted in Main")
                    payload["import_source_identity"] = identity
                    request_hash = hashlib.sha256(dumps({"command": command, "payload": payload, "scenario_id": scenario_id}).encode()).hexdigest()
                    result = app._execute_in(db, command, payload, scenario_id, f"excel:{filename}:v{template_version}:{sheet}:{row}", key, request_hash)
                    results.append({"sheet": sheet, "row": row, "command": command, "source_id": source_id, "effective_date": payload.get("effective_date") or payload.get("start_date") or "", "status": "accepted", **result})
                    counts[command] = counts.get(command, 0) + 1
                    if command == "record_billing":
                        input_billing += Decimal(payload["amount"])
                    elif command == "record_usage":
                        input_units += Decimal(payload["quantity"])
                    elif command == "record_opening_position":
                        opening_positions.append({
                            "contract_id": payload["contract_id"], "cutover_date": payload["effective_date"],
                            "recognized_to_date": f"{sum((Decimal(item['recognized_to_date']) for item in payload['opening_obligations']), Decimal(0)):.2f}",
                            "billed_to_date": payload["billed_to_date"], "contract_asset": payload["contract_asset"],
                            "deferred_revenue": payload["deferred_revenue"], "source_name": payload["source_name"],
                            "obligations": payload["opening_obligations"],
                        })
                except (ValueError, KeyError, TypeError) as exc:
                    raise ValueError(f"{sheet} row {row}: {exc}. No accounting rows were imported.") from exc
            state = app._state(db, scenario_id, selected_period)
            comparison = compare_reports(before["report"], state["report"])
            candidate_periods = {selected_period, *comparison["affected_periods"]}
            candidate_periods.update(row["effective_date"][:7] for row in results if row["effective_date"])
            period_impacts = []
            for candidate in sorted(candidate_periods):
                prior_report = before["report"] if candidate == selected_period else app._state(db, scenario_id, candidate, before["frontier"])["report"]
                next_report = state["report"] if candidate == selected_period else app._state(db, scenario_id, candidate)["report"]
                delta = {key: f"{Decimal(next_report['summary'][key]) - Decimal(prior_report['summary'][key]):.2f}" for key in ("revenue", "billings", "deferred_revenue", "contract_asset", "remaining_revenue")}
                if any(Decimal(value) != 0 for value in delta.values()) or prior_report["journals"] != next_report["journals"]:
                    period_impacts.append({"period": candidate, "delta": delta, "journal_changed": prior_report["journals"] != next_report["journals"]})
            comparison["affected_periods"] = sorted(set(comparison["affected_periods"]) | {row["period"] for row in period_impacts})
            result = {"imported": len(results), "rows": results, "template_version": template_version, "file_hash": file_hash,
                      "controls": {"counts": counts, "source_billing_total": f"{input_billing:.2f}", "source_usage_quantity": str(input_units), "opening_positions": opening_positions}}
            if preview:
                return {"result": result, "before": before["report"], "state": state, "comparison": comparison, "period_impacts": period_impacts, "frontier": before["frontier"]}
            db.commit()
            return {"result": {**result, "backup_path": str(backup)}, "state": state}
        finally:
            if db.in_transaction:
                db.rollback()


def preview_import_bytes(app, data, filename="import.xlsx", scenario_id="main", period=None):
    return _import_bytes(app, data, filename, scenario_id, period, True)


def import_bytes(app, data, filename="import.xlsx", scenario_id="main", period=None, expected_frontier=None, expected_hash=None):
    return _import_bytes(app, data, filename, scenario_id, period, False, expected_frontier, expected_hash)
